from __future__ import annotations

from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook
import pyarrow.parquet as pq
import pytest

from import_profit import build_historical_inputs_candidate as candidate
from test_import_profit_historical_cnf_adapter import GROUPS, workbook
from test_import_profit_historical_dce_adapter import sql_fixture


CONFIG = candidate.REPOSITORY_ROOT / "02_configs" / "import_profit_soybean.yaml"
FIXED_TIME = datetime(2026, 7, 30, 12, 0, tzinfo=timezone.utc)
SAMPLE = "2026-06-10,巴西,2026-12"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def full_cnf(path: Path, business_date: date = date(2026, 6, 10)) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "CNF报价"
    sheet.cell(1, 1, "日期")
    for label, start in GROUPS:
        for month in range(1, 13):
            sheet.cell(1, start + month - 1, f"{label}{month}月")
            sheet.cell(2, start + month - 1, start * 10 + month)
    sheet.cell(2, 1, business_date)
    book.save(path)
    return path


def good_sources(tmp_path: Path, *, warning: bool) -> tuple[Path, Path]:
    if warning:
        excel = workbook(
            tmp_path / "cnf.xlsx",
            dates=(date(2026, 6, 10), date(2026, 6, 11)),
        )
        sql = sql_fixture(tmp_path / "dce.sql")
    else:
        excel = full_cnf(tmp_path / "cnf.xlsx")
        sql = sql_fixture(
            tmp_path / "dce.sql",
            rows=[["2026-06-10", 3001, 3005, 3009, 8001, 8005, 8009]],
        )
    return excel, sql


def test_candidate_with_sample_has_five_atomic_outputs_and_safe_manifest(
    tmp_path: Path,
) -> None:
    excel, sql = good_sources(tmp_path, warning=True)
    excel_before, sql_before = sha(excel), sha(sql)
    output = tmp_path / "candidate"
    result = candidate.build_historical_inputs_candidate(
        excel,
        sql,
        CONFIG,
        output,
        sample_keys=[SAMPLE],
        generated_at=FIXED_TIME,
    )

    assert sha(excel) == excel_before
    assert sha(sql) == sql_before
    assert set(path.name for path in output.iterdir()) == {
        candidate.CNF_FILENAME,
        candidate.DCE_FILENAME,
        candidate.RESOLVED_FILENAME,
        candidate.MANIFEST_FILENAME,
        candidate.QUALITY_FILENAME,
    }
    manifest = json.loads((output / candidate.MANIFEST_FILENAME).read_text("utf-8"))
    quality = json.loads((output / candidate.QUALITY_FILENAME).read_text("utf-8"))
    serialized = json.dumps(manifest, ensure_ascii=False)
    assert str(tmp_path) not in serialized
    assert manifest["synthetic_input"] is False
    assert manifest["source_excel_filename"] == excel.name
    assert manifest["source_sql_filename"] == sql.name
    assert manifest["candidate_status"] == "passed_with_warnings"
    assert quality["candidate_status"] == "passed_with_warnings"
    assert manifest["resolved_sample_record_count"] == 2
    sample = manifest["sample_key_results"][0]
    assert sample["soymeal_contract"] == "M2701"
    assert sample["soyoil_contract"] == "Y2701"
    prices = {
        leg["contract_code"]: leg["price_cny_per_tonne"]
        for leg in sample["dce_legs"]
    }
    assert prices == {"M2701": 3001.0, "Y2701": 8001.0}
    assert pq.read_table(output / candidate.CNF_FILENAME).num_rows == 96
    assert pq.read_table(output / candidate.DCE_FILENAME).num_rows == 10
    resolved = pq.read_table(output / candidate.RESOLVED_FILENAME)
    assert resolved.num_rows == 2
    assert set(resolved.column("contract_identity_status").to_pylist()) == {
        "continuous_inferred"
    }
    assert set(resolved.column("source_delivery_month").to_pylist()) == {1}
    assert set(resolved.column("source_contract_code").to_pylist()) == {None}
    assert result["manifest"] == manifest


def test_candidate_without_sample_has_exactly_four_outputs(tmp_path: Path) -> None:
    excel, sql = good_sources(tmp_path, warning=True)
    output = tmp_path / "candidate"
    result = candidate.build_historical_inputs_candidate(
        excel, sql, CONFIG, output, generated_at=FIXED_TIME
    )
    assert set(path.name for path in output.iterdir()) == {
        candidate.CNF_FILENAME,
        candidate.DCE_FILENAME,
        candidate.MANIFEST_FILENAME,
        candidate.QUALITY_FILENAME,
    }
    assert candidate.RESOLVED_FILENAME not in result["files"]
    assert result["manifest"]["sample_keys"] == []
    assert result["manifest"]["resolved_sample_record_count"] == 0


def test_complete_sources_can_pass_without_warnings(tmp_path: Path) -> None:
    excel, sql = good_sources(tmp_path, warning=False)
    result = candidate.build_historical_inputs_candidate(
        excel,
        sql,
        CONFIG,
        tmp_path / "candidate",
        generated_at=FIXED_TIME,
    )
    assert result["manifest"]["candidate_status"] == "passed"
    assert result["quality_report"]["candidate_status"] == "passed"
    assert result["quality_report"]["warnings"] == []


def test_failed_adapter_produces_no_partial_candidate(tmp_path: Path) -> None:
    excel, _ = good_sources(tmp_path, warning=False)
    sql = sql_fixture(
        tmp_path / "bad.sql",
        rows=[["2026-06-10", 0, 2, 3, 4, 5, 6]],
    )
    output = tmp_path / "candidate"
    with pytest.raises(candidate.HistoricalCandidateBuildError) as error:
        candidate.build_historical_inputs_candidate(
            excel, sql, CONFIG, output, generated_at=FIXED_TIME
        )
    assert error.value.quality_report["candidate_status"] == "failed"
    assert not output.exists()


def test_atomic_replace_failure_cleans_all_partial_files(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    excel, sql = good_sources(tmp_path, warning=False)
    output = tmp_path / "candidate"
    real_replace = candidate.os.replace
    calls = 0

    def failing_replace(source: Path, target: Path) -> None:
        nonlocal calls
        calls += 1
        if calls == 2:
            raise OSError("simulated replace failure")
        real_replace(source, target)

    monkeypatch.setattr(candidate.os, "replace", failing_replace)
    with pytest.raises(OSError, match="simulated"):
        candidate.build_historical_inputs_candidate(
            excel, sql, CONFIG, output, generated_at=FIXED_TIME
        )
    assert not output.exists()


@pytest.mark.parametrize(
    "value",
    [
        "2026-06-10,未知,2026-12",
        "2026/06/10,巴西,2026-12",
        "2026-06-10,巴西,2026-13",
    ],
)
def test_invalid_sample_key_is_rejected(tmp_path: Path, value: str) -> None:
    excel, sql = good_sources(tmp_path, warning=False)
    with pytest.raises(candidate.HistoricalCandidateBuildError):
        candidate.build_historical_inputs_candidate(
            excel,
            sql,
            CONFIG,
            tmp_path / "candidate",
            sample_keys=[value],
            generated_at=FIXED_TIME,
        )


def test_output_inside_repository_is_rejected(tmp_path: Path) -> None:
    excel, sql = good_sources(tmp_path, warning=False)
    with pytest.raises(candidate.HistoricalCandidateBuildError, match="outside"):
        candidate.build_historical_inputs_candidate(
            excel,
            sql,
            CONFIG,
            candidate.REPOSITORY_ROOT / "forbidden-candidate",
            generated_at=FIXED_TIME,
        )
