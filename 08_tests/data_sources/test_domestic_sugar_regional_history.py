from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from agri_research_agent.data_sources.domestic_sugar_history import SugarHistoryError, file_sha256
from agri_research_agent.data_sources.domestic_sugar_regional_history import (
    REGIONS, build_regional_basis, load_regional_history, resolve_regional_history,
)
from agri_research_agent.pipelines.domestic_sugar_regional_history import prepare_regional_sugar_history

# Collection can start in this directory before pytest has imported root tests.
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
pytest_plugins = ["data_sources.test_domestic_sugar_history"]


@pytest.fixture
def regional_workbook(tmp_path):
    path = tmp_path / "钢联三地.xlsx"
    w = openpyxl.Workbook(); s = w.active; s.title = "导出数据"
    s.append(["钢联数据"])
    s.append(["指标名称", *[f"白糖：现货价：{region}（日）" for region in REGIONS]])
    s.append(["频度", "日", "日", "日"])
    s.append(["指标描述", *["当地贸易商报价平均值"] * 3])
    for date, prices in [
        ("2023-01-03", [6100, 6000, 6050]), ("2024-02-29", [6000, 5900, 5950]),
        ("2026-09-30", [5180, 5100, 5160]), ("2026-10-09", [5280, 5180, 5230]),
        ("2026-10-10", [5330, None, 5270]), ("2026-10-12", [5330, 5215, 5290]),
    ]:
        s.append([dt.datetime.fromisoformat(date), *prices])
    w.save(path); w.close()
    return path


@pytest.fixture
def regional_root(regional_workbook, prepared_root, tmp_path):
    root = tmp_path / "regional"
    prepare_regional_sugar_history(regional_workbook, file_sha256(regional_workbook), prepared_root, root)
    return root


def test_exact_same_day_shared_contract_and_missing_weekend(regional_root):
    daily = load_regional_history(regional_root).daily
    oct9 = daily.loc[daily.business_date.eq("2026-10-09")].set_index("region")
    assert oct9.ts_code.eq("SR2701.ZCE").all()
    assert oct9.futures_price.eq(5491).all()
    assert oct9.basis.to_dict() == {"乌鲁木齐": -261, "昆明": -311, "柳州": -211}
    weekend = daily.loc[daily.business_date.eq("2026-10-10")]
    assert weekend.basis.isna().all() and weekend.futures_price.isna().all()
    assert pd.isna(weekend.set_index("region").loc["昆明", "cash_price"])
    early = daily.loc[daily.business_date.eq("2023-01-03")]
    assert early.ts_code.eq("SR2303.ZCE").all()
    assert len(daily) == 18


def test_preparation_preserves_both_inputs_and_is_idempotent(regional_root, regional_workbook, prepared_root):
    spot_hash = file_sha256(regional_workbook)
    old_pointer = file_sha256(prepared_root / "current.json")
    old_files = {str(p): file_sha256(p) for p in prepared_root.rglob("*") if p.is_file()}
    regional_files = {str(p): file_sha256(p) for p in regional_root.rglob("*") if p.is_file()}
    assert prepare_regional_sugar_history(regional_workbook, spot_hash, prepared_root, regional_root)["status"] == "NO_CHANGE"
    assert file_sha256(regional_workbook) == spot_hash and file_sha256(prepared_root / "current.json") == old_pointer
    assert old_files == {str(p): file_sha256(p) for p in prepared_root.rglob("*") if p.is_file()}
    assert regional_files == {str(p): file_sha256(p) for p in regional_root.rglob("*") if p.is_file()}


@pytest.mark.parametrize("cell,value,reason", [("B4", "含税车板价", "口径"), ("C2", "白糖：一级：市场价：昆明（日）", "指标"), ("D8", 0, "价格"), ("B8", "5280-5290", "数值")])
def test_source_definition_and_invalid_prices_are_rejected(regional_workbook, prepared_root, tmp_path, cell, value, reason):
    w = openpyxl.load_workbook(regional_workbook); w.active[cell] = value; w.save(regional_workbook); w.close()
    root = tmp_path / "bad"
    with pytest.raises(SugarHistoryError, match=reason):
        prepare_regional_sugar_history(regional_workbook, file_sha256(regional_workbook), prepared_root, root)
    assert not root.exists()


def test_tampered_content_is_rejected(regional_root):
    path, _ = resolve_regional_history(regional_root)
    path.write_bytes(path.read_bytes() + b" ")
    with pytest.raises(SugarHistoryError, match="身份"):
        load_regional_history(regional_root)


def test_duplicate_mapping_rejected_and_zero_basis_preserved(regional_root):
    path, _ = resolve_regional_history(regional_root)
    p = json.loads(path.read_text(encoding="utf-8"))
    spots, main = pd.DataFrame(p["spot_records"]), pd.DataFrame(p["main_records"])
    spots.loc[spots.business_date.eq("2026-10-09") & spots.region.eq("柳州"), "cash_price"] = 5491
    result = build_regional_basis(spots, main)
    assert result.loc[result.business_date.eq("2026-10-09") & result.region.eq("柳州"), "basis"].iloc[0] == 0
    with pytest.raises(SugarHistoryError, match="重复"):
        build_regional_basis(spots, pd.concat([main, main.iloc[[0]]]))


@pytest.mark.parametrize("field,value,reason", [("source_cell", "C5", "单元格"), ("cash_price", True, "布尔")])
def test_wrong_source_column_and_boolean_prices_rejected(regional_root, field, value, reason):
    path, _ = resolve_regional_history(regional_root)
    payload = json.loads(path.read_text(encoding="utf-8"))
    spots = pd.DataFrame(payload["spot_records"]).astype(object)
    spots.loc[0, field] = value
    with pytest.raises(SugarHistoryError, match=reason):
        build_regional_basis(spots, pd.DataFrame(payload["main_records"]))


def extension(path, close=5518, date="20261012"):
    p = dict(mapping=dict(code=0, data=dict(fields=["ts_code", "trade_date", "mapping_ts_code"], items=[["SR.ZCE", date, "SR2701.ZCE"]])),
        daily=dict(code=0, data=dict(fields=["ts_code", "trade_date", "close"], items=[["SR2701.ZCE", date, close]])))
    path.write_text(json.dumps(p), encoding="utf-8")
    return path


def test_extension_uses_actual_mapped_close_and_rejects_historical_conflict(regional_workbook, prepared_root, tmp_path):
    extra = extension(tmp_path / "extension.json")
    root = tmp_path / "extended"
    prepare_regional_sugar_history(regional_workbook, file_sha256(regional_workbook), prepared_root, root,
        main_extension=extra, extension_sha256=file_sha256(extra))
    rows = load_regional_history(root).daily
    oct12 = rows.loc[rows.business_date.eq("2026-10-12")].set_index("region")
    assert oct12.basis.to_dict() == {"乌鲁木齐": -228, "昆明": -303, "柳州": -188}
    extension(extra, close=5492, date="20261009")
    with pytest.raises(SugarHistoryError, match="冲突"):
        prepare_regional_sugar_history(regional_workbook, file_sha256(regional_workbook), prepared_root, tmp_path / "conflict",
            main_extension=extra, extension_sha256=file_sha256(extra))


def test_different_snapshot_does_not_replace_candidate(regional_root, regional_workbook, prepared_root):
    original = file_sha256(regional_root / "current.json")
    w = openpyxl.load_workbook(regional_workbook); w.active["B8"] = 5300; w.save(regional_workbook); w.close()
    with pytest.raises(SugarHistoryError, match="不同快照"):
        prepare_regional_sugar_history(regional_workbook, file_sha256(regional_workbook), prepared_root, regional_root)
    assert file_sha256(regional_root / "current.json") == original
