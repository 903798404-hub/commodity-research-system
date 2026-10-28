from io import BytesIO
from pathlib import Path
import json
import pytest
from PIL import Image
from agri_research_agent.pipelines.conab_growth_chart import read_growth_chart
from agri_research_agent.pipelines.brazil_soy_scheduled import report_links,weekly_assets,parse_growth

ROOT=Path(__file__).resolve().parents[1]

@pytest.mark.parametrize('day,expected',[
    ('20261002',{'EMERGENCE':64.1,'VEGETATIVE':35.9}),
    ('20261009',{'EMERGENCE':43.,'VEGETATIVE':57.}),
    ('20261016',{'EMERGENCE':52.2,'VEGETATIVE':47.8})])
def test_printed_official_values_not_slice_geometry(day,expected):
    values=read_growth_chart((ROOT/'fixtures/conab-growth'/f'{day}.png').read_bytes())
    assert values==expected
    assert 'FLOWERING' not in values  # An unprinted stage is never zero-filled.

def test_blank_or_unknown_font_does_not_publish():
    buffer=BytesIO();Image.new('RGB',(1040,1070),'white').save(buffer,format='PNG')
    with pytest.raises(ValueError):read_growth_chart(buffer.getvalue())

def test_unknown_legend_does_not_relabel_stages():
    image=Image.open(ROOT/'fixtures/conab-growth/20261016.png').convert('RGB')
    image.paste('white',(0,930,image.width,image.height))
    buffer=BytesIO();image.save(buffer,format='PNG')
    with pytest.raises(ValueError):read_growth_chart(buffer.getvalue())

def test_official_publication_is_distinct_from_reporting_period():
    raw=b'''<span class="documentPublished">Publicado em <span>19/10/2026 19h00</span></span>
    <a href="/conab/report/plantio">Plantio e Colheita 10-10 a 16-10</a>
    <a href="/conab/report/feno">Fenologia das Lavouras 10-10 a 16-10</a>'''
    pub,assets=weekly_assets(raw,'https://www.gov.br/conab/report')
    assert pub=='2026-10-19' and len(assets)==2

@pytest.mark.parametrize('raw',[
    b'<a href="https://evil.example/report">Plantio e Colheita</a>',
    b'<span class="documentPublished">2026-10-19</span>',
    b'<span class="documentPublished">19/10/2026</span><a href="/conab/a">Plantio e Colheita</a>'
])
def test_unapproved_or_incomplete_assets_are_rejected(raw):
    with pytest.raises(ValueError):weekly_assets(raw,'https://www.gov.br/conab/')

def test_index_uses_official_period_and_rejects_external_host():
    rows=report_links(b'<a href="/conab/w">Acompanhamento das Lavouras - 12/10 a 18/10/26</a>','https://www.gov.br/conab/')
    assert rows[0][1].isoformat()=='2026-10-18'
    with pytest.raises(ValueError):report_links(b'<a href="https://evil.example/w">Acompanhamento das Lavouras - 12/10 a 18/10/26</a>','https://www.gov.br/conab/')


def test_harvest_season_does_not_replay_old_planting_weeks(tmp_path,monkeypatch):
    from datetime import date
    from agri_research_agent.pipelines import brazil_soy_scheduled as updater
    baseline={'records':[{'region':'BR','metric':'PLANTED','date':'2026-11-20'},
                         {'region':'BR','metric':'HARVESTED','date':'2027-03-05'}]}
    monkeypatch.setattr(updater,'load_bundle',lambda path:baseline)
    monkeypatch.setattr(updater,'download',lambda *args:b'index')
    monkeypatch.setattr(updater,'report_links',lambda *args:[('https://www.gov.br/conab/january',date(2027,1,17))])
    monkeypatch.setattr(updater,'latest_dates',lambda bundle:{'BR/HARVESTED':'2027-03-05'})
    monkeypatch.setattr(updater,'weekly_assets',lambda *args:pytest.fail('old planting-era report must not be replayed'))
    bundle,evidence,semantic=updater.collect(tmp_path/'baseline',tmp_path/'owned',today=date(2027,3,10))
    assert bundle is baseline and not semantic['business_changed'] and not evidence['sources']


def test_original_growth_image_keeps_values_and_actual_locator():
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SheetImage
    title='Progresso de safra'
    book=Workbook();sheet=book.active;sheet.title=title
    sheet['B96']='Soja - Safra 2026/27'
    sheet['B97']='Semana de 10 a 16 de outubro de 2026'
    sheet.add_image(SheetImage(ROOT/'fixtures/conab-growth/20261016.png'),'B101')
    raw=BytesIO();book.save(raw);book.close()
    source={'source_url':'https://www.gov.br/conab/relatorio.xlsx','sha256':'a'*64,
            'retrieved_at':'2026-10-28T01:00:00Z','published_at':'2026-10-19'}
    rows=parse_growth(raw.getvalue(),source)
    assert {r['metric']:r['value'] for r in rows}=={'EMERGENCE':52.2,'VEGETATIVE':47.8}
    assert all(r['date']=='2026-10-16' and r['source_locator'].startswith(title+'!B96:') for r in rows)


def test_native_official_growth_chart_retains_printed_precision_and_missing_labels():
    import hashlib
    folder=Path(__file__).parent/'fixtures/conab-20261023'
    source=json.loads((folder/'source.json').read_bytes());raw=(folder/'growth.xlsx').read_bytes()
    assert hashlib.sha256(raw).hexdigest()==source['sha256']
    rows=parse_growth(raw,source)
    assert {r['metric']:r['value'] for r in rows}=={'EMERGENCE':39.7,'VEGETATIVE':60.3}
    assert all(r['date']=='2026-10-23' and r['season']=='2026/2027' for r in rows)
    assert all(r['source_locator'].startswith('Fenologia das Lavouras!J') for r in rows)


@pytest.mark.parametrize('change',['formula','blank','label','format','cache_mismatch','wrong_range','duplicate_chart'])
def test_native_growth_rejects_changed_or_inconsistent_source(change):
    from openpyxl import load_workbook
    from agri_research_agent.pipelines.brazil_soy_scheduled import native_growth_values
    book=load_workbook(Path(__file__).parent/'fixtures/conab-20261023/growth.xlsx',data_only=False)
    try:
        sheet=book.active
        if change=='formula':sheet['K96']='=0.396811'
        if change=='blank':sheet['K96']=None
        if change=='label':sheet['J96']='Unknown'
        if change=='format':sheet['K96'].number_format='General'
        if change=='cache_mismatch':sheet['K96']=0.4
        if change=='wrong_range':sheet._charts[3].series[0].val.numRef.f="'Fenologia das Lavouras'!$K$6:$K$11"
        if change=='duplicate_chart':sheet._charts.append(sheet._charts[3])
        with pytest.raises(ValueError):native_growth_values(sheet,96)
    finally:book.close()


@pytest.mark.parametrize('titles',[['Unknown'],['Fenologia das Lavouras','Progresso de safra']])
def test_unknown_or_multiple_growth_sheets_still_rejected(titles):
    from openpyxl import Workbook
    book=Workbook();book.active.title=titles[0]
    for title in titles[1:]:book.create_sheet(title)
    raw=BytesIO();book.save(raw);book.close()
    with pytest.raises(ValueError,match='workbook layout changed'):
        parse_growth(raw.getvalue(),{})
