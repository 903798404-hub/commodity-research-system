from __future__ import annotations

from copy import deepcopy
from datetime import date
import importlib
import io
from pathlib import Path
import sys

from openpyxl import Workbook
import pytest
from streamlit.testing.v1 import AppTest

from agri_research_agent.automation import production_data_delta_brazil as producer
from agri_research_agent.pipelines import brazil_soy as soy
from agri_research_agent.pipelines import brazil_soy_update as update
from agri_research_agent.shared.atomic_storage import atomic_write_json
from agri_research_agent.shared.runtime_context import RuntimeContext, RuntimeMode

ROOT = Path(__file__).resolve().parents[2]
STAMP = "2026-10-07T00:00:00+00:00"


def area_source(tmp_path):
    archive = tmp_path / "raw/brazil_soy/annual-area"
    archive.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Soja"
    ws["A5"], ws["B5"], ws["C6"] = "REGIÃO/UF", "ÁREA (Em mil ha)", "Safra 25/26"
    uf = "AC AL AP AM BA CE DF ES GO MA MT MS MG PA PB PR PE PI RJ RN RS RO RR SC SP SE TO".split()
    for row, state in enumerate(uf, 9):
        ws.cell(row, 1, state)
        ws.cell(row, 3, 200 if state == "MT" else 100)
        ws.cell(row, 9, 9000)  # Production must never be used as the area column.
    ws["A42"], ws["C42"] = "BRASIL", 2800
    raw = io.BytesIO()
    wb.save(raw)
    wb.close()
    (archive / "report.bin").write_bytes(raw.getvalue())
    source = {"schema_version": "brazil-soy-source/1", "source_url": soy.SOURCE_URL,
              "final_url": soy.SOURCE_URL, "sha256": soy.sha256_file(archive / "report.bin"),
              "published_at": "2026-09-15", "retrieved_at": STAMP}
    atomic_write_json(archive / "source.json", source)
    reference = soy.parse_area_reference(io.BytesIO(raw.getvalue()), source)
    return archive, reference


def progress_bundle():
    row = soy.observation("2026/2027", "BR", "PLANTED", date(2026, 10, 2), 9.4,
        source_url=soy.SOURCE_URL, source_sha256="a" * 64, locator="table 1", retrieved_at=STAMP)
    return {"schema_version": soy.SCHEMA, "generated_at": STAMP, "records": [row], "import_notes": []}


def test_area_denominator_is_all_brazil_and_reconciles_uf(tmp_path):
    archive, reference = area_source(tmp_path)
    assert soy.area_share(reference, "MT") == pytest.approx(200 / 2800 * 100)
    assert soy.area_share(reference, "BR") == pytest.approx(1300 / 2800 * 100)
    assert reference["season"] == "2025/2026" and reference["unit"] == "thousand_hectares"
    from openpyxl import load_workbook
    with (archive / "report.bin").open("rb") as handle:
        wb = load_workbook(handle)
    wb["Soja"]["C42"] = 1300  # Incorrectly normalized to the twelve states.
    raw = io.BytesIO()
    wb.save(raw)
    wb.close()
    with pytest.raises(ValueError, match="reconcile"):
        soy.parse_area_reference(io.BytesIO(raw.getvalue()), reference | {"sha256": "a" * 64})


@pytest.mark.parametrize("field,value", [("unit", "tonnes"), ("national_area", 0),
    ("national_area", True), ("published_at", "2099-09-15"), ("source_url", "https://example.com")])
def test_area_reference_rejects_invalid_units_identity_and_denominator(tmp_path, field, value):
    _, reference = area_source(tmp_path)
    reference[field] = value
    with pytest.raises(ValueError):
        soy.validate_bundle(progress_bundle() | {"area_reference": reference})


def test_area_preparation_preserves_progress_and_weekly_merge_preserves_reference(tmp_path):
    archive, reference = area_source(tmp_path)
    atomic_write_json(tmp_path / ".market-data-runtime.json", {"schema_version": 1, "runtime_id": "brazil-area-test",
        "classification": "fixture", "module_id": "brazil-soy", "created_at": STAMP})
    context = RuntimeContext(RuntimeMode.FIXTURE, "brazil-soy", tmp_path)
    original = progress_bundle()
    baseline = tmp_path / soy.STABLE_RELATIVE_PATH
    baseline.parent.mkdir(parents=True)
    atomic_write_json(baseline, original)
    candidate = update.prepare_area_reference(context, baseline, archive / "source.json")
    actual = soy.load_bundle(candidate)
    assert actual["records"] == original["records"] and actual["area_reference"] == reference
    merged, _ = update.merge_observations(actual, original["records"])
    assert merged["area_reference"] == reference
    update.activate_local(context, candidate)
    assert soy.load_bundle(baseline)["area_reference"] == reference


def test_publisher_replays_area_source_and_rejects_tampering_deletion_or_unused_evidence(tmp_path):
    _, reference = area_source(tmp_path)
    original = progress_bundle()
    baseline = tmp_path / "baseline.json"
    candidate = tmp_path / "candidate.json"
    atomic_write_json(baseline, original)
    updated = original | {"area_reference": reference}
    atomic_write_json(candidate, updated)
    config = {"source_root": str(tmp_path), "workbook_path": None, "workbook_sha256": None}
    evidence, inventory = producer.source_evidence(config, updated, original)
    assert evidence["sources"][0]["kind"] == "area_report" and len(inventory) == 2
    host = producer.delivery._host_contract()
    result = host._brazil_observations(candidate, baseline, evidence, [])
    assert result["business_changed"] and result["added"] == result["revised"] == 0
    tampered = deepcopy(updated)
    tampered["area_reference"]["state_areas"]["MT"] = 199
    atomic_write_json(candidate, tampered)
    with pytest.raises(host.DeltaError, match="differs"):
        host._brazil_observations(candidate, baseline, evidence, [])
    atomic_write_json(baseline, updated)
    atomic_write_json(candidate, original)
    with pytest.raises(host.DeltaError, match="deletion"):
        host._brazil_observations(candidate, baseline, evidence, [])
    atomic_write_json(candidate, updated)
    with pytest.raises(host.DeltaError, match="unused"):
        host._brazil_observations(candidate, baseline, evidence, [])
    empty, _ = producer.source_evidence(config, updated, updated)
    assert not host._brazil_observations(candidate, baseline, empty, [])["business_changed"]


def test_page_area_titles_in_both_progress_tabs_leave_growth_national(tmp_path, monkeypatch):
    _, reference = area_source(tmp_path)
    sys.path.insert(0, str(ROOT / "05_apps"))
    page = importlib.import_module("brazil_soy_page")
    monkeypatch.setattr(page, "current_season", lambda: "2026/2027")
    monkeypatch.setenv("PUBLIC_MARKET_DATA_RUNTIME_ROOT", str(tmp_path))
    (tmp_path / soy.STABLE_RELATIVE_PATH).parent.mkdir(parents=True)
    atomic_write_json(tmp_path / soy.STABLE_RELATIVE_PATH, progress_bundle() | {"area_reference": reference})
    app = AppTest.from_string("from brazil_soy_page import render_brazil_soy_page\nrender_brazil_soy_page()", default_timeout=30).run()
    for label in ("播种进度", "收割进度"):
        app.session_state["brazil-metric-tabs"] = label
        app.run()
        assert not app.exception and len(app.get("plotly_chart")) == 13
        titles = [item.value for item in app.markdown if item.value.startswith("#### ")]
        assert "面积覆盖约96%" in titles[0]
        assert "面积占比7.1%" in next(title for title in titles if "马托格罗索 MT" in title)
        assert any("2025/2026" in item.value and "不随历史季切换" in item.value for item in app.caption)
    app.session_state["brazil-metric-tabs"] = "生长进度"
    app.run()
    assert not app.exception and len(app.get("plotly_chart")) == 1
    assert not any(item.value.startswith("#### ") for item in app.markdown)
