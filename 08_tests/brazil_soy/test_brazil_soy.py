from __future__ import annotations

import importlib
import sys
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

from agri_research_agent.pipelines import brazil_soy as soy
from agri_research_agent.pipelines import brazil_soy_update as update
from agri_research_agent.shared.atomic_storage import atomic_write_json
from agri_research_agent.shared.runtime_context import RuntimeContext, RuntimeMode

ROOT = Path(__file__).resolve().parents[2]
STAMP = "2026-10-07T00:00:00+00:00"


def observation(day="2026-10-02", value=9.4, season="2026/2027", region="BR", metric="PLANTED", **extra):
    return {**soy.observation(season, region, metric, date.fromisoformat(day), value, source_url=soy.SOURCE_URL,
        source_sha256="a" * 64, locator="table 1", retrieved_at=STAMP), **extra}


def bundle(*records):
    return dict(schema_version=soy.SCHEMA, generated_at=STAMP, records=list(records), import_notes=[])


def context(path):
    atomic_write_json(path / ".market-data-runtime.json", {"schema_version": 1, "runtime_id": "brazil-test",
        "classification": "fixture", "module_id": "brazil-soy", "created_at": STAMP})
    return RuntimeContext(RuntimeMode.FIXTURE, "brazil-soy", path)


@pytest.mark.parametrize("extra", [dict(value=None), dict(value=True), dict(value=float("nan")), dict(value=101),
    dict(value=-1), dict(season="2026/2028"), dict(date="2027-09-01"), dict(date="2026-08-31"),
    dict(metric="GOOD_EXCELLENT"), dict(region="MT", metric="FLOWERING"), dict(source_url="https://www.gov.br/saude/test"),
    dict(source_sha256="broken"), dict(published_at="2026-09-01"), dict(date_basis="report_cutoff"), dict(retrieved_at="2026-10-07")])
def test_invalid_observation_rejected(extra):
    with pytest.raises((ValueError, TypeError)):
        soy.validate_bundle(bundle(observation(**extra)))


def test_duplicate_records_and_json_keys_rejected(tmp_path):
    with pytest.raises(ValueError, match="duplicate"):
        soy.validate_bundle(bundle(observation(), observation()))
    path = tmp_path / "bad.json"
    path.write_text('{"schema_version":1,"schema_version":2}', encoding="utf-8")
    with pytest.raises(ValueError, match="duplicate"):
        soy.load_bundle(path)


def test_crop_season_cross_year_and_december_harvest_alignment():
    assert soy.current_season(date(2026, 8, 31)) == "2025/2026"
    assert soy.current_season(date(2026, 9, 1)) == "2026/2027"
    assert soy.seasonal_date(date(2025, 12, 27)) < soy.seasonal_date(date(2026, 1, 3))
    frame = soy.observations_frame(bundle(observation("2025-12-27", 0, "2025/2026", metric="HARVESTED"),
        observation("2025-12-28", 1, "2025/2026", metric="HARVESTED"),
        observation("2026-01-03", 3, "2025/2026", metric="HARVESTED")))
    matched = soy.match_history(frame, date(2026, 12, 27), "2026/2027", "2025/2026")
    assert matched["value"] == 0
    assert soy.match_history(frame, date(2027, 1, 12), "2026/2027", "2025/2026") is None


def test_leap_date_history_match_and_partial_mean():
    frame = soy.observations_frame(bundle(observation("2024-02-29", 40, "2023/2024"),
        observation("2023-02-28", 20, "2022/2023")))
    comparison = soy.compare_metric(frame, "BR", "PLANTED", "2023/2024")
    assert comparison["mean"] == 20
    assert len(comparison["samples"]) == 1


def test_official_references_not_fabricated_historical_observations():
    reference = {"last_season": 8.2, "five_season_mean": 8.8, "locator": "C90,F90"}
    frame = soy.observations_frame(bundle(observation(published_at="2026-10-05", date_basis="report_cutoff", reference=reference)))
    comparison = soy.compare_metric(frame, "BR", "PLANTED", "2026/2027")
    assert comparison["mean"] == 8.8 and comparison["last_season"]["value"] == 8.2
    assert not comparison["samples"] and len(frame) == 1
    assert soy.comparison_table(frame, "BR", "2026/2027").iloc[0]["均值依据"] == "官方五季"


def test_missing_current_season_not_replaced_and_current_trace_last():
    frame = soy.observations_frame(bundle(observation("2025-10-02", 8.2, "2025/2026"), observation()))
    assert soy.compare_metric(frame, "BR", "HARVESTED", "2026/2027")["current"] is None
    figure = soy.seasonal_figure(frame, "BR", "PLANTED", "2026/2027", ["2025/2026"])
    assert figure.data[-1].name == "2026/2027（当前季）" and len(figure.data[-1].x) == 1
    assert figure.data[-1].opacity == 1


def make_history(path):
    wb = Workbook()
    plant = wb.active
    plant.title = "巴西种植"
    plant.append(["年度", "日期", *soy.STATE_NAMES, "All States"])
    plant.append(["2024/2025", datetime(2024, 9, 22), *([0] * 4), .003, 0, 0, 0, 0, .01, 0, 0, .002])
    plant.append(["2024/2025", datetime(2024, 9, 29), *([2.4] * 13)])
    harvest = wb.create_sheet("巴西收割")
    harvest.append(["年度", "周数", "Harvest (%)", *soy.STATE_NAMES, "All States"])
    harvest.append(["2025/2026", 1, datetime(2025, 12, 27), *([.01] * 13)])
    growth = wb.create_sheet("巴西大豆作物进度作图")
    growth.append([None])
    growth.append([None, None, None, "Emergencia", "Desenvolvimento Vegetativo", "Floração", "Enchimento de Grãos", "Maturação", "Colheita"])
    growth.append(["年度", "周数", "日期"])
    growth.append(["2025/2026", 1, datetime(2025, 12, 27), None, .8, .2, None, None, "=1-E4-F4"])
    wb.save(path)
    wb.close()


def test_history_units_quarantine_missing_and_formulas(tmp_path):
    path = tmp_path / "source.xlsx"
    make_history(path)
    identity = soy.sha256_file(path)
    data = soy.import_workbook(path)
    assert soy.sha256_file(path) == identity
    frame = soy.observations_frame(data)
    assert set(frame.loc[frame.metric == "HARVESTED", "value"]) == {1}
    assert len(frame.loc[(frame.metric == "PLANTED") & (frame.date == "2024-09-22")]) == 10
    assert set(frame.loc[frame.metric.isin(soy.STAGES), "metric"]) == {"VEGETATIVE", "FLOWERING"}
    assert any("推算公式" in x for x in data["import_notes"])


def test_merge_revision_idempotence_preserves_history():
    base = bundle(observation())
    merged, stats = update.merge_observations(base, [observation(retrieved_at="2026-10-07T01:00:00+00:00")])
    assert merged == base and stats == dict(added=0, revised=0, unchanged=1)
    with pytest.raises(ValueError, match="revision"):
        update.merge_observations(base, [observation(value=10)])
    merged, stats = update.merge_observations(base, [observation(value=10)], allow_revisions=True)
    assert stats["revised"] == 1 and len(merged["records"]) == 1


def test_candidate_identity_baseline_drift_and_backup(tmp_path):
    ctx = context(tmp_path)
    candidate = update.save_candidate(ctx, bundle(observation()), None, {"initial_import": 1})
    stable = update.activate_local(ctx, candidate)
    initial = soy.sha256_file(stable)
    assert update.activate_local(ctx, candidate) == stable
    second = update.save_candidate(ctx, bundle(observation(), observation("2026-09-25", 3.9)), initial, {"added": 1})
    atomic_write_json(stable, bundle(observation(value=4)))
    with pytest.raises(ValueError, match="moved"):
        update.activate_local(ctx, second)
    atomic_write_json(stable, soy.load_bundle(candidate))
    update.activate_local(ctx, second)
    backups = list((tmp_path / "backups/brazil_soy").glob("*.json"))
    assert len(backups) == 1 and soy.sha256_file(backups[0]) == initial


def test_weekly_overlap_preserves_original_source_and_official_references():
    reference = {"last_season": 2, "five_season_mean": 3, "locator": "C90,F90"}
    previous = observation("2026-09-25", 3.9, reference=reference,
                           published_at="2026-09-28", date_basis="report_cutoff")
    baseline = bundle(previous)
    repeated = {**previous, "source_sha256": "b" * 64, "reference": None,
                "published_at": "2026-10-05", "source_locator": "D90"}
    latest = observation(published_at="2026-10-05", date_basis="report_cutoff", source_sha256="b" * 64)
    merged, stats = update.merge_observations(baseline, [repeated, latest])
    assert stats == dict(added=1, revised=0, unchanged=1)
    assert merged["records"][0] == previous
    changed_reference = {**previous, "reference": {**reference, "five_season_mean": 4}}
    with pytest.raises(ValueError, match="revision"):
        update.merge_observations(baseline, [changed_reference])


def test_prepare_update_requires_matching_archived_bytes(tmp_path):
    ctx = context(tmp_path)
    baseline, updates = tmp_path / "base.json", tmp_path / "updates.json"
    atomic_write_json(baseline, bundle(observation()))
    atomic_write_json(updates, bundle(observation("2026-10-03", 10, date_basis="report_cutoff", published_at="2026-10-05")))
    with pytest.raises(ValueError, match="archived"):
        update.prepare_update(ctx, baseline, updates)


def test_page_two_progress_charts_and_growth_remains_national(tmp_path, monkeypatch):
    sys.path.insert(0, str(ROOT / "05_apps"))
    monkeypatch.setenv("PUBLIC_MARKET_DATA_RUNTIME_ROOT", str(tmp_path))
    data = bundle(observation(), observation(region="MT", value=14.36),
        observation(value=64.1, metric="EMERGENCE"), observation(value=35.9, metric="VEGETATIVE"),
        observation("2025-10-02", 8.2, "2025/2026"))
    (tmp_path / soy.STABLE_RELATIVE_PATH).parent.mkdir(parents=True)
    atomic_write_json(tmp_path / soy.STABLE_RELATIVE_PATH, data)
    app = AppTest.from_string("from brazil_soy_page import render_brazil_soy_page\nrender_brazil_soy_page()", default_timeout=30).run()
    assert not app.exception
    assert len(app.get("plotly_chart")) == 3
    assert app.selectbox[1].value == "2026/2027"
    national = app.dataframe[1].value.copy()
    app.selectbox[0].set_value("MT").run()
    assert not app.exception
    assert app.dataframe[1].value.equals(national)
    assert app.dataframe[0].value.iloc[0]["最新"] == "14.4"


def make_official(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Progresso de safra"
    ws.cell(1, 2, "Soja - Safra 2026/27")
    ws.cell(3, 2, "Semeadura")
    ws.cell(4, 2, "Estado")
    ws.cell(4, 6, "Média 5 anos")
    for c, d in ((3, datetime(2025, 10, 3)), (4, datetime(2026, 9, 25)), (5, datetime(2026, 10, 2))):
        ws.cell(6, c, d)
    for row, name in enumerate((*soy.STATE_NAMES, "12 estados"), 7):
        ws.cell(row, 2, name)
        for column, value in ((3, .082), (4, .039), (5, .094), (6, .088)):
            ws.cell(row, column, value)
    wb.save(path)
    wb.close()


def test_official_parser_uses_cell_cutoff_current_season_and_references(tmp_path):
    report = tmp_path / "report.bin"
    make_official(report)
    info = dict(source_url=soy.SOURCE_URL, sha256=soy.sha256_file(report), retrieved_at=STAMP, published_at="2026-10-05")
    atomic_write_json(tmp_path / "source.json", info)
    data = update.parse_progress(tmp_path / "source.json")
    assert len(data["records"]) == 26
    national = [x for x in data["records"] if x["region"] == "BR"]
    assert [(x["date"], x["value"]) for x in national] == [("2026-09-25", 3.9), ("2026-10-02", 9.4)]
    assert national[-1]["reference"]["five_season_mean"] == pytest.approx(8.8)
    report.write_bytes(b"tampered")
    with pytest.raises(ValueError, match="identity"):
        update.parse_progress(tmp_path / "source.json")
