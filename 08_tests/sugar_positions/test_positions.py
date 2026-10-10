from datetime import date
from copy import deepcopy
from io import BytesIO, StringIO
import csv
import importlib.util
import json
from pathlib import Path
import sys

import openpyxl
import pytest
import requests

from agri_research_agent.sugar_positions.model import (
    RANKED_NET_METHOD, domestic_metrics, foreign_metrics, integer, load_members, split_member, positioning_signal,
)
from agri_research_agent.sugar_positions.sources import parse_cftc, parse_czce, parse_ice
from agri_research_agent.sugar_positions.storage import preview_root, publish, read_snapshot

ROOT = Path(__file__).resolve().parents[2]
MEMBERS = load_members(ROOT / "02_configs/sugar_positions.json")
SOURCE = "https://official.example/report"
STAMP = "2026-10-08T01:00:00Z"


@pytest.mark.parametrize("net,change,expected", [
    (10, 2, "偏多 · 增强"), (10, -2, "偏多 · 减弱"),
    (-10, -2, "偏空 · 增强"), (-10, 2, "偏空 · 减弱"),
    (None, 2, "披露不足，情绪无法判断"), (-10, None, "偏空 · 变化不可比"),
    (0, 2, "多空平衡 · 向多变化"), (10, 0, "偏多 · 净持仓不变"),
])
def test_positioning_signal_does_not_confuse_level_change_or_missing(net, change, expected):
    assert positioning_signal(net, change) == expected


def cot_row():
    row = {"cftc_contract_market_code": "080732", "market_and_exchange_names": "SUGAR NO. 11 - ICE FUTURES U.S.",
           "report_date_as_yyyy_mm_dd": "2026-09-29T00:00:00", "open_interest_all": "100"}
    row.update(tot_rept_positions_long_all="80", tot_rept_positions_short="80")
    # Field names verified against the public Socrata response, not guessed from CSV headers.
    for field in ("prod_merc_positions_long", "prod_merc_positions_short", "swap_positions_long_all",
                  "swap__positions_short_all", "m_money_positions_long_all", "m_money_positions_short_all",
                  "other_rept_positions_long", "other_rept_positions_short",
                  "nonrept_positions_long_all", "nonrept_positions_short_all"):
        row[field] = "20"
    for field in ("swap__positions_spread_all", "m_money_positions_spread", "other_rept_positions_spread"):
        row[field] = "0"
    return row


def test_czce_corrupted_zip_is_a_source_validation_failure():
    with pytest.raises(ValueError, match="文件损坏"):
        parse_czce(b"PK broken download", "2026-09-30", SOURCE, STAMP)


def workbook(day="2026-09-30", total_delta=0, missing_header=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([f"品种：白糖SR     日期：{day}"])
    ws.append(["名次", "会员简称", "交易量（手）", "增减量", "会员简称",
        "未知列" if missing_header else "持买仓量", "增减量", "会员简称", "持卖仓量", "增减量"])
    ws.append([1, "东证期货（代客）", "1,000", "1", "东证期货（代客）", "1,000", "-10",
        "永安期货（代客）", "900", "-5"])
    ws.append([2, "高盛期货（代客）", "500", "0", "永安期货（代客）", "400", "1",
        "东证期货（代客）", "200", "-2"])
    ws.append(["合计", None, "1,500", None, None, 1400 + total_delta, None, None, 1100, None])
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def domestic_row(member, side, value, *, day="2026-09-30", rank=1, account="代客", scope="SR"):
    return dict(report_date=day, scope=scope, member=member, side=side, positions=value,
                rank=rank, account=account, raw_member=f"{member}（{account}）")


@pytest.mark.parametrize("value", [None, "", "NaN", "Infinity", -1, "1.5", True])
def test_invalid_positions_are_rejected(value):
    with pytest.raises(ValueError):
        integer(value)


def test_zero_and_missing_have_distinct_semantics():
    assert integer("0") == 0
    assert integer(None, optional=True) is None
    assert integer("-1,200", signed=True) == -1200
    assert split_member(" 高盛期货（代客） ") == ("高盛期货", "代客")


def test_cftc_contract_identity_totals_duplicates():
    row = cot_row()
    parsed = parse_cftc([row], "futures_only", SOURCE, STAMP)
    assert len(parsed) == 5 and all(r["scope"] == "all_expiries" for r in parsed)
    with pytest.raises(ValueError, match="重复"):
        parse_cftc([row, row], "futures_only", SOURCE, STAMP)
    row["cftc_contract_market_code"] = "080734"
    with pytest.raises(ValueError, match="非糖11"):
        parse_cftc([row], "futures_only", SOURCE, STAMP)
    row["cftc_contract_market_code"] = "080732"
    row["m_money_positions_long_all"] = "21"
    with pytest.raises(ValueError, match="总持仓"):
        parse_cftc([row], "futures_only", SOURCE, STAMP)
    assert len(parse_cftc([row], "combined", SOURCE, STAMP)) == 5


def test_combined_category_rounding_still_checks_official_aggregate():
    row = cot_row()
    row["m_money_positions_short_all"] = "18"
    row["tot_rept_positions_short"] = "79"
    assert len(parse_cftc([row], "combined", SOURCE, STAMP)) == 5
    row["tot_rept_positions_short"] = "77"
    with pytest.raises(ValueError, match="官方报告者合计"):
        parse_cftc([row], "combined", SOURCE, STAMP)


def test_ice_report_types_and_year_are_separate():
    base = {"Open_Interest_All": "100"}
    for prefix in ("Prod_Merc", "Swap", "M_Money", "Other_Rept", "NonRept"):
        base[f"{prefix}_Positions_Long_All"] = "20"
        base[f"{prefix}_Positions_Short_All"] = "20"
        if prefix in ("Swap", "M_Money", "Other_Rept"):
            base[f"{prefix}_Positions_Spread_All"] = "0"
    base.update(Market_and_Exchange_Names="ICE White Sugar Futures - ICE Futures Europe",
        CFTC_Commodity_Code="W", **{"As_of_Date_Form_MM/DD/YYYY": "09/29/2026"})
    stream = StringIO()
    writer = csv.DictWriter(stream, fieldnames=[*base, "FutOnly_or_Combined"])
    writer.writeheader()
    writer.writerow(dict(base, FutOnly_or_Combined="FutOnly"))
    writer.writerow(dict(base, Market_and_Exchange_Names="ICE White Sugar Futures and Options- ICE Futures Europe",
                         FutOnly_or_Combined="Combined"))
    raw = stream.getvalue().encode("utf-8")
    result = parse_ice(raw, 2026, SOURCE, STAMP)
    assert len(result) == 10 and {r["report_type"] for r in result} == {"futures_only", "combined"}
    with pytest.raises(ValueError, match="年份"):
        parse_ice(raw, 2025, SOURCE, STAMP)


def test_czce_reads_member_names_not_volume_ranking():
    rows = parse_czce(workbook(), "2026-09-30", SOURCE, STAMP)
    assert len(rows) == 4
    assert all(r["member"] != "高盛期货" for r in rows)
    east_long = next(r for r in rows if r["member"] == "东证期货" and r["side"] == "long")
    assert east_long["positions"] == 1000 and east_long["reported_change"] == -10


def test_czce_dash_placeholders_are_undisclosed_not_zero_members():
    wb = openpyxl.load_workbook(BytesIO(workbook()))
    wb.active.cell(4, 8, "-")
    wb.active.cell(4, 9, "-")
    wb.active.cell(4, 10, "-")
    wb.active.cell(5, 9, 900)
    raw = BytesIO()
    wb.save(raw)
    wb.close()
    rows = parse_czce(raw.getvalue(), "2026-09-30", SOURCE, STAMP)
    assert len(rows) == 3 and not any(r["member"] == "-" for r in rows)


@pytest.mark.parametrize("raw,day,match", [
    (b"<html>blocked</html>", "2026-09-30", "不是 XLSX"),
    (workbook(), "2026-09-29", "日期"),
    (workbook(total_delta=1), "2026-09-30", "合计"),
    (workbook(missing_header=True), "2026-09-30", "列发生变化"),
])
def test_czce_invalid_sources_fail_closed(raw, day, match):
    with pytest.raises(ValueError, match=match):
        parse_czce(raw, day, SOURCE, STAMP)


def test_fixed5_ranked_net_counts_disclosed_sides_without_rewriting_missing_or_zero():
    rows = [domestic_row("高盛期货", "short", 120),
            domestic_row("东证期货", "long", 0, rank=2),
            domestic_row("东证期货", "short", 30, rank=2)]
    original = deepcopy(rows)
    result = {r["group"]: r for r in domestic_metrics(rows, MEMBERS)}
    assert rows == original
    assert result["goldman"]["long"] is None and result["goldman"]["net"] == -120
    assert result["orient"]["long"] == 0
    assert result["orient"]["net"] == -30
    assert result["fixed5"]["net"] == -150
    assert result["fixed5"]["long"] == 0 and result["fixed5"]["short"] == 150
    assert result["yongan"]["long"] is None and result["yongan"]["short"] is None
    assert result["yongan"]["net"] == 0 and "缺少多仓、空仓" in result["yongan"]["coverage"]
    assert all(r["net_method"] == RANKED_NET_METHOD for r in result.values())
    assert result["fixed5"]["coverage"] == "1/5 家两侧已披露"


def test_accounts_and_contract_scopes_do_not_mix():
    rows = [domestic_row("东证期货", "long", 10), domestic_row("东证期货", "short", 3),
        domestic_row("东证期货", "long", 100, account="自营"),
        domestic_row("东证期货", "short", 1, account="自营"),
        domestic_row("东证期货", "long", 999, scope="SR701")]
    result = domestic_metrics(rows, MEMBERS)
    sr = next(r for r in result if r["scope"] == "SR" and r["group"] == "orient")
    assert sr["net"] == 7
    assert next(r for r in result if r["scope"] == "SR701" and r["group"] == "orient")["net"] == 999


@pytest.mark.parametrize("scope,member,long,short,expected", [
    ("P2701", "高盛期货", None, 45991, -45991),
    ("OI", "永安期货", None, 7149, -7149),
    ("OI701", "永安期货", 3269, 5440, -2171),
])
def test_ranked_net_matches_reported_palm_and_rapeseed_examples(scope, member, long, short, expected):
    rows = [domestic_row(member, side, value, scope=scope) for side, value in
            (("long", long), ("short", short)) if value is not None]
    metrics = domestic_metrics(rows, MEMBERS)
    fixed = next(r for r in metrics if r.get("long_raw_name") or r.get("short_raw_name"))
    assert (fixed["long"], fixed["short"], fixed["net"]) == (long, short, expected)
    assert next(r for r in metrics if r["group"] == "fixed5")["net"] == expected


def test_ranked_changes_include_entry_exit_but_do_not_create_reports_or_bridge_large_gaps():
    assert domestic_metrics([], MEMBERS) == []
    rows = [domestic_row("永安期货", "short", 40, day="2026-09-01"),
        domestic_row("永安期货", "long", 30, day="2026-09-02"),
        domestic_row("其他", "long", 10, day="2026-09-03"),
        domestic_row("永安期货", "short", 20, day="2026-09-30")]
    fixed = [r for r in domestic_metrics(rows, MEMBERS) if r["group"] == "yongan"]
    assert [r["net"] for r in fixed] == [-40, 30, 0, -20]
    assert [r["net_change"] for r in fixed] == [None, 70, -30, None]
    assert fixed[2]["long"] is None and fixed[2]["short"] is None
    assert fixed[2]["previous_date"] == "2026-09-02"
    assert fixed[3]["comparison_days"] == 27


def test_alias_ambiguity_is_rejected():
    rows = [domestic_row("高盛期货", "long", 1), domestic_row("乾坤期货", "long", 2)]
    with pytest.raises(ValueError, match="名称映射"):
        domestic_metrics(rows, MEMBERS)


def test_observed_change_preserves_dates_and_does_not_bridge_null_or_large_gap():
    rows = [dict(market="sugar11", report_type="futures_only", group="managed_money",
            report_date=day, long=long, short=short) for day, long, short in [
                ("2026-09-01", 100, 120), ("2026-09-08", 110, 100),
                ("2026-09-29", 140, 100), ("2026-09-30", None, 100), ("2026-10-01", 150, 100)]]
    result = foreign_metrics(rows)
    assert result[0]["net"] == -20 and result[1]["net_change"] == 30
    assert result[1]["previous_date"] == "2026-09-01"
    assert all(r["net_change"] is None for r in result[2:])


def test_rank_change_uses_two_baskets_not_current_member_change():
    rows = [domestic_row("甲", "long", 100, day="2026-09-29"),
            domestic_row("乙", "short", 80, day="2026-09-29"),
            domestic_row("丙", "long", 130), domestic_row("丁", "short", 70)]
    result = [r for r in domestic_metrics(rows, MEMBERS) if r["group"] == "top20"]
    assert result[-1]["net_change"] == 40 and result[-1]["previous_date"] == "2026-09-29"


def test_storage_merge_replaces_whole_revised_ranking_and_preserves_other_sources(tmp_path):
    foreign = parse_cftc([cot_row()], "futures_only", SOURCE, STAMP)
    domestic = parse_czce(workbook(), "2026-09-30", SOURCE, STAMP)
    publish(tmp_path, foreign, domestic, [("cftc", b"[]", SOURCE, "json")], [])
    revised = [domestic_row("新会员", "long", 12), domestic_row("新会员", "short", 3)]
    result = publish(tmp_path, [], revised, [], [])
    assert result["foreign"] == read_snapshot(tmp_path)["foreign"]
    assert len(result["foreign"]) == 5 and len(result["domestic"]) == 2
    assert len(list((tmp_path / "releases").iterdir())) == 2


def test_storage_tamper_and_concurrent_writer_are_rejected(tmp_path):
    result = publish(tmp_path, [], [], [], [])
    path = tmp_path / "releases" / result["release_id"] / "snapshot.json"
    path.write_text("{}", encoding="utf-8")
    with pytest.raises(ValueError, match="数据校验"):
        read_snapshot(tmp_path)
    (tmp_path / "update.lock").write_text("other writer", encoding="utf-8")
    with pytest.raises(FileExistsError):
        publish(tmp_path, [], [], [], [])


def test_preview_writer_rejects_main_checkout_and_redirection(tmp_path):
    with pytest.raises(ValueError, match="linked worktree"):
        preview_root(tmp_path)
    (tmp_path / ".git").write_text("gitdir: test", encoding="utf-8")
    assert preview_root(tmp_path) == tmp_path / "01_data" / "sugar_positions_preview"


def test_partial_source_failure_retains_success_and_is_visible():
    spec = importlib.util.spec_from_file_location("sugar_update", ROOT / "04_scripts/sugar_positions/update_sugar_positions.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    class FakeSource:
        def cftc(self, kind, year):
            if kind == "combined":
                raise requests.Timeout("timeout")
            return parse_cftc([cot_row()], kind, SOURCE, STAMP), b"[]", SOURCE
        def ice(self, year):
            raise ValueError("invalid CSV")
    foreign, domestic, captures, attempts = module.collect(FakeSource(), markets="foreign", start_year=2026,
        start_day=date(2026, 9, 29), end_day=date(2026, 9, 30))
    assert len(foreign) == 5 and len(captures) == 1 and domestic == []
    assert [a["status"] for a in attempts] == ["ok", "failed", "failed"]
