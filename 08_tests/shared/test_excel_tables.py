from datetime import date, datetime

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import pytest

from agri_research_agent.shared.excel_tables import rewritten_tables, write_tables


def test_streamed_generated_tables_keep_existing_excel_values_dates_blanks_and_zero(tmp_path):
    frame = pd.DataFrame({
        "日期": [datetime(2026, 10, 28, 16, 30), date(2026, 10, 29), None],
        "价格": [0.0, np.nan, 5278.25],
        "说明": ["柳州", "", "保留缺口"],
        "开关": [True, False, None],
    })
    old = tmp_path / "pandas.xlsx"
    new = tmp_path / "stream.xlsx"
    frame.to_excel(old, sheet_name="报价", index=False)
    write_tables(new, [("报价", frame)])
    with_values = lambda p: list(load_workbook(p, read_only=True, data_only=False)["报价"].values)
    assert with_values(new) == with_values(old)


def test_rewrite_preserves_sheet_order_and_evidence_without_reading_replaced_history(tmp_path, monkeypatch):
    old = tmp_path / "source.xlsx"
    target = tmp_path / "candidate.xlsx"
    initial = pd.DataFrame({"value": [1, 2]})
    evidence = pd.DataFrame({"source": ["original"], "price": [0]})
    with pd.ExcelWriter(old) as writer:
        evidence.to_excel(writer, sheet_name="证据", index=False)
        initial.to_excel(writer, sheet_name="spread_long", index=False)
    raw = old.read_bytes()
    parse = pd.ExcelFile.parse
    read = []
    def checked(self, sheet_name=0, *args, **kwargs):
        read.append(sheet_name)
        if sheet_name == "spread_long":
            raise AssertionError("Superseded history must not be materialized")
        return parse(self, sheet_name, *args, **kwargs)
    monkeypatch.setattr(pd.ExcelFile, "parse", checked)
    replacement = pd.DataFrame({"value": [3, 0, None]})
    write_tables(target, rewritten_tables(old, {"spread_long": replacement}))
    assert read == ["证据"] and old.read_bytes() == raw
    book = load_workbook(target, read_only=True)
    assert book.sheetnames == ["证据", "spread_long"]
    assert list(book["证据"].values) == [("source", "price"), ("original", 0)]
    assert list(book["spread_long"].values) == [("value",), (3,), (0,), (None,)]


def test_rewrite_missing_history_fails_without_modifying_input(tmp_path):
    source = tmp_path / "source.xlsx"
    pd.DataFrame({"value": [1]}).to_excel(source, index=False)
    raw = source.read_bytes()
    with pytest.raises(ValueError, match="missing"):
        list(rewritten_tables(source, {"spread_long": pd.DataFrame()}))
    assert source.read_bytes() == raw


def test_empty_evidence_table_has_no_phantom_row(tmp_path):
    old, new = tmp_path / 'old.xlsx', tmp_path / 'new.xlsx'
    pd.DataFrame().to_excel(old, sheet_name='failures', index=False)
    write_tables(new, [('failures', pd.DataFrame())])
    for path in [old, new]:
        book = load_workbook(path, read_only=True)
        assert list(book['failures'].values) == []
        book.close()
