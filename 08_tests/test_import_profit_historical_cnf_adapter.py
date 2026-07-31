from __future__ import annotations

from datetime import date
import hashlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
import pytest

from agri_research_agent.import_profit.historical_cnf_adapter import (
    HISTORICAL_CNF_SCHEMA,
    HistoricalCnfError,
    adapt_historical_cnf_excel,
    shipment_year_for,
)


GROUPS = (("美湾", 2), ("美西", 14), ("巴西", 26), ("阿根廷", 38))


def workbook(
    path: Path,
    *,
    dates: tuple[date, ...] = (date(2026, 7, 28), date(2026, 12, 31)),
    bad_header: tuple[int, str] | None = None,
) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "CNF报价"
    sheet.cell(1, 1, "日期")
    for label, start in GROUPS:
        for month in range(1, 13):
            sheet.cell(1, start + month - 1, f"{label}{month}月")
    if bad_header is not None:
        sheet.cell(1, bad_header[0], bad_header[1])
    for row, business_date in enumerate(dates, start=2):
        sheet.cell(row, 1, business_date)
    if dates:
        sheet["B2"] = 12.5
        sheet["C2"] = 0
        sheet["D2"] = -1.5
        sheet["E2"] = "#N/A"
        sheet["E2"].data_type = "e"
        sheet["F2"] = "=NA()"
        sheet["G2"] = "#N/A"
    book.save(path)
    return path


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_four_origin_headers_values_nulls_and_source_are_normalized(
    tmp_path: Path,
) -> None:
    path = workbook(tmp_path / "cnf.xlsx")
    before = sha(path)
    result = adapt_historical_cnf_excel(path)

    assert sha(path) == before
    assert len(result.records) == 2 * 4 * 12
    assert len({record.key for record in result.records}) == len(result.records)
    assert list(HISTORICAL_CNF_SCHEMA.names) == [
        "business_date",
        "commodity",
        "origin",
        "shipment_year",
        "shipment_month",
        "cnf_cents_per_bushel",
        "source",
        "updated_at",
        "batch_id",
    ]
    values = {
        (record.origin, record.shipment_month): record.cnf_cents_per_bushel
        for record in result.records
        if record.business_date == date(2026, 7, 28)
    }
    assert values[("us_gulf", 1)] == 12.5
    assert values[("us_gulf", 2)] == 0
    assert values[("us_gulf", 3)] == -1.5
    assert values[("us_gulf", 4)] is None
    assert values[("us_gulf", 5)] is None
    assert values[("us_gulf", 6)] is None
    assert set(record.origin for record in result.records) == {
        "brazil",
        "us_gulf",
        "us_pnw",
        "argentina",
    }
    assert all(record.source == "historical_excel" for record in result.records)
    assert result.quality_report["value_counts"] == {
        "numeric": 3,
        "null": 93,
        "excel_error": 2,
        "true_blank": 91,
        "zero": 1,
        "negative": 1,
        "invalid": 0,
    }
    assert result.quality_report["complete_null_business_date_count"] == 1
    assert result.quality_report["candidate_status"] == "passed_with_warnings"
    assert [record.key for record in result.records] == sorted(
        record.key for record in result.records
    )


@pytest.mark.parametrize(
    ("business_date", "month", "expected"),
    [
        (date(2026, 7, 28), 8, 2026),
        (date(2026, 7, 28), 7, 2027),
        (date(2026, 7, 28), 1, 2027),
        (date(2026, 7, 28), 12, 2026),
        (date(2026, 12, 31), 12, 2027),
        (date(2026, 12, 31), 1, 2027),
        (date(2026, 1, 2), 1, 2027),
        (date(2026, 1, 2), 2, 2026),
    ],
)
def test_rolling_shipment_year_rule(
    business_date: date, month: int, expected: int
) -> None:
    assert shipment_year_for(business_date, month) == expected


def test_business_key_samples_use_year_rule_even_when_quote_is_null(
    tmp_path: Path,
) -> None:
    result = adapt_historical_cnf_excel(workbook(tmp_path / "cnf.xlsx"))
    july = next(
        item
        for item in result.records
        if item.business_date == date(2026, 7, 28)
        and item.origin == "brazil"
        and item.shipment_month == 7
    )
    august = next(
        item
        for item in result.records
        if item.business_date == date(2026, 7, 28)
        and item.origin == "brazil"
        and item.shipment_month == 8
    )
    assert july.shipment_year == 2027
    assert august.shipment_year == 2026
    assert july.cnf_cents_per_bushel is None
    assert august.cnf_cents_per_bushel is None


def test_duplicate_business_date_is_fatal(tmp_path: Path) -> None:
    path = workbook(
        tmp_path / "duplicate.xlsx",
        dates=(date(2026, 7, 28), date(2026, 7, 28)),
    )
    with pytest.raises(HistoricalCnfError) as error:
        adapt_historical_cnf_excel(path)
    assert error.value.quality_report["duplicate_business_date_count"] == 1
    assert error.value.quality_report["candidate_status"] == "failed"


@pytest.mark.parametrize(
    "bad_header",
    [
        (2, "墨西哥1月"),
        (13, "美湾13月"),
    ],
)
def test_wrong_origin_or_month_header_is_fatal(
    tmp_path: Path, bad_header: tuple[int, str]
) -> None:
    path = workbook(tmp_path / "bad.xlsx", bad_header=bad_header)
    with pytest.raises(HistoricalCnfError) as error:
        adapt_historical_cnf_excel(path)
    assert error.value.quality_report["header_error_count"] == 1


def test_missing_sheet_is_fatal(tmp_path: Path) -> None:
    path = tmp_path / "wrong-sheet.xlsx"
    book = Workbook()
    book.active.title = "其他"
    book.save(path)
    with pytest.raises(HistoricalCnfError, match="required worksheet"):
        adapt_historical_cnf_excel(path)


@pytest.mark.parametrize("month", [0, 13])
def test_invalid_shipment_month_is_rejected(month: int) -> None:
    with pytest.raises(HistoricalCnfError):
        shipment_year_for(date(2026, 7, 28), month)
