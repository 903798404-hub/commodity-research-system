from __future__ import annotations

import json
import logging
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from bs4 import BeautifulSoup


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "03_src"))

from oil_world_data.workbook_builder import (  # noqa: E402
    WorkbookBuildError,
    _normalize_table,
    build_workbook,
    normalize_cell,
    safe_sheet_name,
)


class WorkbookBuilderTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="oil_world_workbook_builder_")
        self.root = Path(self.temporary.name)
        self.source = self.root / "source"
        self.output = self.root / "output"
        self.source.mkdir()
        (self.source / "stats").mkdir()
        self.config = self.root / "config.json"
        self.config.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "builder_version": "test",
                    "start_pattern": "__* START.htm",
                    "stats_directory": "stats",
                    "excel_engine": "openpyxl",
                    "metadata_sheet": "__cell_metadata",
                    "footnotes_sheet": "__footnotes",
                    "required_report_ids": ["AN_TEST"],
                }
            ),
            encoding="utf-8",
        )
        (self.source / "__March 2026 START.htm").write_text(
            '<table><tr><td><h3>Section</h3><a href="Index.htm">Index</a></td></tr></table>',
            encoding="iso-8859-1",
        )
        (self.source / "Index.htm").write_text(
            '<b>Oilseeds</b><a href="stats/AN_TEST.HTM">Test report</a>',
            encoding="iso-8859-1",
        )
        (self.source / "stats" / "AN_TEST.HTM").write_text(
            '<b>TEST BALANCE (1000 T)</b><table>'
            '<tr><td>Metric</td><td>25/26F</td><td>24/25</td></tr>'
            '<tr><td>Production</td><td>0*</td><td>-12.5</td></tr>'
            '<tr><td>Stocks/Use</td><td>10.5%</td><td>.</td></tr>'
            '</table><p>* Forecast estimate</p>',
            encoding="iso-8859-1",
        )
        self.logger = logging.getLogger(f"test_builder_{id(self)}")
        self.logger.handlers.clear()
        self.logger.addHandler(logging.NullHandler())

    def tearDown(self):
        self.temporary.cleanup()

    def build(self):
        return build_workbook(
            release="2026-03",
            source_dir=self.source,
            output_dir=self.output,
            config_path=self.config,
            cli_path=Path(__file__),
            logger=self.logger,
        )

    def test_zero_negative_percentage_and_star_are_preserved(self):
        manifest = self.build()
        self.assertEqual(manifest["starred_numeric_count"], 1)
        self.assertEqual(manifest["footnote_record_count"], 1)
        workbook = openpyxl.load_workbook(manifest["output_path"], data_only=True)
        try:
            sheet = workbook["AN_TEST"]
            values = [cell.value for row in sheet.iter_rows() for cell in row]
            self.assertIn(0, values)
            self.assertIn(-12.5, values)
            self.assertIn(10.5, values)
            metadata = workbook["__cell_metadata"]
            rows = list(metadata.iter_rows(values_only=True))
            header = rows[0]
            marker_index = header.index("footnote_marker")
            raw_index = header.index("raw_value")
            self.assertTrue(any(row[raw_index] == "0*" and row[marker_index] == "*" for row in rows[1:]))
            self.assertEqual(metadata.sheet_state, "hidden")
        finally:
            workbook.close()

    def test_existing_candidate_is_not_overwritten(self):
        self.build()
        with self.assertRaisesRegex(WorkbookBuildError, "拒绝覆盖"):
            self.build()

    def test_missing_stats_reference_stops_before_candidate(self):
        (self.source / "stats" / "AN_TEST.HTM").unlink()
        with self.assertRaisesRegex(WorkbookBuildError, "stats引用不完整"):
            self.build()
        self.assertFalse((self.output / "油世界季度表-March 2026.xlsx").exists())

    def test_normalization_and_safe_sheet_names(self):
        self.assertEqual(normalize_cell("0.0*")[:2], (0.0, "*"))
        self.assertEqual(normalize_cell("-")[0], None)
        self.assertEqual(normalize_cell("(12.5)")[0], -12.5)
        used: set[str] = set()
        first = safe_sheet_name("A/B:C*D?E[F]", used)
        second = safe_sheet_name("A/B:C*D?E[F]", used)
        self.assertNotRegex(first, r"[\[\]:*?/\\]")
        self.assertNotEqual(first, second)
        self.assertLessEqual(len(second), 31)

    def test_marker_columns_are_folded_and_nan_columns_are_removed(self):
        table = BeautifulSoup(
            """
            <table>
              <tr><td>Country</td><td>25/26</td><td>p</td><td>24/25</td><td></td></tr>
              <tr><td>Ukraine</td><td>0</td><td>*</td><td>12</td><td></td></tr>
            </table>
            """,
            "html.parser",
        ).table
        normalized = _normalize_table(table)
        self.assertIsNotNone(normalized)
        values, raw, markers = normalized
        self.assertEqual(values.shape[1], 3)
        self.assertEqual(values.iat[0, 1], "25/26")
        self.assertEqual(markers.iat[0, 1], "p")
        self.assertEqual(values.iat[1, 1], 0)
        self.assertEqual(markers.iat[1, 1], "*")
        self.assertEqual(raw.iat[1, 1], "0*")


if __name__ == "__main__":
    unittest.main()
