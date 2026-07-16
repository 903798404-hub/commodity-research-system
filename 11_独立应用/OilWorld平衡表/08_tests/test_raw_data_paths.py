from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "03_src"))

from oil_world_data.paths import (  # noqa: E402
    RAW_DATA_ROOT_ENV,
    raw_source_reference,
    resolve_configured_source_workbook,
    resolve_raw_data_root,
)
from oil_world_data.release_pipeline import find_source_workbook  # noqa: E402


class OilWorldRawDataPathTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory(prefix="oil_world_paths_")
        self.root = Path(self.temporary.name) / "project"
        (self.root / "02_configs").mkdir(parents=True)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_environment_override_has_priority(self) -> None:
        environment_root = Path(self.temporary.name) / "environment_raw"
        local_root = Path(self.temporary.name) / "local_raw"
        (self.root / "02_configs" / "local_paths.json").write_text(
            json.dumps({RAW_DATA_ROOT_ENV: str(local_root)}), encoding="utf-8"
        )
        with patch.dict(os.environ, {RAW_DATA_ROOT_ENV: str(environment_root)}):
            self.assertEqual(resolve_raw_data_root(self.root), environment_root.resolve())

    def test_local_override_is_used_without_environment(self) -> None:
        local_root = Path(self.temporary.name) / "local_raw"
        (self.root / "02_configs" / "local_paths.json").write_text(
            json.dumps({RAW_DATA_ROOT_ENV: str(local_root)}), encoding="utf-8"
        )
        with patch.dict(os.environ, {}, clear=True):
            self.assertEqual(resolve_raw_data_root(self.root), local_root.resolve())

    def test_legacy_directory_remains_a_compatibility_fallback(self) -> None:
        with patch.dict(os.environ, {}, clear=True):
            self.assertEqual(resolve_raw_data_root(self.root), (self.root / "01_原始资料").resolve())

    def test_baseline_config_uses_raw_root_but_sandbox_path_stays_local(self) -> None:
        raw_root = Path(self.temporary.name) / "raw"
        with patch.dict(os.environ, {RAW_DATA_ROOT_ENV: str(raw_root)}):
            baseline = resolve_configured_source_workbook(
                self.root, "01_原始资料/2026-06/source.xlsx"
            )
            sandbox = resolve_configured_source_workbook(self.root, "source/source.xlsx")
        self.assertEqual(baseline, (raw_root / "2026-06" / "source.xlsx").resolve())
        self.assertEqual(sandbox, (self.root / "source" / "source.xlsx").resolve())

    def test_release_lookup_and_reference_use_external_root(self) -> None:
        raw_root = Path(self.temporary.name) / "raw"
        workbook = raw_root / "2026-09" / "Oil World-September 2026.xlsx"
        workbook.parent.mkdir(parents=True)
        openpyxl.Workbook().save(workbook)
        with patch.dict(os.environ, {RAW_DATA_ROOT_ENV: str(raw_root)}):
            resolved = find_source_workbook(self.root, "2026-09")
            reference = raw_source_reference(self.root, resolved)
        self.assertEqual(resolved, workbook)
        self.assertEqual(reference, "raw/2026-09/Oil World-September 2026.xlsx")

    def test_current_june_workbook_is_readable(self) -> None:
        workbook = find_source_workbook(ROOT, "2026-06")
        book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
        try:
            self.assertGreater(len(book.sheetnames), 0)
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
