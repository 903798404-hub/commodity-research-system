from __future__ import annotations

import hashlib
import json
import shutil
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "03_src"))

from oil_world_data.release_pipeline import (  # noqa: E402
    ReleasePipelineError,
    _annual_column_run,
    _atomic_replace,
    _comparison_record,
    _rebase_audit_for_workbook,
    update_release,
)
import openpyxl  # noqa: E402


BASELINE_TREE_HASH = "8ec74563583a2002d934aeda53689c6e19da30e1496ada999a0422faefff89aa"


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def tree_hash(root: Path) -> str:
    digest = hashlib.sha256()
    for path in sorted(item for item in root.rglob("*") if item.is_file()):
        digest.update(path.relative_to(root).as_posix().encode())
        digest.update(b"\0")
        digest.update(hashlib.sha256(path.read_bytes()).digest())
    return digest.hexdigest()


class FixtureBuilder:
    def __init__(self, value_delta: float = 0):
        self.value_delta = value_delta

    def __call__(self, project_root: Path, release: str, workbook: Path, staging: Path):
        destination = staging / "fixture_release"
        baseline = project_root / "01_data" / "releases" / "2026-06"
        shutil.copytree(baseline, destination)
        index = read_json(destination / "index.json")
        index["release"] = release
        write_json(destination / "index.json", index)
        manifest = read_json(destination / "manifest.json")
        manifest["release"] = release
        write_json(destination / "manifest.json", manifest)
        quality = read_json(destination / "quality_report.json")
        quality["release"] = release
        quality["passed"] = True
        write_json(destination / "quality_report.json", quality)
        for item in index["files"]:
            path = destination / item["path"]
            payload = read_json(path)
            payload["release"] = release
            for metric in payload["metrics"]:
                if metric["mapping_status"] in {"direct", "derived"}:
                    metric["values"] = {
                        period: (value + self.value_delta if value is not None else None)
                        for period, value in metric["values"].items()
                    }
            write_json(path, payload)
        return destination, {"fixture": True, "release": release}


class OilWorldReleasePipelineTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="oil_world_pipeline_test_")
        self.root = Path(self.temporary.name)
        for relative in [
            Path("01_data/releases/2026-06"),
            Path("public/data/oil_world/releases/2026-06"),
        ]:
            shutil.copytree(ROOT / relative, self.root / relative)
        public_root = self.root / "public/data/oil_world"
        write_json(public_root / "latest.json", {"release": "2026-06"})
        write_json(
            public_root / "releases.json",
            {
                "releases": [
                    {"release": "2026-06", "label": "June 2026", "available": True}
                ]
            },
        )

    def tearDown(self):
        self.temporary.cleanup()

    def add_source(self, release: str) -> Path:
        year, month = release.split("-")
        names = {
            "03": "March", "06": "June", "09": "September", "12": "December"
        }
        path = self.root / "01_原始资料" / release / f"油世界季度表-{names[month]} {year}.xlsx"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(f"fixture-{release}".encode())
        return path

    def test_backfill_keeps_latest_and_builds_direct_adjacent_comparison(self):
        self.add_source("2026-03")
        report = update_release(
            self.root,
            "2026-03",
            backfill=True,
            snapshot_builder=FixtureBuilder(-10),
            import_time="2026-07-15T00:00:00+00:00",
        )
        self.assertEqual(report["latest_before"], "2026-06")
        self.assertEqual(report["latest_after"], "2026-06")
        self.assertEqual(read_json(self.root / "public/data/oil_world/latest.json")["release"], "2026-06")
        releases = read_json(self.root / "public/data/oil_world/releases.json")["releases"]
        self.assertEqual([item["release"] for item in releases], ["2026-03", "2026-06"])
        self.assertEqual(releases[0]["next_release"], "2026-06")
        self.assertEqual(releases[1]["previous_release"], "2026-03")
        comparison = self.root / "public/data/oil_world/comparisons/2026-03_to_2026-06/index.json"
        self.assertTrue(comparison.exists())
        self.assertGreater(read_json(comparison)["calculated_revision_count"], 0)

    def test_future_update_advances_latest_and_compares_only_to_direct_previous(self):
        self.add_source("2026-03")
        update_release(self.root, "2026-03", backfill=True, snapshot_builder=FixtureBuilder(-10))
        self.add_source("2026-09")
        report = update_release(self.root, "2026-09", snapshot_builder=FixtureBuilder(20))
        self.assertEqual(report["latest_after"], "2026-09")
        self.assertEqual(read_json(self.root / "public/data/oil_world/latest.json")["release"], "2026-09")
        comparisons = self.root / "public/data/oil_world/comparisons"
        self.assertTrue((comparisons / "2026-03_to_2026-06").is_dir())
        self.assertTrue((comparisons / "2026-06_to_2026-09").is_dir())
        self.assertFalse((comparisons / "2026-03_to_2026-09").exists())

    def test_earlier_insert_rebuilds_chronological_adjacency(self):
        self.add_source("2026-03")
        update_release(self.root, "2026-03", backfill=True, snapshot_builder=FixtureBuilder(-10))
        self.add_source("2025-12")
        update_release(self.root, "2025-12", backfill=True, snapshot_builder=FixtureBuilder(-20))
        comparisons = self.root / "public/data/oil_world/comparisons"
        pairs = sorted(path.name for path in comparisons.iterdir() if path.is_dir())
        self.assertEqual(pairs, ["2025-12_to_2026-03", "2026-03_to_2026-06"])

    def test_validate_only_builds_preview_without_formal_writes(self):
        self.add_source("2026-03")
        latest_before = (self.root / "public/data/oil_world/latest.json").read_bytes()
        releases_before = (self.root / "public/data/oil_world/releases.json").read_bytes()
        report = update_release(
            self.root,
            "2026-03",
            backfill=True,
            validate_only=True,
            snapshot_builder=FixtureBuilder(-10),
        )
        self.assertTrue(report["validate_only"])
        self.assertFalse((self.root / "01_data/releases/2026-03").exists())
        self.assertFalse((self.root / "public/data/oil_world/comparisons").exists())
        self.assertEqual((self.root / "public/data/oil_world/latest.json").read_bytes(), latest_before)
        self.assertEqual((self.root / "public/data/oil_world/releases.json").read_bytes(), releases_before)

    def test_duplicate_release_is_rejected_without_pointer_change(self):
        self.add_source("2026-06")
        latest_before = (self.root / "public/data/oil_world/latest.json").read_bytes()
        with self.assertRaises(ReleasePipelineError):
            update_release(self.root, "2026-06", backfill=True, snapshot_builder=FixtureBuilder())
        self.assertEqual((self.root / "public/data/oil_world/latest.json").read_bytes(), latest_before)

    def test_failed_build_does_not_change_pointer(self):
        self.add_source("2026-09")
        latest_before = (self.root / "public/data/oil_world/latest.json").read_bytes()

        def fail_builder(*_args):
            raise ReleasePipelineError("fixture failure")

        with self.assertRaises(ReleasePipelineError):
            update_release(self.root, "2026-09", snapshot_builder=fail_builder)
        self.assertEqual((self.root / "public/data/oil_world/latest.json").read_bytes(), latest_before)
        self.assertFalse((self.root / "01_data/releases/2026-09").exists())

    def test_source_directory_must_contain_exactly_one_matching_workbook(self):
        source = self.add_source("2026-09")
        source.with_name("Oil World-September 2026.xlsx").write_bytes(b"second fixture")
        with self.assertRaisesRegex(ReleasePipelineError, "只能包含一个"):
            update_release(self.root, "2026-09", snapshot_builder=FixtureBuilder())

    def test_quarter_revision_rules_preserve_null_zero_and_ratio_units(self):
        identity = ("Palm System", "Palm Oil", "Global")
        direct_previous = {"mapping_status": "direct", "unit": "1000 T", "values": {"2025/26": 0}, "source_report_id": ["A"]}
        direct_current = {"mapping_status": "direct", "unit": "1000 T", "values": {"2025/26": 5}, "source_report_id": ["B"]}
        record = _comparison_record("2026-03", "2026-06", identity, "Production", "2025/26", direct_previous, direct_current)
        self.assertEqual(record["quarter_revision"], 5)
        self.assertEqual(record["comparison_status"], "calculated")

        direct_current["values"]["2025/26"] = None
        missing_value = _comparison_record("2026-03", "2026-06", identity, "Production", "2025/26", direct_previous, direct_current)
        self.assertIsNone(missing_value["quarter_revision"])

        ratio_previous = {"mapping_status": "direct", "unit": "%", "values": {"2025/26": 10}, "source_report_id": ["A"]}
        ratio_current = {"mapping_status": "direct", "unit": "%", "values": {"2025/26": 11.25}, "source_report_id": ["B"]}
        ratio = _comparison_record("2026-03", "2026-06", identity, "Stocks/Use Ratio", "2025/26", ratio_previous, ratio_current)
        self.assertEqual(ratio["quarter_revision"], 1.25)
        self.assertEqual(ratio["unit"], "percentage points")

        for status in ("missing", "not_applicable", "conflict"):
            item = {"mapping_status": status, "unit": "1000 T", "values": {"2025/26": 1}, "source_report_id": []}
            unavailable = _comparison_record("2026-03", "2026-06", identity, "Imports", "2025/26", item, item)
            self.assertIsNone(unavailable["quarter_revision"])
            self.assertEqual(unavailable["comparison_status"], "mapping_status_ineligible")

        derived_current = dict(direct_current, mapping_status="derived", values={"2025/26": 5})
        changed = _comparison_record(
            "2026-03", "2026-06", identity, "Production", "2025/26", direct_previous, derived_current
        )
        self.assertIsNone(changed["quarter_revision"])
        self.assertEqual(changed["comparison_status"], "mapping_status_changed")

    def test_real_2026_06_snapshot_remains_byte_identical(self):
        self.assertEqual(tree_hash(ROOT / "01_data/releases/2026-06"), BASELINE_TREE_HASH)
        self.assertEqual(tree_hash(ROOT / "public/data/oil_world/releases/2026-06"), BASELINE_TREE_HASH)

    def test_safe_column_shift_and_one_year_history_difference_are_accepted(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "AN_FIXTURE"
        sheet["D2"] = "Sept Aug 25/26F"
        sheet["E2"] = "Sept Aug 24/25"
        sheet["F2"] = "Sept Aug 23/24"
        sheet["G2"] = "Sept Aug 22/23"
        sheet["H2"] = "Sept Aug 21/22"
        sheet["I2"] = "Apr Sept 2026F"
        columns, labels = _annual_column_run(
            sheet,
            header_row=2,
            old_start=2,
            old_end=6,
            expected_labels=[
                "Sept Aug 26/27F",
                "Sept Aug 25/26",
                "Sept Aug 24/25",
                "Sept Aug 23/24",
                "Sept Aug 22/23",
            ],
        )
        self.assertEqual(columns, [4, 5, 6, 7, 8])
        self.assertEqual(labels[-1], "Sept Aug 21/22")

    def test_three_repeated_annual_groups_are_resolved_inside_the_metric_block(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "AN13993"
        blocks = ((3, "PRODUCTION"), (7, "Y I E L D"), (11, "HARVEST'D AREA"))
        for start, block in blocks:
            sheet.cell(2, start, block)
            for offset, label in enumerate(("25/26", "24/25", "23/24")):
                sheet.cell(3, start + offset, label)
            sheet.cell(3, start + 3, "20/21- 24/25")

        expectations = {
            "Production": ([3, 4, 5], "PRODUCTION"),
            "Yield": ([7, 8, 9], "Y I E L D"),
            "Area Harvested": ([11, 12, 13], "HARVEST'D AREA"),
        }
        for metric, (expected_columns, block) in expectations.items():
            columns, labels = _annual_column_run(
                sheet,
                header_row=3,
                old_start=expected_columns[0],
                old_end=expected_columns[-1],
                expected_labels=["26/27", "25/26", "24/25"],
                report_id="AN13993",
                metric=metric,
                region="Global",
                expected_block=block,
            )
            self.assertEqual(columns, expected_columns)
            self.assertEqual(labels, ["25/26", "24/25", "23/24"])

    def test_metric_block_can_move_and_have_one_fewer_year(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "AN13993"
        sheet["D2"] = "PRODUCTION"
        sheet["D3"] = "25/26"
        sheet["E3"] = "24/25"
        columns, labels = _annual_column_run(
            sheet,
            header_row=3,
            old_start=3,
            old_end=5,
            expected_labels=["26/27", "25/26", "24/25"],
            report_id="AN13993",
            metric="Production",
            region="Global",
            unit="1000 T",
            expected_block="PRODUCTION",
        )
        self.assertEqual(columns, [4, 5])
        self.assertEqual(labels, ["25/26", "24/25"])

    def test_ambiguous_annual_groups_fail_with_candidate_coordinates(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "AN13993"
        for cell, label in (("B3", "25/26"), ("C3", "24/25"), ("F3", "25/26"), ("G3", "24/25")):
            sheet[cell] = label
        with self.assertRaisesRegex(ReleasePipelineError, r"B3:C3.*F3:G3"):
            _annual_column_run(
                sheet,
                header_row=3,
                old_start=4,
                old_end=5,
                expected_labels=["26/27", "25/26", "24/25"],
                report_id="AN13993",
                metric="Production",
                region="Global",
                unit="1000 T",
            )

    def test_real_june_rebase_preserves_all_649_mappings_and_other_reports(self):
        audit_path = ROOT / "07_docs" / "报表映射审计" / "2026-06_report_mapping_audit.json"
        workbook_path = ROOT / "01_原始资料" / "2026-06" / "油世界季度表-June 2026.xlsx"
        audit = read_json(audit_path)
        original = deepcopy(audit["coverage_matrix"])
        rebased = _rebase_audit_for_workbook(deepcopy(audit), workbook_path)["coverage_matrix"]
        self.assertEqual(len(rebased), 649)
        self.assertEqual(
            [row["mapping_status"] for row in rebased],
            [row["mapping_status"] for row in original],
        )
        self.assertEqual(
            [row["source_cell_or_range"] for row in rebased],
            [row["source_cell_or_range"] for row in original],
        )
        other_reports = {
            row["report_id"]
            for row in audit["report_catalog"]
            if row["decision"] == "adopted" and row["report_id"] != "AN13993"
        }
        self.assertEqual(len(other_reports), 29)

    def test_atomic_directory_install_retries_transient_permission_error(self):
        with (
            patch(
                "oil_world_data.release_pipeline.os.replace",
                side_effect=[PermissionError("transient lock"), None],
            ) as replace,
            patch("oil_world_data.release_pipeline.time.sleep") as sleep,
        ):
            _atomic_replace(Path("source"), Path("destination"), attempts=3)
        self.assertEqual(replace.call_count, 2)
        sleep.assert_called_once_with(0.1)

    def test_atomic_directory_install_does_not_mask_persistent_failure(self):
        with (
            patch(
                "oil_world_data.release_pipeline.os.replace",
                side_effect=PermissionError("persistent lock"),
            ) as replace,
            patch("oil_world_data.release_pipeline.time.sleep") as sleep,
            self.assertRaises(PermissionError),
        ):
            _atomic_replace(Path("source"), Path("destination"), attempts=3)
        self.assertEqual(replace.call_count, 3)
        self.assertEqual(sleep.call_count, 2)


if __name__ == "__main__":
    unittest.main()
