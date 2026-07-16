from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import shutil
import sys
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = PROJECT_ROOT / "03_src"
if str(SOURCE_ROOT) not in sys.path:
    sys.path.insert(0, str(SOURCE_ROOT))

from oil_world_data.generator import build_release  # noqa: E402
from oil_world_data.paths import resolve_raw_data_root  # noqa: E402


SOURCE_RANGE = re.compile(
    r"(?P<sheet>[A-Za-z0-9_]+)!(?P<start_col>[A-Z]+)(?P<start_row>\d+):(?P<end_col>[A-Z]+)(?P<end_row>\d+)"
)
FIELDS = [
    "report_id", "sheet", "cell", "product", "region", "metric", "period",
    "raw_html_value", "old_excel_value", "new_excel_value", "difference_type",
    "mapped_to_release", "impact_level", "quality_note",
]


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def equal(left: Any, right: Any) -> bool:
    if left is None and right is None:
        return True
    if isinstance(left, float) and math.isnan(left):
        left = None
    if isinstance(right, float) and math.isnan(right):
        right = None
    if left is None or right is None:
        return left is right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(float(left), float(right), rel_tol=0, abs_tol=1e-9)
    return left == right


def classify(old: Any, new: Any, raw: Any = "", marker: str = "") -> str:
    if equal(old, new):
        return "footnote_preserved" if "*" in marker else "identical"
    if old is None and isinstance(new, (int, float)) and float(new) == 0:
        return "zero_restored"
    if old is None:
        return "missing_in_old"
    if new is None:
        return "missing_in_new"
    if isinstance(old, (int, float)) and isinstance(new, (int, float)):
        return "numeric_difference"
    def numeric_text(value: Any) -> float | None:
        text = str(value).strip().replace(",", "")
        text = re.sub(r"\*+", "", text).strip()
        text = text.removesuffix("%").strip()
        negative = text.startswith("(") and text.endswith(")")
        if negative:
            text = text[1:-1].strip()
        try:
            number = float(text)
        except ValueError:
            return None
        return -abs(number) if negative else number
    old_number, new_number = numeric_text(old), numeric_text(new)
    if old_number is not None and new_number is not None and math.isclose(old_number, new_number, rel_tol=0, abs_tol=1e-9):
        if "*" in str(raw) or "*" in marker:
            return "footnote_preserved"
        return "format_only"
    if re.sub(r"\s+", " ", str(old)).strip() == re.sub(r"\s+", " ", str(new)).strip():
        return "format_only"
    return "structure_difference"


def load_metadata(workbook: openpyxl.Workbook) -> dict[tuple[str, str], tuple[Any, Any, str]]:
    metadata: dict[tuple[str, str], tuple[Any, Any, str]] = {}
    sheet = workbook["__cell_metadata"]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    indexes = {name: header.index(name) for name in ["sheet", "cell", "raw_value", "normalized_value", "footnote_marker"]}
    for row in rows:
        key = (str(row[indexes["sheet"]]), str(row[indexes["cell"]]))
        metadata[key] = (
            row[indexes["raw_value"]], row[indexes["normalized_value"]], str(row[indexes["footnote_marker"]] or "")
        )
    return metadata


def build_candidate_snapshot(candidate: Path, destination: Path) -> tuple[Path, dict[str, Any]]:
    config = read_json(PROJECT_ROOT / "02_configs" / "release_2026-06.json")
    shutil.copytree(PROJECT_ROOT / config["audit_directory"], destination / config["audit_directory"])
    source = destination / "source" / candidate.name
    source.parent.mkdir(parents=True)
    shutil.copy2(candidate, source)
    source_hash = sha256(candidate)
    audit_path = destination / config["audit_directory"] / config["audit_files"]["mapping_audit"]
    audit = read_json(audit_path)
    audit["source_workbook_sha256"] = source_hash
    write_json(audit_path, audit)
    config.update(
        {
            "source_workbook": f"source/{candidate.name}",
            "source_workbook_sha256": source_hash,
            "expected_annual_change_count": None,
        }
    )
    (destination / "02_configs").mkdir()
    write_json(destination / "02_configs" / "release_2026-06.json", config)
    result = build_release(destination)
    return destination / "01_data" / "releases" / "2026-06", result


def metric_map(payload: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {row["metric"]: row for row in payload["metrics"]}


def compare_snapshots(candidate_root: Path, official_root: Path) -> dict[str, Any]:
    candidate_index = read_json(candidate_root / "index.json")
    official_index = read_json(official_root / "index.json")
    candidate_files = {(x["system"], x["product"], x["region"]): x["path"] for x in candidate_index["files"]}
    official_files = {(x["system"], x["product"], x["region"]): x["path"] for x in official_index["files"]}
    changed_combinations: list[dict[str, Any]] = []
    metric_differences: list[dict[str, Any]] = []
    for key in sorted(set(candidate_files) | set(official_files)):
        if key not in candidate_files or key not in official_files:
            changed_combinations.append({"key": key, "reason": "combination_missing"})
            continue
        candidate = read_json(candidate_root / candidate_files[key])
        official = read_json(official_root / official_files[key])
        cm, om = metric_map(candidate), metric_map(official)
        for metric in sorted(set(cm) | set(om)):
            if metric not in cm or metric not in om:
                metric_differences.append({"key": key, "metric": metric, "field": "metric_presence"})
                continue
            for field in ["mapping_status", "unit", "periods", "values", "annual_change"]:
                if cm[metric].get(field) != om[metric].get(field):
                    metric_differences.append(
                        {"key": key, "metric": metric, "field": field, "old": om[metric].get(field), "new": cm[metric].get(field)}
                    )
        if metric_differences and any(item["key"] == key for item in metric_differences):
            changed_combinations.append({"key": key, "reason": "business_metric_difference"})
    return {
        "candidate_index": {key: candidate_index.get(key) for key in ["combination_count", "mapping_record_count", "numeric_observation_count", "annual_change_count", "status_counts"]},
        "official_index": {key: official_index.get(key) for key in ["combination_count", "mapping_record_count", "numeric_observation_count", "annual_change_count", "status_counts"]},
        "combination_keys_equal": set(candidate_files) == set(official_files),
        "changed_combination_count": len(changed_combinations),
        "changed_combinations": changed_combinations,
        "metric_difference_count": len(metric_differences),
        "metric_differences": metric_differences,
    }


def main() -> int:
    output = PROJECT_ROOT / "06_outputs" / "workbook_rebuild" / "2026-06"
    original_root = resolve_raw_data_root(PROJECT_ROOT)
    old_path = next((original_root / "2026-06").glob("*.xlsx"))
    new_path = next(output.glob("*.xlsx"))
    audit = read_json(next((PROJECT_ROOT / "07_docs").glob("*/2026-06_report_mapping_audit.json")))
    coverage = audit["coverage_matrix"]
    config = read_json(PROJECT_ROOT / "02_configs" / "workbook_rebuild.json")
    # Normal worksheets provide O(1) random cell access. Read-only worksheets
    # rescan XML for each .cell() call and are prohibitively slow for a complete
    # 564-sheet coordinate comparison.
    old_book = openpyxl.load_workbook(old_path, read_only=False, data_only=True)
    new_book = openpyxl.load_workbook(new_path, read_only=False, data_only=True)
    metadata = load_metadata(new_book)
    old_visible = old_book.sheetnames
    new_visible = [name for name in new_book.sheetnames if new_book[name].sheet_state == "visible"]
    rows: list[dict[str, Any]] = []
    mapping_results: list[dict[str, Any]] = []
    for mapping in coverage:
        periods = [part.strip() for part in str(mapping.get("market_years") or "").split("|") if part.strip()]
        cell_results: list[str] = []
        for match in SOURCE_RANGE.finditer(str(mapping.get("source_cell_or_range") or "")):
            sheet = match.group("sheet")
            min_col, min_row, max_col, max_row = range_boundaries(
                f"{match.group('start_col')}{match.group('start_row')}:{match.group('end_col')}{match.group('end_row')}"
            )
            for row_number in range(min_row, max_row + 1):
                for offset, column in enumerate(range(min_col, max_col + 1)):
                    cell = f"{get_column_letter(column)}{row_number}"
                    old = old_book[sheet].cell(row_number, column).value if sheet in old_book.sheetnames else None
                    new = new_book[sheet].cell(row_number, column).value if sheet in new_book.sheetnames else None
                    raw, _, marker = metadata.get((sheet, cell), ("", new, ""))
                    difference = classify(old, new, raw, marker)
                    cell_results.append(difference)
                    rows.append(
                        {
                            "report_id": mapping.get("primary_report_id", sheet), "sheet": sheet, "cell": cell,
                            "product": mapping.get("commodity", ""), "region": mapping.get("country_or_region", ""),
                            "metric": mapping.get("metric", ""), "period": periods[offset] if offset < len(periods) else "",
                            "raw_html_value": raw, "old_excel_value": old, "new_excel_value": new,
                            "difference_type": difference, "mapped_to_release": True,
                            "impact_level": "none" if difference in {"identical", "footnote_preserved", "format_only"} else "high",
                            "quality_note": "星号保存在__cell_metadata。" if difference == "footnote_preserved" else "",
                        }
                    )
        mapping_results.append(
            {
                "system": mapping.get("system"), "product": mapping.get("commodity"),
                "region": mapping.get("country_or_region"), "metric": mapping.get("metric"),
                "mapping_status": mapping.get("mapping_status"),
                "source_cells_compared": len(cell_results),
                "difference_types": dict(Counter(cell_results)),
                "business_values_equal": all(item in {"identical", "footnote_preserved", "format_only", "zero_restored"} for item in cell_results),
            }
        )
    # Scan every visible sheet for non-mapped value changes. Blank trailing rows are summarized as format-only dimensions.
    mapped_keys = {(row["sheet"], row["cell"]) for row in rows}
    visible_value_differences = 0
    dimension_changes: list[dict[str, Any]] = []
    for sheet in sorted(set(old_visible) & set(new_visible)):
        old_sheet, new_sheet = old_book[sheet], new_book[sheet]
        if (old_sheet.max_row, old_sheet.max_column) != (new_sheet.max_row, new_sheet.max_column):
            dimension_changes.append(
                {"sheet": sheet, "old": [old_sheet.max_row, old_sheet.max_column], "new": [new_sheet.max_row, new_sheet.max_column]}
            )
            rows.append(
                {
                    "report_id": sheet, "sheet": sheet, "cell": "used_range", "product": "", "region": "", "metric": "",
                    "period": "", "raw_html_value": "", "old_excel_value": f"{old_sheet.max_row}x{old_sheet.max_column}",
                    "new_excel_value": f"{new_sheet.max_row}x{new_sheet.max_column}", "difference_type": "format_only",
                    "mapped_to_release": False, "impact_level": "none",
                    "quality_note": "统一生成器保留了表格后的空白分隔行；可见数据坐标未发生业务变化。",
                }
            )
        for row_number in range(1, max(old_sheet.max_row, new_sheet.max_row) + 1):
            for column in range(1, max(old_sheet.max_column, new_sheet.max_column) + 1):
                cell = f"{get_column_letter(column)}{row_number}"
                if (sheet, cell) in mapped_keys:
                    continue
                old = old_sheet.cell(row_number, column).value
                new = new_sheet.cell(row_number, column).value
                if equal(old, new):
                    continue
                raw, _, marker = metadata.get((sheet, cell), ("", new, ""))
                difference = classify(old, new, raw, marker)
                visible_value_differences += 1
                rows.append(
                    {
                        "report_id": sheet, "sheet": sheet, "cell": cell, "product": "", "region": "", "metric": "",
                        "period": "", "raw_html_value": raw, "old_excel_value": old, "new_excel_value": new,
                        "difference_type": difference, "mapped_to_release": False,
                        "impact_level": "low" if difference in {"zero_restored", "footnote_preserved", "format_only"} else "high",
                        "quality_note": "正式映射范围外的工作簿差异。",
                    }
                )
    old_book.close()
    new_book.close()
    with tempfile.TemporaryDirectory(prefix="oil_world_june_compare_") as temporary:
        snapshot, build_result = build_candidate_snapshot(new_path, Path(temporary))
        snapshot_comparison = compare_snapshots(snapshot, PROJECT_ROOT / "01_data" / "releases" / "2026-06")
    counts = Counter(row["difference_type"] for row in rows)
    result = {
        "old_workbook": {"path": str(old_path), "size": old_path.stat().st_size, "sha256": sha256(old_path), "sheet_count": len(old_visible)},
        "new_workbook": {"path": str(new_path), "size": new_path.stat().st_size, "sha256": sha256(new_path), "sheet_count": len(new_visible) + 2, "visible_sheet_count": len(new_visible)},
        "visible_sheet_names_equal": old_visible == new_visible,
        "hidden_trace_sheets": ["__cell_metadata", "__footnotes"],
        "required_reports_ok": sum(report in new_visible and report in old_visible for report in config["required_report_ids"]),
        "dimension_change_count": len(dimension_changes),
        "dimension_changes": dimension_changes,
        "visible_value_difference_count": visible_value_differences,
        "mapping_record_count": len(mapping_results),
        "mapping_records_business_equal": sum(row["business_values_equal"] for row in mapping_results),
        "mapping_status_counts": dict(Counter(row["mapping_status"] for row in mapping_results)),
        "difference_type_counts": dict(counts),
        "candidate_snapshot_build": build_result,
        "snapshot_comparison": snapshot_comparison,
    }
    with (output / "difference_report.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    write_json(output / "difference_report.json", rows)
    write_json(output / "june_comparison_report.json", result)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    severe = sum(counts[name] for name in ["numeric_difference", "structure_difference", "missing_in_new"])
    severe += snapshot_comparison["metric_difference_count"] + snapshot_comparison["changed_combination_count"]
    return 2 if severe else 0


if __name__ == "__main__":
    raise SystemExit(main())
