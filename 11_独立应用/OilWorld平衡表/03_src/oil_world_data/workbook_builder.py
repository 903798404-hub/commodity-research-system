from __future__ import annotations

import hashlib
import importlib.metadata
import json
import logging
import math
import os
import re
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from io import StringIO
from numbers import Number
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from bs4 import BeautifulSoup
from openpyxl.utils import get_column_letter


MONTH_NAMES = {
    1: "January",
    2: "February",
    3: "March",
    4: "April",
    5: "May",
    6: "June",
    7: "July",
    8: "August",
    9: "September",
    10: "October",
    11: "November",
    12: "December",
}
INVALID_SHEET_CHARACTERS = re.compile(r"[\[\]:*?/\\]")
NUMERIC_TEXT = re.compile(
    r"^(?P<open>\()?\s*(?P<number>[+-]?(?:\d[\d,]*\.?\d*|\.\d+))\s*(?P<percent>%)?\s*(?P<close>\))?$"
)
TRAILING_MARKER = re.compile(r"(?P<marker>\*+)\s*$")
BLANK_TOKENS = {"", ".", "-", "nan", "na", "n/a", "none", "null"}
IGNORE_PHRASES = [
    "ista mielke",
    "oil world",
    "information provider",
    "more details see",
    "http",
    "production",
    "stocks",
    "supply",
    "demand",
    "crush",
    "exports",
    "imports",
    "oil",
    "world",
    "balance",
    "commodity",
]


class WorkbookBuildError(RuntimeError):
    pass


@dataclass(frozen=True)
class CatalogEntry:
    section: str
    category: str
    report_title: str
    source_href: str
    source_file: Path


@dataclass
class ExtractedReport:
    values: pd.DataFrame
    raw_values: pd.DataFrame
    markers: pd.DataFrame
    input_files: set[Path]
    footnotes: list[str]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def atomic_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(temporary, path)


def release_label(release: str) -> str:
    match = re.fullmatch(r"(20\d{2})-(0[1-9]|1[0-2])", release)
    if not match:
        raise WorkbookBuildError(f"发布期必须使用 YYYY-MM：{release}")
    year, month = int(match.group(1)), int(match.group(2))
    return f"{MONTH_NAMES[month]} {year}"


def output_filename(release: str) -> str:
    return f"油世界季度表-{release_label(release)}.xlsx"


def read_soup(path: Path) -> BeautifulSoup:
    if not path.is_file():
        raise WorkbookBuildError(f"HTML文件不存在：{path}")
    return BeautifulSoup(path.read_text(encoding="iso-8859-1", errors="ignore"), "html.parser")


def normalize_href(value: str) -> str:
    return value.replace("\\", "/").strip()


def resolve_local_href(base_directory: Path, href: str) -> Path:
    normalized = normalize_href(href)
    if re.match(r"^[a-z]+://", normalized, flags=re.I):
        raise WorkbookBuildError(f"不允许联网读取HTML：{href}")
    return (base_directory / Path(*[part for part in normalized.split("/") if part not in {"", "."}])).resolve()


def build_catalog(
    source_dir: Path,
    config: dict[str, Any],
    report_aliases: dict[str, str] | None = None,
) -> tuple[Path, list[CatalogEntry], set[Path]]:
    start_files = sorted(source_dir.glob(config["start_pattern"]))
    if len(start_files) != 1:
        raise WorkbookBuildError(
            f"START文件必须且只能有一个，当前发现 {len(start_files)} 个：{source_dir}"
        )
    start_file = start_files[0]
    start_soup = read_soup(start_file)
    entries: list[CatalogEntry] = []
    inputs = {start_file}
    missing_index: list[str] = []
    for section in start_soup.find_all("h3"):
        parent = section.find_parent("td")
        if not parent:
            continue
        for index_link in parent.find_all("a"):
            href = normalize_href(index_link.get("href") or "")
            if not href.casefold().endswith(".htm"):
                continue
            index_path = resolve_local_href(source_dir, href)
            if not index_path.is_file():
                missing_index.append(href)
                continue
            inputs.add(index_path)
            index_name = index_link.get_text(strip=True)
            index_soup = read_soup(index_path)
            for report_link in index_soup.find_all("a"):
                report_href = normalize_href(report_link.get("href") or "")
                if "stats/" not in report_href.casefold() or not report_href.casefold().endswith(".htm"):
                    continue
                report_file = source_dir / config["stats_directory"] / Path(report_href).name
                previous = report_link.find_previous(["b", "strong"])
                category = index_name
                if previous:
                    category_text = previous.get_text(strip=True)
                    if category_text and "Copyright" not in category_text and "OIL WORLD" not in category_text:
                        if category_text.casefold() in {
                            "oilseeds",
                            "oils & fats",
                            "oilmeals",
                            "oils & fats and biodiesel",
                        }:
                            category = f"{index_name} - {category_text}"
                        else:
                            category = category_text
                entries.append(
                    CatalogEntry(
                        section=section.get_text(strip=True),
                        category=category,
                        report_title=report_link.get_text(strip=True),
                        source_href=report_href,
                        source_file=report_file.resolve(),
                    )
                )
    if missing_index:
        raise WorkbookBuildError(f"目录HTML缺失：{len(missing_index)} 个；示例：{missing_index[:10]}")
    if not entries:
        raise WorkbookBuildError("START和目录页中没有识别到stats报表")
    missing_stats = sorted({entry.source_file for entry in entries if not entry.source_file.is_file()})
    if missing_stats:
        examples = [path.name for path in missing_stats[:20]]
        raise WorkbookBuildError(f"stats引用不完整：缺失 {len(missing_stats)} 个报表；示例：{examples}")
    inputs.update(entry.source_file for entry in entries)
    report_aliases = report_aliases or {}
    required = set(config["required_report_ids"])
    referenced = {report_aliases.get(entry.source_file.stem, entry.source_file.stem) for entry in entries}
    missing_required = sorted(required - referenced)
    if missing_required:
        raise WorkbookBuildError(f"30张正式采用报表中缺失：{missing_required}")
    return start_file, entries, inputs


def split_marker(text: str) -> tuple[str, str]:
    match = TRAILING_MARKER.search(text)
    if not match:
        return text.strip(), ""
    return text[: match.start()].strip(), match.group("marker")


def normalize_cell(value: Any) -> tuple[Any, str, str]:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None, "", ""
    raw = str(value).replace("\xa0", " ").strip()
    base, marker = split_marker(raw)
    if base.casefold() in BLANK_TOKENS:
        return None, marker, raw
    match = NUMERIC_TEXT.fullmatch(base)
    if match:
        number = float(match.group("number").replace(",", ""))
        if match.group("open") and match.group("close"):
            number = -abs(number)
        return number, marker, raw
    return base, marker, raw


def _normalize_table(table: Any) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame] | None:
    try:
        frames = pd.read_html(StringIO(str(table)), header=None, keep_default_na=False)
    except (ValueError, ImportError):
        return None
    if not frames or frames[0].empty or len(frames[0]) <= 1:
        return None
    source = frames[0]
    normalized = source.map(lambda value: normalize_cell(value)[0])
    raw = source.map(lambda value: normalize_cell(value)[2])
    markers = source.map(lambda value: normalize_cell(value)[1])
    # Oil World commonly stores an estimate/footnote marker in a narrow cell
    # immediately following the value. Preserve that marker on the value before
    # the marker-only column is removed as layout noise.
    for row in range(normalized.shape[0]):
        for column in range(1, normalized.shape[1]):
            current = normalized.iat[row, column]
            marker = markers.iat[row, column]
            raw_current = raw.iat[row, column]
            previous = normalized.iat[row, column - 1]
            inline_marker = ""
            if marker and (current is None or pd.isna(current)):
                inline_marker = marker
            elif isinstance(current, str) and current.casefold() in {"p", "f", "e"}:
                inline_marker = current
            if not inline_marker or previous is None or pd.isna(previous):
                continue
            existing_marker = markers.iat[row, column - 1]
            markers.iat[row, column - 1] = f"{existing_marker}{inline_marker}"
            previous_raw = raw.iat[row, column - 1]
            raw.iat[row, column - 1] = f"{previous_raw}{inline_marker}"
            normalized.iat[row, column] = None
            raw.iat[row, column] = raw_current
            markers.iat[row, column] = ""
    valid_columns = [0]
    for column in range(1, normalized.shape[1]):
        meaningful = False
        for cell in normalized.iloc[:, column]:
            if cell is None or pd.isna(cell):
                continue
            if isinstance(cell, Number) and not isinstance(cell, bool):
                meaningful = True
                break
            if column == 1 and isinstance(cell, str) and len(cell.strip()) > 1:
                lowered = cell.strip().casefold()
                if not any(phrase in lowered for phrase in IGNORE_PHRASES):
                    meaningful = True
                    break
        if meaningful:
            valid_columns.append(column)
    normalized = normalized.iloc[:, valid_columns].copy()
    raw = raw.iloc[:, valid_columns].copy()
    markers = markers.iloc[:, valid_columns].copy()
    if normalized.shape[1] <= 1:
        return None
    for frame in (normalized, raw, markers):
        frame.columns = range(frame.shape[1])
    return normalized, raw, markers


def _footnote_lines(soup: BeautifulSoup) -> list[str]:
    result: list[str] = []
    for line in soup.get_text("\n").splitlines():
        text = re.sub(r"\s+", " ", line).strip()
        if re.match(r"^(?:\*+|\([A-Za-z0-9]+\)|[A-Za-z0-9]+\))\s*\S", text) and not re.fullmatch(
            r"\*+", text
        ):
            result.append(text[:500])
    return list(dict.fromkeys(result))


def extract_report(path: Path, visited: set[Path] | None = None) -> ExtractedReport | None:
    visited = set() if visited is None else visited
    path = path.resolve()
    if path in visited:
        raise WorkbookBuildError(f"HTML框架形成循环引用：{path}")
    visited.add(path)
    soup = read_soup(path)
    input_files = {path}
    tables = soup.find_all("table")
    if not tables:
        frame = soup.find("frame", src=re.compile(r"_body\.htm", re.I))
        if not frame:
            frame = soup.find("a", href=re.compile(r"_body\.htm", re.I))
        if frame:
            href = frame.get("src") or frame.get("href")
            body_path = resolve_local_href(path.parent, href)
            if not body_path.is_file():
                raise WorkbookBuildError(f"_body引用缺失：{path.name} -> {href}")
            extracted = extract_report(body_path, visited)
            if extracted:
                extracted.input_files.add(path)
            return extracted
        return None

    value_frames: list[pd.DataFrame] = []
    raw_frames: list[pd.DataFrame] = []
    marker_frames: list[pd.DataFrame] = []
    for table in tables:
        title = ""
        previous = table.previous_element
        for _ in range(40):
            if not previous:
                break
            if getattr(previous, "name", None) in {"u", "strong", "b", "font"}:
                text = previous.get_text(strip=True)
                if len(text) > 5:
                    title = text
                    break
            previous = previous.previous_element
        normalized = _normalize_table(table)
        if normalized is None:
            continue
        values, raw, markers = normalized
        if title:
            width = values.shape[1]
            value_frames.append(pd.DataFrame([[title] + [None] * (width - 1)]))
            raw_frames.append(pd.DataFrame([[title] + [""] * (width - 1)]))
            marker_frames.append(pd.DataFrame([[""] * width]))
        value_frames.append(values)
        raw_frames.append(raw)
        marker_frames.append(markers)
        width = values.shape[1]
        value_frames.append(pd.DataFrame([[None] * width]))
        raw_frames.append(pd.DataFrame([[""] * width]))
        marker_frames.append(pd.DataFrame([[""] * width]))
    if not value_frames:
        return None

    def concatenate(frames: list[pd.DataFrame]) -> pd.DataFrame:
        normalized_frames = []
        for frame in frames:
            copied = frame.copy()
            copied.columns = range(copied.shape[1])
            normalized_frames.append(copied)
        return pd.concat(normalized_frames, ignore_index=True)

    values = concatenate(value_frames)
    raw_values = concatenate(raw_frames)
    markers = concatenate(marker_frames)
    values.iloc[:, 0] = values.iloc[:, 0].apply(
        lambda value: (
            None
            if value is None or pd.isna(value)
            else re.sub(r"[\.\s]+$", "", str(value))
        )
    )
    return ExtractedReport(
        values=values,
        raw_values=raw_values,
        markers=markers,
        input_files=input_files,
        footnotes=_footnote_lines(soup),
    )


def safe_sheet_name(source_name: str, used_names: set[str]) -> str:
    base = INVALID_SHEET_CHARACTERS.sub("_", source_name).strip("'") or "Sheet"
    base = base[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in {name.casefold() for name in used_names}:
        marker = f"_{suffix}"
        candidate = f"{base[: 31 - len(marker)]}{marker}"
        suffix += 1
    used_names.add(candidate)
    return candidate


def scan_html_tokens(paths: set[Path]) -> dict[str, Any]:
    zero_count = starred_numeric_count = 0
    footnotes: list[dict[str, str]] = []
    for path in sorted(paths):
        soup = read_soup(path)
        for table in soup.find_all("table"):
            normalized_table = _normalize_table(table)
            if normalized_table is None:
                continue
            values, _, markers = normalized_table
            for row in range(values.shape[0]):
                for column in range(values.shape[1]):
                    normalized = values.iat[row, column]
                    marker = markers.iat[row, column]
                    if isinstance(normalized, Number) and not isinstance(normalized, bool):
                        if float(normalized) == 0:
                            zero_count += 1
                        if "*" in str(marker):
                            starred_numeric_count += 1
        for line in _footnote_lines(soup):
            footnotes.append({"source_file": path.name, "text": line})
    unique_footnotes = list({(item["source_file"], item["text"]): item for item in footnotes}.values())
    return {
        "html_zero_count": zero_count,
        "starred_numeric_count": starred_numeric_count,
        "footnote_definition_count": len(unique_footnotes),
        "footnotes": unique_footnotes,
    }


def validate_xlsx(path: Path, required_reports: list[str], metadata_sheet: str) -> dict[str, Any]:
    if not path.is_file() or path.stat().st_size <= 0:
        raise WorkbookBuildError("候选XLSX不存在或为0字节")
    if not zipfile.is_zipfile(path):
        raise WorkbookBuildError("候选XLSX不是有效ZIP容器")
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise WorkbookBuildError(f"XLSX ZIP成员校验失败：{bad_member}")
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    try:
        missing = sorted(set(required_reports) - set(workbook.sheetnames))
        if missing:
            raise WorkbookBuildError(f"候选工作簿缺少正式采用报表：{missing}")
        if metadata_sheet not in workbook.sheetnames:
            raise WorkbookBuildError("候选工作簿缺少单元格来源元数据工作表")
        if workbook[metadata_sheet].sheet_state != "hidden":
            raise WorkbookBuildError("单元格来源元数据工作表必须隐藏")
        return {
            "sheet_count": len(workbook.sheetnames),
            "visible_sheet_count": sum(workbook[name].sheet_state == "visible" for name in workbook.sheetnames),
            "required_report_count": len(required_reports),
            "missing_required_reports": missing,
        }
    finally:
        workbook.close()


def dependency_versions() -> dict[str, str]:
    result: dict[str, str] = {}
    for package in ["pandas", "beautifulsoup4", "openpyxl", "lxml"]:
        try:
            result[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            result[package] = "missing"
    return result


def build_workbook(
    *,
    release: str,
    source_dir: Path,
    output_dir: Path,
    config_path: Path,
    cli_path: Path,
    logger: logging.Logger,
) -> dict[str, Any]:
    source_dir = source_dir.resolve()
    output_dir = output_dir.resolve()
    config_path = config_path.resolve()
    if not source_dir.is_dir():
        raise WorkbookBuildError(f"输入目录不存在：{source_dir}")
    try:
        output_dir.relative_to(source_dir)
    except ValueError:
        pass
    else:
        raise WorkbookBuildError("输出目录不得位于正式原始资料目录内")
    config = json.loads(config_path.read_text(encoding="utf-8"))
    label = release_label(release)
    report_aliases = config.get("report_id_aliases", {}).get(release, {})
    start_file, catalog, initial_inputs = build_catalog(source_dir, config, report_aliases)
    if label.casefold() not in start_file.stem.casefold():
        raise WorkbookBuildError(f"START文件名与发布期不匹配：{start_file.name} / {release}")
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / output_filename(release)
    manifest_path = output_dir / "manifest.json"
    if target.exists() or manifest_path.exists():
        raise WorkbookBuildError(f"候选输出已存在，默认拒绝覆盖：{target}")
    logger.info("START=%s", start_file)
    logger.info("目录记录=%d，唯一stats报表=%d", len(catalog), len({item.source_file for item in catalog}))

    used_names = {"目录", config["metadata_sheet"], config["footnotes_sheet"]}
    sheet_names: dict[Path, str] = {}
    for entry in catalog:
        source_id = entry.source_file.stem
        canonical_id = report_aliases.get(source_id, source_id)
        sheet_names.setdefault(entry.source_file, safe_sheet_name(canonical_id, used_names))
    missing_named = sorted(set(config["required_report_ids"]) - set(sheet_names.values()))
    if missing_named:
        raise WorkbookBuildError(f"工作表命名后无法匹配正式采用报表：{missing_named}")

    handle, temporary_name = tempfile.mkstemp(prefix=f".{release}_", suffix=".xlsx", dir=output_dir)
    os.close(handle)
    temporary = Path(temporary_name)
    temporary.unlink()
    metadata_rows: list[dict[str, Any]] = []
    footnote_rows: list[dict[str, str]] = []
    parsed_reports = 0
    parsed_sheet_names: set[str] = set()
    skipped_reports: list[str] = []
    input_files = set(initial_inputs)
    try:
        catalog_frame = pd.DataFrame(
            [
                {
                    "板块": item.section,
                    "品种/国家": item.category,
                    "报表名称": item.report_title,
                    "数据文件名": item.source_file.name,
                }
                for item in catalog
            ]
        )
        with pd.ExcelWriter(temporary, engine=config["excel_engine"]) as writer:
            catalog_frame.to_excel(writer, sheet_name="目录", index=False)
            for report_file in dict.fromkeys(item.source_file for item in catalog):
                extracted = extract_report(report_file)
                if extracted is None:
                    canonical_id = sheet_names[report_file]
                    if canonical_id in set(config["required_report_ids"]):
                        raise WorkbookBuildError(f"正式采用报表没有可生成的数据表：{report_file.name}")
                    skipped_reports.append(report_file.name)
                    logger.warning("跳过无数据目录页：%s", report_file.name)
                    continue
                sheet_name = sheet_names[report_file]
                extracted.values.to_excel(writer, sheet_name=sheet_name, index=False, header=False)
                parsed_reports += 1
                parsed_sheet_names.add(sheet_name)
                input_files.update(extracted.input_files)
                for row_index in range(extracted.values.shape[0]):
                    for column_index in range(extracted.values.shape[1]):
                        raw = extracted.raw_values.iat[row_index, column_index]
                        marker = extracted.markers.iat[row_index, column_index]
                        normalized = extracted.values.iat[row_index, column_index]
                        if raw in {None, ""} and not marker:
                            continue
                        if isinstance(normalized, float) and math.isnan(normalized):
                            normalized = None
                        metadata_rows.append(
                            {
                                "sheet": sheet_name,
                                "cell": f"{get_column_letter(column_index + 1)}{row_index + 1}",
                                "raw_value": raw,
                                "normalized_value": normalized,
                                "footnote_marker": marker,
                                "source_file": report_file.name,
                            }
                        )
                for definition in extracted.footnotes:
                    footnote_rows.append(
                        {"sheet": sheet_name, "source_file": report_file.name, "footnote_text": definition}
                    )
            pd.DataFrame(
                metadata_rows,
                columns=["sheet", "cell", "raw_value", "normalized_value", "footnote_marker", "source_file"],
            ).to_excel(writer, sheet_name=config["metadata_sheet"], index=False)
            pd.DataFrame(
                footnote_rows,
                columns=["sheet", "source_file", "footnote_text"],
            ).drop_duplicates().to_excel(writer, sheet_name=config["footnotes_sheet"], index=False)

            workbook = writer.book
            menu = workbook["目录"]
            for row_number, entry in enumerate(catalog, start=2):
                sheet_name = sheet_names[entry.source_file]
                cell = menu.cell(row=row_number, column=5, value=entry.source_file.name)
                if sheet_name in parsed_sheet_names:
                    cell.hyperlink = f"#'{sheet_name}'!A1"
                    cell.style = "Hyperlink"
            workbook[config["metadata_sheet"]].sheet_state = "hidden"
            workbook[config["footnotes_sheet"]].sheet_state = "hidden"
            menu.freeze_panes = "A2"

        validation = validate_xlsx(temporary, config["required_report_ids"], config["metadata_sheet"])
        unique_footnotes = list(
            {
                (item["source_file"], item["footnote_text"]): {
                    "source_file": item["source_file"],
                    "text": item["footnote_text"],
                }
                for item in footnote_rows
            }.values()
        )
        source_stats = {
            "html_zero_count": sum(
                isinstance(item["normalized_value"], Number)
                and not isinstance(item["normalized_value"], bool)
                and float(item["normalized_value"]) == 0
                for item in metadata_rows
            ),
            "starred_numeric_count": sum(
                isinstance(item["normalized_value"], Number)
                and not isinstance(item["normalized_value"], bool)
                and "*" in str(item["footnote_marker"])
                for item in metadata_rows
            ),
            "footnote_definition_count": len(unique_footnotes),
            "footnotes": unique_footnotes,
        }
        output_hash = sha256(temporary)
        generated_at = datetime.now(timezone.utc).isoformat()
        manifest = {
            "schema_version": 1,
            "builder_version": config["builder_version"],
            "release": release,
            "release_label": label,
            "generated_at": generated_at,
            "source_directory": str(source_dir),
            "start_file": start_file.name,
            "catalog_record_count": len(catalog),
            "unique_report_count": len({item.source_file for item in catalog}),
            "parsed_report_count": parsed_reports,
            "skipped_report_count": len(skipped_reports),
            "skipped_reports": skipped_reports,
            "required_report_count": len(config["required_report_ids"]),
            "required_reports": config["required_report_ids"],
            "report_id_aliases": report_aliases,
            "sheet_count": validation["sheet_count"],
            "visible_sheet_count": validation["visible_sheet_count"],
            "metadata_record_count": len(metadata_rows),
            "footnote_record_count": len({tuple(item.values()) for item in footnote_rows}),
            "html_zero_count": source_stats["html_zero_count"],
            "starred_numeric_count": source_stats["starred_numeric_count"],
            "input_file_count": len(input_files),
            "input_file_hashes": {
                str(path.relative_to(source_dir)).replace("\\", "/"): sha256(path)
                for path in sorted(input_files)
            },
            "config_file": str(config_path),
            "config_sha256": sha256(config_path),
            "cli_script_sha256": sha256(cli_path),
            "core_script_sha256": sha256(Path(__file__)),
            "dependencies": dependency_versions(),
            "output_file": target.name,
            "output_size": temporary.stat().st_size,
            "output_sha256": output_hash,
            "quality_status": "passed",
        }
        footnotes_payload = {
            "release": release,
            "count": source_stats["footnote_definition_count"],
            "definitions": source_stats["footnotes"],
        }
        os.replace(temporary, target)
        try:
            atomic_json(manifest_path, manifest)
            atomic_json(output_dir / "footnotes.json", footnotes_payload)
        except Exception:
            target.unlink(missing_ok=True)
            raise
        logger.info("候选工作簿生成成功：%s", target)
        logger.info("size=%d sha256=%s sheets=%d", target.stat().st_size, output_hash, validation["sheet_count"])
        return {**manifest, "output_path": str(target)}
    finally:
        temporary.unlink(missing_ok=True)
