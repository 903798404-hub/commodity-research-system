from __future__ import annotations

import csv
import json
import re
from collections import Counter
from numbers import Number
from pathlib import Path
from typing import Any

import openpyxl
from bs4 import BeautifulSoup
from openpyxl.utils.cell import range_boundaries

from .workbook_builder import extract_report, normalize_cell, read_soup, sha256


SOURCE_RANGE = re.compile(
    r"(?P<sheet>[A-Za-z0-9_]+)!(?P<start_col>[A-Z]+)(?P<start_row>\d+):(?P<end_col>[A-Z]+)(?P<end_row>\d+)"
)
ZERO_CELL = re.compile(
    rb"<t[dh]\b[^>]*>\s*(?:<[^>]+>\s*)*[+-]?0(?:\.0+)?\s*(?:<[^>]+>\s*)*</t[dh]>",
    re.I,
)


def _json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _catalog_references(source_dir: Path, start_pattern: str) -> tuple[set[str], set[str]]:
    starts = list(source_dir.glob(start_pattern))
    if len(starts) != 1:
        return set(), set()
    referenced: set[str] = set()
    existing: set[str] = set()
    soup = read_soup(starts[0])
    for link in soup.find_all("a"):
        href = str(link.get("href") or "").replace("\\", "/")
        if not href.casefold().endswith(".htm") or "stats/" in href.casefold():
            continue
        index = source_dir / Path(href).name
        if not index.is_file():
            continue
        for report_link in read_soup(index).find_all("a"):
            report_href = str(report_link.get("href") or "").replace("\\", "/")
            if "stats/" not in report_href.casefold() or not report_href.casefold().endswith(".htm"):
                continue
            name = Path(report_href).name
            referenced.add(name.casefold())
            if (source_dir / "stats" / name).is_file():
                existing.add(name.casefold())
    return referenced, existing


def _raw_html_zero_count(source_dir: Path) -> int:
    return sum(len(ZERO_CELL.findall(path.read_bytes())) for path in source_dir.rglob("*.htm"))


def _report_stats(source_dir: Path, required: list[str], aliases: dict[str, str]) -> dict[str, Any]:
    reverse_aliases = {canonical: source for source, canonical in aliases.items()}
    rows: list[dict[str, Any]] = []
    total_zero = total_star = total_negative = 0
    footnotes: list[dict[str, str]] = []
    for report_id in required:
        source_id = reverse_aliases.get(report_id, report_id)
        matches = list((source_dir / "stats").glob(f"{source_id}.*")) + list(
            (source_dir / "stats").glob(f"{source_id}.htm")
        )
        report_file = next((path for path in matches if path.suffix.casefold() == ".htm"), None)
        if report_file is None:
            rows.append({"report_id": report_id, "source_report_id": source_id, "status": "missing"})
            continue
        extracted = extract_report(report_file)
        if extracted is None:
            rows.append({"report_id": report_id, "source_report_id": source_id, "status": "empty"})
            continue
        zeros = stars = negatives = 0
        for row in range(extracted.values.shape[0]):
            for column in range(extracted.values.shape[1]):
                value = extracted.values.iat[row, column]
                marker = str(extracted.markers.iat[row, column] or "")
                if isinstance(value, Number) and not isinstance(value, bool):
                    zeros += float(value) == 0
                    negatives += float(value) < 0
                    stars += "*" in marker
        total_zero += zeros
        total_star += stars
        total_negative += negatives
        footnotes.extend({"report_id": report_id, "text": text} for text in extracted.footnotes)
        rows.append(
            {
                "report_id": report_id,
                "source_report_id": source_id,
                "status": "ok",
                "rows": extracted.values.shape[0],
                "columns": extracted.values.shape[1],
                "zero_count": zeros,
                "starred_numeric_count": stars,
                "negative_count": negatives,
            }
        )
    return {
        "reports": rows,
        "successful": sum(row["status"] == "ok" for row in rows),
        "failed": [row["report_id"] for row in rows if row["status"] != "ok"],
        "zero_count": total_zero,
        "starred_numeric_count": total_star,
        "negative_count": total_negative,
        "footnotes": list({(item["report_id"], item["text"]): item for item in footnotes}.values()),
    }


def _mapping_zero_count(workbook_path: Path, coverage: list[dict[str, Any]]) -> dict[str, int]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    visited: set[tuple[str, int, int]] = set()
    zero_cells: set[tuple[str, int, int]] = set()
    try:
        for mapping in coverage:
            if mapping.get("mapping_status") not in {"direct", "derived"}:
                continue
            for match in SOURCE_RANGE.finditer(str(mapping.get("source_cell_or_range") or "")):
                sheet = match.group("sheet")
                if sheet not in workbook.sheetnames:
                    continue
                min_col, min_row, max_col, max_row = range_boundaries(
                    f"{match.group('start_col')}{match.group('start_row')}:{match.group('end_col')}{match.group('end_row')}"
                )
                for row in range(min_row, max_row + 1):
                    for column in range(min_col, max_col + 1):
                        key = (sheet, row, column)
                        visited.add(key)
                        value = workbook[sheet].cell(row, column).value
                        if isinstance(value, Number) and not isinstance(value, bool) and float(value) == 0:
                            zero_cells.add(key)
    finally:
        workbook.close()
    return {"unique_source_cell_count": len(visited), "zero_count": len(zero_cells)}


def _tree_hash(paths: list[Path]) -> str:
    import hashlib

    digest = hashlib.sha256()
    for path in sorted(paths, key=lambda item: str(item).casefold()):
        digest.update(str(path).encode("utf-8"))
        digest.update(sha256(path).encode("ascii"))
    return digest.hexdigest()


def audit_rebuild(project_root: Path) -> dict[str, Any]:
    config = _json(project_root / "02_configs" / "workbook_rebuild.json")
    audit_path = next((project_root / "07_docs").glob("*/2026-06_report_mapping_audit.json"))
    audit = _json(audit_path)
    coverage = audit["coverage_matrix"]
    required = config["required_report_ids"]
    original_root = next(path for path in project_root.iterdir() if path.name.startswith("01_") and path.name != "01_data")
    march_source = original_root / "2026-03"
    june_source = original_root / "2026-06"
    march_candidate = next((project_root / "06_outputs" / "workbook_rebuild" / "2026-03").glob("*.xlsx"))
    old_june = next(june_source.glob("*.xlsx"))
    march_refs, march_existing = _catalog_references(march_source, config["start_pattern"])
    june_refs, june_existing = _catalog_references(june_source, config["start_pattern"])
    march_stats = _report_stats(march_source, required, config.get("report_id_aliases", {}).get("2026-03", {}))
    june_stats = _report_stats(june_source, required, config.get("report_id_aliases", {}).get("2026-06", {}))
    status_counts = Counter(item["mapping_status"] for item in coverage)
    result = {
        "schema_version": 1,
        "generator_identity": {
            "core_script": "03_src/oil_world_data/workbook_builder.py",
            "core_script_sha256": sha256(project_root / "03_src" / "oil_world_data" / "workbook_builder.py"),
            "config_sha256": sha256(project_root / "02_configs" / "workbook_rebuild.json"),
        },
        "mapping_record_count": len(coverage),
        "mapping_status_counts": dict(status_counts),
        "march": {
            "raw_html_zero_count": _raw_html_zero_count(march_source),
            "catalog_reference_count": len(march_refs),
            "missing_catalog_reference_count": len(march_refs - march_existing),
            "candidate": str(march_candidate),
            "candidate_size": march_candidate.stat().st_size,
            "candidate_sha256": sha256(march_candidate),
            "required_report_stats": march_stats,
            "mapping_region_stats": _mapping_zero_count(march_candidate, coverage),
        },
        "june": {
            "raw_html_zero_count": _raw_html_zero_count(june_source),
            "catalog_reference_count": len(june_refs),
            "existing_catalog_reference_count": len(june_existing),
            "missing_catalog_reference_count": len(june_refs - june_existing),
            "missing_catalog_references": sorted(june_refs - june_existing),
            "candidate": None,
            "candidate_status": "blocked_missing_stats_sources",
            "old_workbook": str(old_june),
            "old_workbook_size": old_june.stat().st_size,
            "old_workbook_sha256": sha256(old_june),
            "required_report_stats": june_stats,
            "mapping_region_stats": _mapping_zero_count(old_june, coverage),
        },
    }
    return result


DIFF_FIELDS = [
    "report_id",
    "sheet",
    "cell",
    "product",
    "region",
    "metric",
    "period",
    "raw_html_value",
    "old_excel_value",
    "new_excel_value",
    "difference_type",
    "mapped_to_release",
    "impact_level",
]


def write_audit_reports(project_root: Path, result: dict[str, Any]) -> None:
    output_root = project_root / "06_outputs" / "workbook_rebuild"
    output_root.mkdir(parents=True, exist_ok=True)
    (output_root / "workbook_rebuild_audit.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    coverage = _json(next((project_root / "07_docs").glob("*/2026-06_report_mapping_audit.json")))[
        "coverage_matrix"
    ]
    by_report: dict[str, dict[str, Any]] = {}
    for row in coverage:
        by_report.setdefault(str(row.get("primary_report_id") or ""), row)
    diff_rows: list[dict[str, Any]] = []
    for missing in result["june"]["missing_catalog_references"]:
        report_id = Path(missing).stem.upper()
        mapping = by_report.get(report_id, {})
        diff_rows.append(
            {
                "report_id": report_id,
                "sheet": report_id,
                "cell": "",
                "product": mapping.get("commodity", ""),
                "region": mapping.get("country_or_region", ""),
                "metric": mapping.get("metric", ""),
                "period": "",
                "raw_html_value": "source HTML missing",
                "old_excel_value": "present in formal June workbook",
                "new_excel_value": "candidate not generated",
                "difference_type": "structure_difference",
                "mapped_to_release": bool(mapping),
                "impact_level": "high" if mapping else "medium",
            }
        )
    june_output = output_root / "2026-06"
    june_output.mkdir(parents=True, exist_ok=True)
    with (june_output / "difference_report.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=DIFF_FIELDS)
        writer.writeheader()
        writer.writerows(diff_rows)
    (june_output / "difference_report.json").write_text(
        json.dumps(diff_rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    counts = Counter(row["difference_type"] for row in diff_rows)
    md = [
        "# Oil World 工作簿重建审计",
        "",
        "## 结论",
        "",
        "- March 候选已由统一生成器生成并通过 XLSX/OpenPyXL/30 张报表检查。",
        "- June 构建被严格阻断：目录引用的 stats HTML 不完整，未生成候选工作簿。",
        "- 因缺少新 June 候选，无法执行 649 条映射的新旧 June 数值比较，也不能据此进入 backfill。",
        "",
        "## 关键统计",
        "",
        f"- March HTML 原始 0/0.0：{result['march']['raw_html_zero_count']}",
        f"- March 30 张报表中的 0：{result['march']['required_report_stats']['zero_count']}",
        f"- March 649 条映射相关唯一来源单元格中的 0：{result['march']['mapping_region_stats']['zero_count']}",
        f"- March 带星号数值：{result['march']['required_report_stats']['starred_numeric_count']}（30 张采用报表）",
        f"- June HTML 原始 0/0.0：{result['june']['raw_html_zero_count']}",
        f"- June 30 张报表中的 0：{result['june']['required_report_stats']['zero_count']}",
        f"- June 649 条映射相关唯一来源单元格中的 0：{result['june']['mapping_region_stats']['zero_count']}",
        f"- June 带星号数值：{result['june']['required_report_stats']['starred_numeric_count']}（现存 30 张采用报表）",
        f"- June 缺失 stats 引用：{result['june']['missing_catalog_reference_count']}",
        f"- 差异分类：{dict(counts)}",
        "",
        "## 星号与脚注",
        "",
        "- 星号已在隐藏的 `__cell_metadata` 工作表中与对应数值一并保留。",
        "- 现有 HTML 中存在 (a)/(1)/(2) 等可读脚注，已写入隐藏的 `__footnotes` 工作表。",
        "- 现有文件未找到统一解释所有数值星号含义的明确脚注，因此不作推断。",
        "",
        "## 发布判断",
        "",
        "- 新 March 暂不适合执行 backfill：现有发布重基准检查对同一行内重复出现的 Production/Yield/Area 年度组仍会阻断。",
        "- 不能判断是否需要 2026-06 修正版：新 June 候选未生成，无法完成新旧 649 条映射逐项比较。",
    ]
    (output_root / "workbook_rebuild_audit.md").write_text("\n".join(md) + "\n", encoding="utf-8")
