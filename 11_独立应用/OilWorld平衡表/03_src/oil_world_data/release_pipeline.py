from __future__ import annotations

import json
import os
import re
import shutil
import tempfile
import time
from collections.abc import Callable
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .generator import BuildError, build_release
from .paths import raw_source_reference, resolve_raw_data_root


RELEASE_PATTERN = re.compile(r"^20\d{2}-(0[1-9]|1[0-2])$")
SOURCE_RANGE_PATTERN = re.compile(
    r"(?P<sheet>[A-Za-z0-9_]+)!(?P<start_col>[A-Z]+)(?P<start_row>\d+):(?P<end_col>[A-Z]+)(?P<end_row>\d+)"
)
ELIGIBLE_MAPPING_STATUSES = {"direct", "derived"}
MONTH_LABELS = {
    1: "January",
    2: "February",
    3: "March",
    4: "April",
    5: "May",
    6: "June",
    7: "July",
    8: "August",
    9: "September",
    10: "October",
    11: "November",
    12: "December",
}


class ReleasePipelineError(BuildError):
    """A release was rejected before any official pointer was changed."""


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(temporary, path)


def _sha256(path: Path) -> str:
    import hashlib

    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _release_key(release: str) -> tuple[int, int]:
    if not RELEASE_PATTERN.fullmatch(release):
        raise ReleasePipelineError(f"发布期必须使用 YYYY-MM：{release}")
    year, month = release.split("-")
    return int(year), int(month)


def _release_label(release: str) -> str:
    year, month = _release_key(release)
    return f"{MONTH_LABELS[month]} {year}"


def find_source_workbook(project_root: Path, release: str) -> Path:
    source_directory = resolve_raw_data_root(project_root) / release
    if not source_directory.is_dir():
        raise ReleasePipelineError(f"原始发布期目录不存在：{source_directory}")
    workbooks = sorted(
        path for path in source_directory.iterdir() if path.is_file() and path.suffix.casefold() == ".xlsx"
    )
    if len(workbooks) != 1:
        raise ReleasePipelineError(
            f"{source_directory} 必须且只能包含一个主 Excel，当前发现 {len(workbooks)} 个"
        )
    year, month = _release_key(release)
    expected_tokens = (str(year), MONTH_LABELS[month].casefold())
    filename = workbooks[0].name.casefold()
    if not all(token in filename for token in expected_tokens):
        raise ReleasePipelineError(
            f"主 Excel 文件名与发布期不匹配：预期包含 {MONTH_LABELS[month]} {year}，实际为 {workbooks[0].name}"
        )
    return workbooks[0]


def _canonical_period(value: Any) -> str | None:
    text = str(value or "").strip()
    matches = list(re.finditer(r"(?<!\d)(20)?(\d{2})\s*/\s*(\d{2})(?!\d)", text))
    # A complete annual column has exactly one market-year token. Labels such as
    # ``20/21-24/25`` are multi-year averages and delimit adjacent annual blocks.
    if len(matches) != 1:
        return None
    match = matches[0]
    start = int(match.group(2))
    start_year = int(f"20{start:02d}")
    return f"{start_year}/{match.group(3)}"


def _basis(value: str) -> str:
    return re.sub(
        r"\s+",
        " ",
        re.sub(r"(?:20)?\d{2}\s*/\s*\d{2}F?", "", value, flags=re.I),
    ).strip(" -–/")


def _normalized_block_label(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def _block_label_for_columns(sheet: Any, header_row: int, start: int, end: int) -> str:
    """Return the unique section label immediately above an annual column group."""
    if header_row <= 1:
        return ""
    labels: list[str] = []
    normalized: set[str] = set()
    for column in range(start, end + 1):
        label = str(sheet.cell(header_row - 1, column).value or "").strip()
        key = _normalized_block_label(label)
        if label and key not in normalized:
            labels.append(label)
            normalized.add(key)
    return labels[0] if len(labels) == 1 else ""


def _candidate_description(
    sheet: Any,
    header_row: int,
    group: list[tuple[int, str]],
) -> str:
    start, end = group[0][0], group[-1][0]
    coordinates = f"{get_column_letter(start)}{header_row}:{get_column_letter(end)}{header_row}"
    block = _block_label_for_columns(sheet, header_row, start, end) or "<unlabelled>"
    years = " | ".join(item[1] for item in group)
    return f"{coordinates} [{block}] ({years})"


def _unconfigured_annual_group(
    sheet: Any,
    configured_header_row: int,
    old_start: int,
    old_end: int,
    context: str,
) -> tuple[int, list[int], list[str]]:
    """Locate an annual group when the baseline audit contains blank merged headers."""
    candidates: list[tuple[int, list[tuple[int, str]]]] = []
    first_row = max(1, configured_header_row - 2)
    last_row = min(sheet.max_row, configured_header_row + 5)
    for header_row in range(first_row, last_row + 1):
        annual_cells = [
            (column, str(sheet.cell(header_row, column).value or "").strip())
            for column in range(1, sheet.max_column + 1)
            if _canonical_period(sheet.cell(header_row, column).value) is not None
        ]
        groups: list[list[tuple[int, str]]] = []
        for annual_cell in annual_cells:
            basis = _basis(annual_cell[1]).casefold()
            previous_basis = _basis(groups[-1][-1][1]).casefold() if groups else ""
            if (
                not groups
                or annual_cell[0] != groups[-1][-1][0] + 1
                or basis != previous_basis
            ):
                groups.append([annual_cell])
            else:
                groups[-1].append(annual_cell)
        candidates.extend((header_row, group) for group in groups if len(group) >= 2)
    if not candidates:
        raise ReleasePipelineError(f"{context}: no complete annual group near header row")

    def score(candidate: tuple[int, list[tuple[int, str]]]) -> tuple[int, int, int]:
        header_row, group = candidate
        overlap = max(0, min(old_end, group[-1][0]) - max(old_start, group[0][0]) + 1)
        row_distance = abs(header_row - configured_header_row)
        column_distance = abs(group[0][0] - old_start)
        return overlap, -row_distance, -column_distance

    best_score = max(score(candidate) for candidate in candidates)
    best = [candidate for candidate in candidates if score(candidate) == best_score]
    if len(best) != 1:
        descriptions = "; ".join(
            _candidate_description(sheet, row, group) for row, group in best
        )
        raise ReleasePipelineError(
            f"{context}: annual column group is ambiguous; candidates: {descriptions}"
        )
    header_row, group = best[0]
    return header_row, [item[0] for item in group], [item[1] for item in group]


def _annual_column_run(
    sheet: Any,
    header_row: int,
    old_start: int,
    old_end: int,
    expected_labels: list[str],
    *,
    report_id: str = "",
    metric: str = "",
    region: str = "",
    unit: str = "",
    expected_block: str = "",
    preserve_expected_labels: bool = False,
) -> tuple[list[int], list[str]]:
    expected_basis = _basis(expected_labels[0]) if expected_labels else ""
    candidates: list[tuple[int, str]] = []
    for column in range(1, sheet.max_column + 1):
        raw = sheet.cell(header_row, column).value
        label = str(raw or "").strip()
        if _canonical_period(label) is None:
            continue
        label_basis = _basis(label)
        if expected_basis and label_basis and label_basis.casefold() != expected_basis.casefold():
            continue
        candidates.append((column, label))

    groups: list[list[tuple[int, str]]] = []
    for candidate in candidates:
        previous_basis = _basis(groups[-1][-1][1]).casefold() if groups else ""
        candidate_basis = _basis(candidate[1]).casefold()
        if (
            not groups
            or candidate[0] != groups[-1][-1][0] + 1
            or candidate_basis != previous_basis
        ):
            groups.append([candidate])
        else:
            groups[-1].append(candidate)
    groups = [group for group in groups if len(group) >= 2]
    target_width = len(expected_labels)
    if target_width:
        windowed_groups: list[list[tuple[int, str]]] = []
        for group in groups:
            if len(group) > target_width:
                windowed_groups.extend(
                    group[offset : offset + target_width]
                    for offset in range(len(group) - target_width + 1)
                )
            else:
                windowed_groups.append(group)
        groups = windowed_groups
    if not groups:
        raise ReleasePipelineError(
            f"{sheet.title} 第 {header_row} 行无法识别完整年度列，拒绝继续发布"
        )

    source_block = expected_block or _block_label_for_columns(
        sheet, header_row, old_start, old_end
    )
    normalized_source_block = _normalized_block_label(source_block)

    def score(group: list[tuple[int, str]]) -> tuple[int, int, int, int]:
        group_block = _block_label_for_columns(
            sheet, header_row, group[0][0], group[-1][0]
        )
        block_match = int(
            bool(normalized_source_block)
            and _normalized_block_label(group_block) == normalized_source_block
        )
        overlap = max(0, min(old_end, group[-1][0]) - max(old_start, group[0][0]) + 1)
        distance = abs(group[0][0] - old_start)
        width_delta = abs(len(group) - len(expected_labels))
        return block_match, overlap, -distance, -width_delta

    best_score = max(score(group) for group in groups)
    best_groups = [group for group in groups if score(group) == best_score]
    if len(best_groups) != 1:
        context = " / ".join(
            item for item in (report_id or sheet.title, metric, region, unit) if item
        )
        candidates_text = "; ".join(
            _candidate_description(sheet, header_row, group) for group in best_groups
        )
        raise ReleasePipelineError(
            f"{context}: annual column group is ambiguous; candidates: {candidates_text}"
        )
    selected = best_groups[0]
    columns = [item[0] for item in selected]
    labels = [item[1] for item in selected]
    canonical = [_canonical_period(label) for label in labels]
    if len(set(canonical)) != len(canonical):
        raise ReleasePipelineError(f"{sheet.title} 第 {header_row} 行存在重复市场年度")
    starts = [int(period.split("/")[0]) for period in canonical if period]
    if starts != sorted(starts, reverse=True):
        raise ReleasePipelineError(f"{sheet.title} 第 {header_row} 行市场年度顺序无法确认")
    if abs(len(columns) - len(expected_labels)) > 1:
        raise ReleasePipelineError(
            f"{sheet.title} 完整年度列数量变化超过安全范围：{len(expected_labels)} -> {len(columns)}"
        )
    return columns, list(expected_labels) if preserve_expected_labels else labels


def _rebase_audit_for_workbook(audit: dict[str, Any], workbook_path: Path) -> dict[str, Any]:
    """Re-anchor audited annual columns while keeping every audited data row unchanged."""
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    cache: dict[
        tuple[str, int, int, int, tuple[str, ...], str, str, str, str, bool],
        tuple[list[int], list[str]],
    ] = {}
    try:
        for row in audit["coverage_matrix"]:
            if row.get("mapping_status") not in ELIGIBLE_MAPPING_STATUSES:
                continue
            expression = str(row.get("source_cell_or_range") or "")
            matches = list(SOURCE_RANGE_PATTERN.finditer(expression))
            if not matches:
                raise ReleasePipelineError(
                    f"{row.get('commodity')} / {row.get('country_or_region')} / {row.get('metric')} 缺少可追溯来源单元格"
                )
            expected_labels = [part.strip() for part in str(row.get("market_years") or "").split("|") if part.strip()]
            header_row = int(row.get("header_row") or 0)
            if header_row < 1:
                raise ReleasePipelineError(f"{row.get('primary_report_id')} 缺少有效表头行")
            replacement_by_span: dict[tuple[int, int], str] = {}
            canonical_sequences: list[list[str | None]] = []
            displayed_labels: list[str] | None = None
            for match in matches:
                sheet_name = match.group("sheet")
                if sheet_name not in workbook.sheetnames:
                    raise ReleasePipelineError(f"必需工作表缺失：{sheet_name}")
                old_start = column_index_from_string(match.group("start_col"))
                old_end = column_index_from_string(match.group("end_col"))
                report_id = str(row.get("primary_report_id") or sheet_name)
                metric = str(row.get("metric") or "")
                region = str(row.get("country_or_region") or "")
                unit = str(row.get("original_unit") or row.get("standard_unit") or "")
                derivation_text = " ".join(
                    str(row.get(field) or "")
                    for field in ("original_metric", "derivation_formula", "derivation_components")
                ).casefold()
                preserve_expected_labels = (
                    "prior-year" in derivation_text or "t-1" in derivation_text
                )
                expected_block = _block_label_for_columns(
                    workbook[sheet_name], header_row, old_start, old_end
                )
                key = (
                    sheet_name,
                    header_row,
                    old_start,
                    old_end,
                    tuple(expected_labels),
                    report_id,
                    metric,
                    region,
                    unit,
                    preserve_expected_labels,
                )
                if key not in cache:
                    if expected_labels:
                        cache[key] = _annual_column_run(
                            workbook[sheet_name],
                            header_row,
                            old_start,
                            old_end,
                            expected_labels,
                            report_id=report_id,
                            metric=metric,
                            region=region,
                            unit=unit,
                            expected_block=expected_block,
                            preserve_expected_labels=preserve_expected_labels,
                        )
                    else:
                        context = " / ".join(
                            item for item in (report_id, metric, region, unit) if item
                        )
                        _, columns, labels = _unconfigured_annual_group(
                            workbook[sheet_name],
                            header_row,
                            old_start,
                            old_end,
                            context,
                        )
                        cache[key] = columns, labels
                columns, labels = cache[key]
                displayed_labels = displayed_labels or labels
                canonical_sequences.append([_canonical_period(label) for label in labels])
                replacement_by_span[match.span()] = (
                    f"{sheet_name}!{get_column_letter(columns[0])}{match.group('start_row')}:"
                    f"{get_column_letter(columns[-1])}{match.group('end_row')}"
                )
            if any(sequence != canonical_sequences[0] for sequence in canonical_sequences[1:]):
                raise ReleasePipelineError(f"{row.get('primary_report_id')} 派生组成项的市场年度无法安全对齐")
            parts: list[str] = []
            position = 0
            for span, replacement in sorted(replacement_by_span.items()):
                parts.append(expression[position : span[0]])
                parts.append(replacement)
                position = span[1]
            parts.append(expression[position:])
            row["source_cell_or_range"] = "".join(parts)
            row["market_years"] = " | ".join(displayed_labels or expected_labels)
    finally:
        workbook.close()
    return audit


def build_snapshot_from_audit(
    project_root: Path,
    release: str,
    workbook_path: Path,
    staging_root: Path,
) -> tuple[Path, dict[str, Any]]:
    """Build one release in an isolated project, never in the official release tree."""
    sandbox = staging_root / "builder"
    baseline_config_path = project_root / "02_configs" / "release_2026-06.json"
    baseline_config = _read_json(baseline_config_path)
    audit_source = project_root / baseline_config["audit_directory"]
    audit_destination = sandbox / baseline_config["audit_directory"]
    shutil.copytree(audit_source, audit_destination)
    source_destination = sandbox / "source" / workbook_path.name
    source_destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(workbook_path, source_destination)

    source_hash = _sha256(workbook_path)
    audit_path = audit_destination / baseline_config["audit_files"]["mapping_audit"]
    audit = _read_json(audit_path)
    if source_hash != baseline_config["source_workbook_sha256"]:
        audit = _rebase_audit_for_workbook(audit, source_destination)
    audit["source_workbook_sha256"] = source_hash
    _write_json(audit_path, audit)

    config = dict(baseline_config)
    config.update(
        {
            "release": release,
            "release_label": _release_label(release),
            "source_workbook": f"source/{workbook_path.name}",
            "source_workbook_sha256": source_hash,
            "expected_annual_change_count": None,
            "spot_check_seed": int(release.replace("-", "")),
        }
    )
    config_directory = sandbox / "02_configs"
    config_directory.mkdir(parents=True, exist_ok=True)
    _write_json(config_directory / "release_2026-06.json", config)
    result = build_release(sandbox)
    snapshot = sandbox / "01_data" / "releases" / release
    return snapshot, result


def _release_files(index: dict[str, Any]) -> dict[tuple[str, str, str], str]:
    return {
        (item["system"], item["product"], item["region"]): item["path"]
        for item in index.get("files", [])
    }


def _metric_map(payload: dict[str, Any]) -> dict[tuple[str, str, str], dict[str, Any]]:
    return {
        (metric["metric"], metric.get("period_family", "marketing_year"), metric.get("source_role", "balance")): metric
        for metric in payload.get("metrics", [])
    }


def _report_id(metric: dict[str, Any] | None) -> list[str]:
    return list((metric or {}).get("source_report_id") or [])


def _comparison_record(
    previous_release: str,
    current_release: str,
    identity: tuple[str, str, str],
    metric_name: str,
    period: str,
    previous_metric: dict[str, Any] | None,
    current_metric: dict[str, Any] | None,
    period_family: str = "marketing_year",
    source_role: str = "balance",
) -> dict[str, Any]:
    previous_status = (previous_metric or {}).get("mapping_status")
    current_status = (current_metric or {}).get("mapping_status")
    previous_unit = (previous_metric or {}).get("unit")
    current_unit = (current_metric or {}).get("unit")
    unit = current_unit or previous_unit or ""
    previous_value = (previous_metric or {}).get("values", {}).get(period)
    current_value = (current_metric or {}).get("values", {}).get(period)
    revision: float | None = None
    note = ""
    if not previous_metric or not current_metric:
        comparison_status = "metric_missing"
        note = "相邻发布期中有一期缺少该指标。"
    elif previous_unit != current_unit:
        comparison_status = "unit_mismatch"
        note = "相邻发布期单位不一致，未计算季度修正。"
    elif previous_status != current_status:
        comparison_status = "mapping_status_changed"
        note = "相邻发布期映射状态不一致，未计算季度修正。"
    elif current_status not in ELIGIBLE_MAPPING_STATUSES:
        comparison_status = "mapping_status_ineligible"
        note = f"{current_status} 状态不计算季度修正。"
    elif previous_value is None or current_value is None:
        comparison_status = "value_missing"
        note = "任意一期为空时季度修正保持为空。"
    else:
        comparison_status = "calculated"
        revision = round(float(current_value) - float(previous_value), 8)
    revision_unit = "percentage points" if metric_name == "Stocks/Use Ratio" else unit
    system, product, region = identity
    return {
        "previous_release": previous_release,
        "current_release": current_release,
        "system": system,
        "product": product,
        "region": region,
        "metric": metric_name,
        "period": period,
        "period_family": period_family,
        "source_role": source_role,
        "unit": revision_unit,
        "previous_value": previous_value,
        "current_value": current_value,
        "quarter_revision": revision,
        "previous_mapping_status": previous_status,
        "current_mapping_status": current_status,
        "previous_report_id": _report_id(previous_metric),
        "current_report_id": _report_id(current_metric),
        "comparison_status": comparison_status,
        "quality_note": note,
    }


def build_comparison(
    previous_release: str,
    current_release: str,
    previous_root: Path,
    current_root: Path,
    destination: Path,
) -> dict[str, Any]:
    previous_index = _read_json(previous_root / "index.json")
    current_index = _read_json(current_root / "index.json")
    previous_files = _release_files(previous_index)
    current_files = _release_files(current_index)
    if set(previous_files) != set(current_files):
        raise ReleasePipelineError("相邻发布期的商品—地区固定范围发生变化")
    files: list[dict[str, Any]] = []
    total = calculated = 0
    for identity in sorted(current_files):
        previous_payload = _read_json(previous_root / previous_files[identity])
        current_payload = _read_json(current_root / current_files[identity])
        previous_metrics = _metric_map(previous_payload)
        current_metrics = _metric_map(current_payload)
        metric_order = list(dict.fromkeys([*current_metrics, *previous_metrics]))
        records: list[dict[str, Any]] = []
        for metric_name, period_family, source_role in metric_order:
            metric_key = (metric_name, period_family, source_role)
            previous_metric = previous_metrics.get(metric_key)
            current_metric = current_metrics.get(metric_key)
            periods = sorted(
                set((previous_metric or {}).get("periods", [])) | set((current_metric or {}).get("periods", [])),
                reverse=True,
            )
            for period in periods:
                record = _comparison_record(
                    previous_release,
                    current_release,
                    identity,
                    metric_name,
                    period,
                    previous_metric,
                    current_metric,
                    period_family,
                    source_role,
                )
                records.append(record)
                total += 1
                calculated += record["quarter_revision"] is not None
        relative_path = current_files[identity]
        _write_json(
            destination / relative_path,
            {
                "schema_version": 1,
                "previous_release": previous_release,
                "current_release": current_release,
                "system": identity[0],
                "product": identity[1],
                "region": identity[2],
                "records": records,
            },
        )
        files.append(
            {"system": identity[0], "product": identity[1], "region": identity[2], "path": relative_path}
        )
    index = {
        "schema_version": 1,
        "previous_release": previous_release,
        "current_release": current_release,
        "files": files,
        "combination_count": len(files),
        "comparison_record_count": total,
        "calculated_revision_count": calculated,
        "null_revision_count": total - calculated,
    }
    _write_json(destination / "index.json", index)
    return index


def _legacy_release_entry(project_root: Path, release: str) -> dict[str, Any]:
    release_root = project_root / "public" / "data" / "oil_world" / "releases" / release
    manifest = _read_json(release_root / "manifest.json")
    index = _read_json(release_root / "index.json")
    return {
        "release": release,
        "label": index.get("release_label", _release_label(release)),
        "source_file": manifest.get("source_workbook", ""),
        "source_file_hash": manifest.get("source_workbook_sha256", ""),
        "release_type": "baseline" if release == "2026-06" else "regular_update",
        "import_time": manifest.get("generated_at", ""),
        "previous_release": None,
        "next_release": None,
        "quality_status": "passed",
        "mapping_record_count": index["mapping_record_count"],
        "numeric_value_count": index["numeric_observation_count"],
        "annual_change_count": index.get("annual_change_count", 0),
        "available": True,
    }


def _load_release_entries(project_root: Path) -> list[dict[str, Any]]:
    path = project_root / "public" / "data" / "oil_world" / "releases.json"
    raw = _read_json(path)
    entries = []
    for item in raw.get("releases", []):
        if "source_file_hash" in item:
            entries.append(dict(item))
        else:
            entries.append(_legacy_release_entry(project_root, item["release"]))
    return sorted(entries, key=lambda item: _release_key(item["release"]))


def _snapshot_root(project_root: Path, release: str, staged_release: str, staged_root: Path) -> Path:
    if release == staged_release:
        return staged_root
    return project_root / "01_data" / "releases" / release


def _validate_snapshot(snapshot: Path, release: str) -> dict[str, Any]:
    index = _read_json(snapshot / "index.json")
    manifest = _read_json(snapshot / "manifest.json")
    quality = _read_json(snapshot / "quality_report.json")
    if index.get("release") != release or manifest.get("release") != release:
        raise ReleasePipelineError("暂存快照发布期标识不一致")
    if not quality.get("passed"):
        raise ReleasePipelineError("数据质量检查未通过")
    stable_keys: set[tuple[str, str, str, str, str, str, str, str]] = set()
    for item in index.get("files", []):
        payload = _read_json(snapshot / item["path"])
        for metric in payload.get("metrics", []):
            for period, value in metric.get("values", {}).items():
                if value is not None and not isinstance(value, (int, float)):
                    raise ReleasePipelineError("数据值必须是数字或空值")
                key = (
                    payload["system"], payload["product"], payload["region"],
                    metric["metric"], period, metric["unit"],
                    metric.get("period_family", "marketing_year"), metric.get("source_role", "balance"),
                )
                if key in stable_keys:
                    raise ReleasePipelineError(f"重复稳定键：{key}")
                stable_keys.add(key)
                if value is not None and not metric.get("source_cells", {}).get(period):
                    raise ReleasePipelineError(f"来源单元格不可追溯：{key}")
    return {"index": index, "manifest": manifest, "quality": quality}


def _copy_new_directory(source: Path, destination: Path) -> None:
    if destination.exists():
        raise ReleasePipelineError(f"不可变目录已存在，拒绝覆盖：{destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.parent / f".{destination.name}.installing"
    if temporary.exists():
        shutil.rmtree(temporary)
    shutil.copytree(source, temporary)
    _atomic_replace(temporary, destination)


def _atomic_replace(source: Path, destination: Path, attempts: int = 5) -> None:
    """Retry transient Windows/OneDrive directory locks without masking persistent failures."""
    for attempt in range(attempts):
        try:
            os.replace(source, destination)
            return
        except PermissionError:
            if attempt + 1 == attempts:
                raise
            time.sleep(0.1 * (2**attempt))


def _replace_comparison_root(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.parent / f".{destination.name}.installing"
    backup = destination.parent / f".{destination.name}.backup"
    for path in (temporary, backup):
        if path.exists():
            shutil.rmtree(path)
    shutil.copytree(source, temporary)
    if destination.exists():
        _atomic_replace(destination, backup)
    try:
        _atomic_replace(temporary, destination)
    except Exception:
        if backup.exists():
            _atomic_replace(backup, destination)
        raise
    if backup.exists():
        shutil.rmtree(backup)


SnapshotBuilder = Callable[[Path, str, Path, Path], tuple[Path, dict[str, Any]]]


def update_release(
    project_root: Path,
    release: str,
    *,
    backfill: bool = False,
    validate_only: bool = False,
    snapshot_builder: SnapshotBuilder = build_snapshot_from_audit,
    import_time: str | None = None,
) -> dict[str, Any]:
    project_root = project_root.resolve()
    _release_key(release)
    workbook_path = find_source_workbook(project_root, release)
    source_hash_before = _sha256(workbook_path)
    entries = _load_release_entries(project_root)
    existing_releases = [item["release"] for item in entries]
    latest_before = max(existing_releases, key=_release_key)
    if release in existing_releases:
        raise ReleasePipelineError(f"发布期已存在，默认拒绝覆盖：{release}")
    if backfill and _release_key(release) > _release_key(latest_before):
        raise ReleasePipelineError("backfill 仅用于早于或等于当前 latest 的历史发布期")
    if not backfill and _release_key(release) <= _release_key(latest_before):
        raise ReleasePipelineError("普通发布必须晚于当前 latest；历史版本请使用 --backfill")

    with tempfile.TemporaryDirectory(prefix=f"oil_world_{release}_") as temporary_directory:
        staging = Path(temporary_directory)
        staged_snapshot, build_result = snapshot_builder(project_root, release, workbook_path, staging)
        validated = _validate_snapshot(staged_snapshot, release)
        source_hash_after = _sha256(workbook_path)
        if source_hash_after != source_hash_before:
            raise ReleasePipelineError("原始 Excel 在只读构建期间发生变化")
        index = validated["index"]
        timestamp = import_time or datetime.now(timezone.utc).isoformat()
        new_entry = {
            "release": release,
            "label": index.get("release_label", _release_label(release)),
            "source_file": raw_source_reference(project_root, workbook_path),
            "source_file_hash": source_hash_after,
            "release_type": "backfill" if backfill else "regular_update",
            "import_time": timestamp,
            "previous_release": None,
            "next_release": None,
            "quality_status": "passed",
            "mapping_record_count": index["mapping_record_count"],
            "numeric_value_count": index["numeric_observation_count"],
            "annual_change_count": index.get("annual_change_count", 0),
            "available": True,
        }
        all_entries = sorted([*entries, new_entry], key=lambda item: _release_key(item["release"]))
        for position, item in enumerate(all_entries):
            item["previous_release"] = all_entries[position - 1]["release"] if position else None
            item["next_release"] = all_entries[position + 1]["release"] if position + 1 < len(all_entries) else None
        latest_after = all_entries[-1]["release"]

        comparison_staging = staging / "comparisons"
        comparison_summaries: list[dict[str, Any]] = []
        for previous, current in zip(all_entries, all_entries[1:], strict=False):
            pair = f"{previous['release']}_to_{current['release']}"
            summary = build_comparison(
                previous["release"],
                current["release"],
                _snapshot_root(project_root, previous["release"], release, staged_snapshot),
                _snapshot_root(project_root, current["release"], release, staged_snapshot),
                comparison_staging / pair,
            )
            comparison_summaries.append(summary)

        report = {
            "release": release,
            "mode": "backfill" if backfill else "regular_update",
            "validate_only": validate_only,
            "source_file": str(workbook_path),
            "source_file_hash": source_hash_after,
            "latest_before": latest_before,
            "latest_after": latest_after,
            "previous_release": new_entry["previous_release"],
            "next_release": new_entry["next_release"],
            "release_count": len(all_entries),
            "comparison_pairs": [
                f"{item['previous_release']}_to_{item['current_release']}" for item in comparison_summaries
            ],
            "mapping_record_count": index["mapping_record_count"],
            "numeric_value_count": index["numeric_observation_count"],
            "annual_change_count": index.get("annual_change_count", 0),
            "quality_status": "passed",
            "build_result": build_result,
        }
        manifest = validated["manifest"]
        manifest.update(
            {
                "release_type": new_entry["release_type"],
                "source_workbook": new_entry["source_file"],
                "source_workbook_sha256": source_hash_after,
                "import_time": timestamp,
                "previous_release": new_entry["previous_release"],
                "next_release": new_entry["next_release"],
                "quality_status": "passed",
            }
        )
        _write_json(staged_snapshot / "manifest.json", manifest)
        _write_json(staged_snapshot / "update_report.json", report)
        _write_json(staging / "update_report.json", report)
        if validate_only:
            return report

        internal_release = project_root / "01_data" / "releases" / release
        public_release = project_root / "public" / "data" / "oil_world" / "releases" / release
        if internal_release.exists() or public_release.exists():
            raise ReleasePipelineError(f"发布目录已存在，拒绝覆盖：{release}")
        _copy_new_directory(staged_snapshot, internal_release)
        try:
            _copy_new_directory(staged_snapshot, public_release)
            _replace_comparison_root(comparison_staging, project_root / "01_data" / "comparisons")
            _replace_comparison_root(
                comparison_staging, project_root / "public" / "data" / "oil_world" / "comparisons"
            )
            _write_json(
                project_root / "public" / "data" / "oil_world" / "releases.json",
                {"schema_version": 2, "releases": all_entries},
            )
            _write_json(
                project_root / "public" / "data" / "oil_world" / "latest.json",
                {"release": latest_after},
            )
        except Exception:
            for created in (public_release, internal_release):
                if created.exists():
                    shutil.rmtree(created)
            raise
        return report
