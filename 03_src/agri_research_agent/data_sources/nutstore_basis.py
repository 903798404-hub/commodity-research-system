"""Read the explicitly selected Nutstore workbook without writing its tree."""
from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import statistics

from openpyxl import load_workbook

PROTECTED_ROOT = Path(r"D:\坚果云\123")
SOURCE_PATH = PROTECTED_ROOT / "国内基差及一口价" / "基差数据.xlsx"
SOURCE_SHEET = "油脂油料价格 basis_price"
COLUMNS = ("日期", "品种", "地区", "省份", "工厂", "合同情况", "基差合同年",
           "基差合同月", "期货合约", "基差", "现货价", "成交量", "报价类别", "期货收盘价")
PRODUCTS = {
    "豆粕": ("豆粕", "soybean_meal", "DCE", "M"),
    "菜粕": ("菜粕", "rapeseed_meal", "CZCE", "RM"),
    "大豆油": ("一豆", "soybean_oil", "DCE", "Y"),
    "菜籽油": ("三菜", "rapeseed_oil_3", "CZCE", "OI"),
    "棕榈油": ("24度", "palm_oil_24", "DCE", "P"),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _unlinked(path: Path) -> Path:
    for node in (path, *path.parents):
        if node.is_symlink() or node.is_junction():
            raise ValueError("Nutstore paths must not traverse links or junctions")
    return path.resolve()


def assert_external_output(path: str | Path) -> Path:
    candidate = Path(path).absolute()
    resolved = _unlinked(candidate)
    protected = PROTECTED_ROOT.absolute()
    if candidate.is_relative_to(protected) or resolved.is_relative_to(protected.resolve()):
        raise ValueError("The entire Nutstore 123 tree is forbidden for writes")
    return resolved


def _number(value: object, row_number: int) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        result = Decimal(str(value))
    except InvalidOperation:
        raise ValueError(f"Non-numeric raw quote at row {row_number}") from None
    if not result.is_finite():
        raise ValueError(f"Non-finite raw quote at row {row_number}")
    return result


@dataclass(frozen=True)
class NutstoreBasisSnapshot:
    rows: tuple[dict, ...]
    report: dict


def read_nutstore_basis(
    source: str | Path = SOURCE_PATH, *, after: date, today: date | None = None,
) -> NutstoreBasisSnapshot:
    """Aggregate new dates only; original rows remain identified by file SHA."""
    path = Path(source).absolute()
    if path != SOURCE_PATH.absolute() or _unlinked(path) != SOURCE_PATH.resolve():
        raise ValueError("Only the explicitly allowed Nutstore workbook may be read")
    before_sha = sha256(path)
    groups = defaultdict(list)
    counters = Counter()
    source_max = None
    today = today or date.today()
    # Read-only mode, no save, no recalculation and no external link refresh.
    with path.open("rb") as handle:
        workbook = load_workbook(handle, read_only=True, data_only=False, keep_links=False)
        try:
            if SOURCE_SHEET not in workbook.sheetnames:
                raise ValueError("Nutstore raw source sheet is missing")
            stream = workbook[SOURCE_SHEET].iter_rows(values_only=True)
            if tuple(next(stream, ())) != COLUMNS:
                raise ValueError("Nutstore raw source columns differ from the audited contract")
            for row_number, values in enumerate(stream, 2):
                if all(value is None for value in values):
                    continue
                counters["raw_rows"] += 1
                if len(values) != len(COLUMNS) or any(
                    isinstance(value, str) and value.startswith("=") for value in values
                ):
                    raise ValueError(f"Formula or schema mismatch in raw row {row_number}")
                raw = dict(zip(COLUMNS, values, strict=True))
                business_date = raw["日期"]
                if isinstance(business_date, datetime):
                    business_date = business_date.date()
                if type(business_date) is not date or business_date > today:
                    raise ValueError(f"Invalid business date at raw row {row_number}")
                source_max = max(source_max or business_date, business_date)
                if business_date <= after:
                    counters["preserved_history_raw_rows"] += 1
                    continue
                if raw["品种"] not in PRODUCTS or not str(raw["地区"] or "").strip():
                    raise ValueError(f"Unknown product or empty region at row {row_number}")
                quote_type = raw["报价类别"]
                if quote_type not in {"现货基差", "一口价"}:
                    counters["excluded_quote_rows"] += 1
                    continue
                field = "基差" if quote_type == "现货基差" else "现货价"
                value = _number(raw[field], row_number)
                if value is None:
                    counters["missing_value_rows"] += 1
                    continue
                counters["zero_value_rows"] += int(value == 0)
                contract = None
                if quote_type == "现货基差":
                    code = str(raw["期货合约"] or "").strip()
                    # Excel may store native YYMM as a number.
                    if code.endswith(".0"):
                        code = code[:-2]
                    if not re.fullmatch(r"[0-9]{2}(0[1-9]|1[0-2])", code):
                        counters["invalid_contract_rows"] += 1
                        continue
                    contract = (2000 + int(code[:2]), int(code[2:]))
                    if contract < (business_date.year, business_date.month):
                        counters["expired_contract_rows"] += 1
                        continue
                key = (business_date, raw["品种"], str(raw["地区"]).strip(), quote_type)
                groups[key].append((contract, value, row_number))
        finally:
            workbook.close()
    if sha256(path) != before_sha:
        raise ValueError("Nutstore source changed during read; candidate rejected")
    output = []
    for (business_date, product, region, quote_type), observations in sorted(groups.items()):
        nearest = min(item[0] for item in observations) if quote_type == "现货基差" else None
        selected = [item for item in observations if item[0] == nearest]
        commodity, product_id, exchange, instrument = PRODUCTS[product]
        region_id = hashlib.sha256(region.encode("utf-8")).hexdigest()[:12]
        quote_id = "spot" if quote_type == "现货基差" else "cash"
        value = statistics.median(item[1] for item in selected).quantize(Decimal("0.0001"))
        group_sha = hashlib.sha256(json.dumps(
            [(item[0], str(item[1]), item[2]) for item in selected],
            ensure_ascii=False, sort_keys=True,
        ).encode("utf-8")).hexdigest()
        contract_code = None if nearest is None else f"{nearest[0] % 100:02d}{nearest[1]:02d}"
        series_id = f"market.basis.domestic.china.nutstore.{product_id}.{region_id}.{quote_id}"
        output.append({
            "schema_version": "public-domestic-basis-current/4", "segment": "LIVE_NUTSTORE",
            "series_id": series_id, "provider_dataset_id": "nutstore:basis_price",
            "provider_series_id": f"nutstore:{product}:{region}:{quote_type}",
            "source_series_id": f"nutstore:{product}:{region}:{quote_type}", "provider": "Nutstore",
            "business_date": business_date, "date": business_date, "commodity": commodity,
            "region": region, "product": product_id, "consumer_product": commodity,
            "location": region, "region_id": region_id,
            "quote_type": "基差报价" if nearest is not None else "一口价",
            "canonical_quote_type": quote_type, "delivery_month": "现货",
            "futures_contract": contract_code, "value": value,
            "cash_price": value if nearest is None else None, "futures_price": None,
            "basis": value if nearest is not None else None,
            "currency": "CNY", "unit": "CNY/metric_tonne",
            "underlying_futures_reference": None if nearest is None else f"{exchange}:{instrument}{nearest[0]}{nearest[1]:02d}",
            "source_sheet": SOURCE_SHEET,
            "source_locator": f"nutstore-file-sha256:{before_sha}/sheet:{SOURCE_SHEET}",
            "source_row_count": len(selected), "far_contract_row_count": len(observations) - len(selected),
            "source_group_sha256": group_sha, "aggregation_method": "nearest_nonexpired_contract_median" if nearest else "cash_price_median",
            "query_identity": "nutstore-read-only-raw-sheet/1", "snapshot_identity": before_sha,
            "captured_at": None, "mapping_version": "nutstore-basis/1",
            "evidence_type": "LOCAL_SYNCED_WORKBOOK", "live_status": "SOURCE_FILE_VERIFIED",
        })
    return NutstoreBasisSnapshot(tuple(output), {
        "source_sha256": before_sha, "source_path": str(path), "source_sheet": SOURCE_SHEET,
        "source_max_date": source_max.isoformat() if source_max else None,
        "append_after": after.isoformat(), "candidate_rows": len(output),
        "counts": dict(counters), "source_write_policy": "entire_123_tree_forbidden",
        "date_fill_policy": "none", "missing_value_policy": "null_not_zero",
        "series_max_dates": {
            series: max(row["business_date"] for row in output if row["series_id"] == series).isoformat()
            for series in sorted({row["series_id"] for row in output})
        },
    })
