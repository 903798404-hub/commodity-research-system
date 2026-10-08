"""Normalize CGC export facts and compute exact-grain-week comparisons."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

SCHEMA = "canola-exports/1"
MODULE_ID = "canola-exports"
STABLE = Path("processed/canola_exports/weekly.json")
STATUS = Path("update_status/canola_exports.json")
BASE_URL = "https://www.grainscanada.gc.ca/en/grain-research/statistics/grain-statistics-weekly/"
PORTS = {"Vancouver", "Prince Rupert", "Churchill", "Thunder Bay", "Bay & Lakes", "St. Lawrence"}
PERIODS = {"Current Week": "weekly_mt", "Crop Year": "cumulative_mt"}
HEADER = {"Crop Year": "crop_year", "Grain Week": "grain_week",
          "Week Ending Date": "week_ending_date", "Region": "region"}


def digest(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def source_url(year: str) -> str:
    crop_year(year)
    return f"{BASE_URL}{year[:4]}-{year[-2:]}/gsw-shg-en.csv"


def crop_year(value: str) -> int:
    if not isinstance(value, str) or not re.fullmatch(r"20\d{2}-20\d{2}", value):
        raise ValueError("invalid crop year")
    start, end = map(int, value.split("-"))
    if end != start + 1:
        raise ValueError("invalid crop year span")
    return start


def tonnes(value: str) -> float | None:
    text = value.strip()
    if text in {"", "-", ".", "..", "...", "N/A"}:
        return None
    if not re.fullmatch(r"-?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?", text):
        raise ValueError(f"invalid Ktonnes: {text!r}")
    try:
        number = Decimal(text.replace(",", "")) * 1000
    except InvalidOperation as exc:
        raise ValueError("invalid Ktonnes") from exc
    if number < 0 or not number.is_finite():
        raise ValueError("invalid negative/nonfinite export")
    return float(number)


def parse_csv(payload: bytes, year: str) -> list[dict]:
    start = crop_year(year)
    reader = csv.DictReader(io.StringIO(payload.decode("utf-8-sig")))
    fields = [HEADER.get(x, x) for x in (reader.fieldnames or [])]
    required = {"crop_year", "grain_week", "week_ending_date", "worksheet", "metric",
                "period", "grain", "grade", "region", "Ktonnes"}
    if len(fields) != len(set(fields)) or set(fields) != required:
        raise ValueError("unexpected CGC CSV header")
    groups = defaultdict(dict)
    dates = {}
    for original in reader:
        row = {HEADER.get(k, k): v for k, v in original.items()}
        if None in row or any(v is None for v in row.values()):
            raise ValueError("malformed CSV row")
        if row["grain"] != "Canola":
            continue
        if row["crop_year"] != year:
            raise ValueError("source crop year mismatch")
        port = row["worksheet"] == "Terminal Disposition" and row["metric"] == "Export Destinations"
        direct = (row["worksheet"] == "Primary Shipment Distribution"
                  and row["metric"] == "Shipment Distribution" and row["region"] == "Export Destinations")
        if not (port or direct):
            continue
        if row["period"] not in PERIODS or row["grade"]:
            raise ValueError("unexpected export period/grade")
        location = row["region"] if port else "Direct"
        if port and location not in PORTS:
            raise ValueError("unexpected export region")
        week = int(row["grain_week"])
        ending = datetime.strptime(row["week_ending_date"], "%d/%m/%Y").date()
        if not 1 <= week <= 53 or not date(start, 8, 1) <= ending <= date(start + 1, 8, 7):
            raise ValueError("invalid grain week/date")
        if week in dates and dates[week] != ending.isoformat():
            raise ValueError("conflicting week ending dates")
        dates[week] = ending.isoformat()
        key = (week, PERIODS[row["period"]])
        if location in groups[key]:
            raise ValueError("duplicate export component")
        groups[key][location] = tonnes(row["Ktonnes"])
    records = []
    for week in sorted(dates):
        item = {"crop_year": year, "grain_week": week, "week_ending": dates[week], "source_sha256": digest(payload)}
        for metric in PERIODS.values():
            cells = groups[(week, metric)]
            missing = sorted((PORTS | {"Direct"}) - set(cells))
            item[metric] = None if missing or any(v is None for v in cells.values()) else round(sum(cells.values()), 3)
            item[f"{metric}_missing_components"] = missing
        records.append(item)
    if not records:
        raise ValueError("no canola exports in source")
    for previous, current in zip(records, records[1:]):
        if previous["week_ending"] >= current["week_ending"]:
            current["date_quality"] = "non_monotonic_source_date"
    return records


def report_values(payload: bytes, year: str, week: int) -> dict:
    """Read explicit Summary cells, never infer zero from absent CSV components."""
    with io.BytesIO(payload) as handle:
        workbook = openpyxl.load_workbook(handle, data_only=True, read_only=True)
        try:
            cover = {str(cell).strip() for row in workbook.worksheets[0].iter_rows(values_only=True) for cell in row if cell is not None}
            if year not in cover or f"Week {week}" not in cover:
                raise ValueError("CGC Excel crop year/week mismatch")
            sheet = workbook["Summary"]
            rows = list(sheet.iter_rows(values_only=True))
            columns = {i for row in rows for i, value in enumerate(row) if value == "Canola"}
            if len(columns) != 1 or not any("000" in str(cell) for row in rows for cell in row):
                raise ValueError("CGC Excel canola column/units mismatch")
            column = columns.pop()
            in_exports = False
            result = {}
            for index, row in enumerate(rows, 1):
                label = str(row[0] or "").strip()
                if label == "Exports":
                    in_exports = True
                elif in_exports and label in {"Current Week", "To Date"}:
                    value = row[column]
                    if type(value) not in {int, float} or not math.isfinite(value) or value < 0:
                        raise ValueError("CGC Excel has no explicit numeric export")
                    metric = "weekly_mt" if label == "Current Week" else "cumulative_mt"
                    result[metric] = {"value": round(value * 1000, 3),
                                      "cell": f"Summary!{openpyxl.utils.get_column_letter(column+1)}{index}"}
                elif in_exports and label == "Domestic Disappearance":
                    break
            if set(result) != set(PERIODS.values()):
                raise ValueError("CGC Excel export summary incomplete")
            return result
        finally:
            workbook.close()


def reconcile_report(records: list[dict], year: str, week: int, url: str, raw: bytes) -> None:
    prefix = source_url(year).rsplit("/", 1)[0] + "/"
    if not url.startswith(prefix) or not url.endswith(".xlsx") or "/../" in url:
        raise ValueError("nonofficial crop-year report URL")
    record = next(r for r in records if r["grain_week"] == week)
    values = report_values(raw, year, week)
    filled = {}
    for metric in PERIODS.values():
        if record[metric] is None:
            record[metric] = values[metric]["value"]
            filled[metric] = values[metric]
    if filled:
        record["report_evidence"] = {"source_url": url, "sha256": digest(raw), "cells": filled}


def validate_bundle(bundle: dict) -> dict:
    if bundle.get("schema_version") != SCHEMA or not bundle.get("records"):
        raise ValueError("invalid canola export bundle")
    datetime.fromisoformat(bundle["generated_at"])
    sources = bundle["sources"]
    seen = set()
    for item in bundle["records"]:
        year, week = item["crop_year"], item["grain_week"]
        start = crop_year(year)
        if type(week) is not int or not 1 <= week <= 53 or (year, week) in seen:
            raise ValueError("duplicate/invalid bundle week")
        seen.add((year, week))
        ending = date.fromisoformat(item["week_ending"])
        if not date(start, 8, 1) <= ending <= date(start + 1, 8, 7):
            raise ValueError("invalid bundle date")
        source = sources[year]
        if source["source_url"] != source_url(year) or source["sha256"] != item["source_sha256"]:
            raise ValueError("bundle source identity mismatch")
        if not re.fullmatch("[0-9a-f]{64}", source["sha256"]):
            raise ValueError("invalid source digest")
        if datetime.fromisoformat(source["retrieved_at"]).utcoffset() is None:
            raise ValueError("source retrieval time requires timezone")
        for metric in PERIODS.values():
            value = item[metric]
            if value is not None and (type(value) not in {int, float} or not math.isfinite(value) or value < 0):
                raise ValueError("invalid export value")
        evidence = item.get("report_evidence")
        if evidence is not None:
            prefix = source_url(year).rsplit("/", 1)[0] + "/"
            if (not evidence["source_url"].startswith(prefix) or not evidence["source_url"].endswith(".xlsx")
                    or not re.fullmatch("[0-9a-f]{64}", evidence["sha256"])):
                raise ValueError("invalid report evidence identity")
            for metric, cell in evidence["cells"].items():
                if (metric not in PERIODS.values() or cell["value"] != item[metric]
                        or not re.fullmatch(r"Summary![A-Z]+[1-9]\d*", cell["cell"])):
                    raise ValueError("invalid report evidence value/locator")
    for year in sources:
        track = sorted((r for r in bundle["records"] if r["crop_year"] == year), key=lambda r: r["grain_week"])
        if not track:
            raise ValueError("empty source track")
        for previous, current in zip(track, track[1:]):
            if previous["week_ending"] >= current["week_ending"] and current.get("date_quality") != "non_monotonic_source_date":
                raise ValueError("unmarked source date inconsistency")
    return bundle


def load_bundle(path: Path) -> dict:
    return validate_bundle(json.loads(path.read_text(encoding="utf-8")))


def percentage(current: float | None, previous: float | None) -> float | None:
    return None if current is None or previous in {None, 0} else (current / previous - 1) * 100


def page_payload(bundle: dict, selected_years: list[str] | None = None) -> dict:
    validate_bundle(bundle)
    years = sorted(bundle["sources"], reverse=True)
    current = years[0]
    selected = selected_years if selected_years is not None else years[:7]
    tracks = {}
    for year in years:
        by_week = {r["grain_week"]: r for r in bundle["records"] if r["crop_year"] == year}
        points = []
        for week in range(1, max(by_week) + 1):
            record = by_week.get(week, {"crop_year": year, "grain_week": week, "week_ending": None,
                                        "weekly_mt": None, "cumulative_mt": None})
            window = [by_week.get(w, {}).get("weekly_mt") for w in range(week - 3, week + 1)]
            rolling = sum(window) if week >= 4 and all(v is not None for v in window) else None
            points.append({**record, "four_week_mt": rolling})
        tracks[year] = points
    latest = tracks[current][-1]
    week = latest["grain_week"]
    start = crop_year(current) - 1
    previous = f"{start}-{start+1}"
    prior = next((p for p in tracks.get(previous, []) if p["grain_week"] == week), {})
    samples = [p["cumulative_mt"] for y in selected for p in tracks.get(y, [])
               if p["grain_week"] == week and p["cumulative_mt"] is not None]
    cumulative = latest["cumulative_mt"]
    return {"years": years, "current_year": current, "previous_year": previous,
            "tracks": {y: tracks[y] for y in selected if y in tracks}, "latest": latest,
            "previous": prior, "cumulative_yoy": percentage(cumulative, prior.get("cumulative_mt")),
            "four_week_yoy": percentage(latest["four_week_mt"], prior.get("four_week_mt")),
            "rank": (1 + sum(v > cumulative for v in samples)) if cumulative is not None and current in selected else None,
            "rank_samples": len(samples), "generated_at": bundle["generated_at"], "sources": bundle["sources"],
            "revision_count": len(bundle.get("revisions", []))}
