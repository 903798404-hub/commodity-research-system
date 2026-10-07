"""Manual CONAB archiving, candidate preparation and isolated local activation."""
from __future__ import annotations

import json
import shutil
import uuid
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import requests
from filelock import FileLock
from openpyxl import load_workbook

from agri_research_agent.shared.atomic_storage import atomic_write_bytes, atomic_write_json
from agri_research_agent.shared.runtime_context import RuntimeContext, RuntimeMode, assert_runtime_write
from .brazil_soy import (
    SCHEMA, STATES, STATE_NAMES, STABLE_RELATIVE_PATH, check_source_url, import_workbook, load_bundle,
    normalize_season, observation, sha256_file, strict_json, utc_now, validate_bundle,
)


def local_write(context: RuntimeContext, path: Path) -> Path:
    if context.mode not in {RuntimeMode.FIXTURE, RuntimeMode.ISOLATED_DEV} or context.module_id != "brazil-soy":
        raise ValueError("Brazil preparation only accepts its isolated local runtime")
    return assert_runtime_write(context, path)


def latest_dates(bundle: dict) -> dict:
    result = {}
    for row in bundle["records"]:
        key = f"{row['season']}/{row['region']}/{row['metric']}"
        result[key] = max(result.get(key, ""), row["date"])
    return result


def save_candidate(context: RuntimeContext, bundle: dict, baseline: str | None, changes: dict) -> Path:
    validate_bundle(bundle)
    directory = local_write(context, context.runtime_root / "candidates/brazil_soy" / uuid.uuid4().hex)
    directory.mkdir(parents=True, exist_ok=False)
    candidate = directory / "soy_weekly.json"
    atomic_write_json(local_write(context, candidate), bundle)
    atomic_write_json(local_write(context, directory / "candidate.json"), {
        "schema_version": "brazil-soy-candidate/1", "created_at": utc_now(),
        "baseline_sha256": baseline, "payload_sha256": sha256_file(candidate), "changes": changes,
        "latest_dates": latest_dates(bundle)})
    return candidate


def candidate_from_workbook(context: RuntimeContext, path: Path) -> Path:
    bundle = import_workbook(path)
    candidate = save_candidate(context, bundle, None, {"initial_import": len(bundle["records"])})
    snapshot = local_write(context, candidate.parent / "source.xlsm")
    shutil.copyfile(path, snapshot)
    if sha256_file(snapshot) != bundle["records"][0]["source_sha256"]:
        raise ValueError("workbook changed while copying")
    return candidate


def prepare_area_reference(context: RuntimeContext, baseline: Path, source: Path) -> Path:
    from .brazil_soy import parse_area_reference
    local_write(context, source)
    source_root = (context.runtime_root / "raw/brazil_soy").resolve()
    if source_root not in source.resolve().parents:
        raise ValueError("area source must be in this local report archive")
    identity = sha256_file(baseline)
    bundle = load_bundle(baseline)
    info = strict_json(source)
    report = source.parent / "report.bin"
    if sha256_file(report) != info["sha256"]:
        raise ValueError("archived area report identity mismatch")
    with report.open("rb") as handle:
        reference = parse_area_reference(handle, info)
    previous = bundle.get("area_reference")
    if previous and (reference["season"], reference["published_at"]) < (previous["season"], previous["published_at"]):
        raise ValueError("area reference cannot move backwards")
    if sha256_file(baseline) != identity:
        raise ValueError("baseline changed while preparing area reference")
    return save_candidate(context, {**bundle, "generated_at": utc_now(), "area_reference": reference},
                          identity, {"area_reference_changed": previous != reference})


def fetch_report(context: RuntimeContext, url: str, published_at: str) -> Path:
    """Archive one explicitly requested report; validate every redirect before requesting it."""
    from datetime import date
    if date.fromisoformat(published_at) > date.today():
        raise ValueError("publication date is in the future")
    check_source_url(url)
    directory = local_write(context, context.runtime_root / "raw/brazil_soy" / uuid.uuid4().hex)
    directory.mkdir(parents=True, exist_ok=False)
    original = url
    with requests.Session() as session:
        for _ in range(6):
            check_source_url(url)
            response = session.get(url, timeout=(10, 30), stream=True, allow_redirects=False)
            if response.is_redirect:
                url = urljoin(url, response.headers["Location"])
                response.close()
                continue
            response.raise_for_status()
            payload = bytearray()
            try:
                for chunk in response.iter_content(65536):
                    payload.extend(chunk)
                    if len(payload) > 20_000_000:
                        raise ValueError("official report exceeds 20 MB")
            finally:
                response.close()
            break
        else:
            raise ValueError("too many official redirects")
    if not payload:
        raise ValueError("empty official report")
    report = directory / "report.bin"
    atomic_write_bytes(local_write(context, report), bytes(payload))
    source = directory / "source.json"
    atomic_write_json(local_write(context, source), {"schema_version": "brazil-soy-source/1",
        "source_url": original, "final_url": url, "sha256": sha256_file(report),
        "retrieved_at": utc_now(), "published_at": published_at})
    return source


def parse_progress(source: Path) -> dict:
    """Read current and preceding-week soybean cells; official references remain references."""
    import re
    info = strict_json(source)
    report = source.parent / "report.bin"
    if sha256_file(report) != info["sha256"]:
        raise ValueError("archived report identity mismatch")
    check_source_url(info["source_url"])
    with report.open("rb") as handle:
        wb = load_workbook(handle, read_only=True, data_only=False, keep_links=False)
        try:
            if wb.sheetnames != ["Progresso de safra"]:
                raise ValueError("CONAB progress sheet layout changed")
            ws = wb.active
            rows = list(ws.iter_rows())
            records = []
            for index, row in enumerate(rows):
                title = str(row[1].value or "").strip() if len(row) >= 6 else ""
                match = re.fullmatch(r"Soja\s*-\s*Safra\s+(20\d{2}/(?:20)?\d{2})", title)
                if not match:
                    continue
                season = normalize_season(match.group(1))
                mode = str(rows[index + 2][1].value).strip()
                metric = {"Semeadura": "PLANTED", "Colheita": "HARVESTED"}.get(mode)
                if not metric:
                    raise ValueError("unknown soybean progress section")
                header = rows[index + 3]
                if str(header[1].value).strip() != "Estado" or "5" not in str(header[5].value):
                    raise ValueError("CONAB column layout changed")
                dates = rows[index + 5]
                if not all(isinstance(dates[c].value, datetime) for c in (2, 3, 4)):
                    raise ValueError("CONAB cutoff cells are not dates")
                region_rows = rows[index + 6:index + 19]
                if len(region_rows) != 13 or [str(r[1].value).strip() for r in region_rows[:12]] != list(STATE_NAMES):
                    raise ValueError("CONAB state ordering changed")
                if not re.fullmatch(r"12\s+estados", str(region_rows[-1][1].value).strip(), re.IGNORECASE):
                    raise ValueError("CONAB aggregate coverage changed")
                for row_offset, region_row in enumerate(region_rows):
                    region = STATES[row_offset] if row_offset < 12 else "BR"
                    source_row = index + 7 + row_offset
                    for column in (3, 4):
                        cell = region_row[column]
                        if cell.value is None:
                            continue
                        if cell.data_type == "f" or type(cell.value) not in (float, int) or not 0 <= cell.value <= 1:
                            raise ValueError("CONAB fraction is invalid or formula-based")
                        reference = None
                        if column == 4:
                            refs = [region_row[c] for c in (2, 5)]
                            if any(c.data_type == "f" or (c.value is not None and (type(c.value) not in (float, int) or not 0 <= c.value <= 1)) for c in refs):
                                raise ValueError("CONAB reference fraction is invalid")
                            reference = {"last_season": refs[0].value * 100 if refs[0].value is not None else None,
                                "five_season_mean": refs[1].value * 100 if refs[1].value is not None else None,
                                "locator": f"{ws.title}!C{source_row},F{source_row}"}
                        records.append(observation(season, region, metric, dates[column].value.date(), cell.value * 100,
                            source_url=info["source_url"], source_sha256=info["sha256"],
                            locator=f"{ws.title}!{'DE'[column - 3]}{source_row};cutoff={'DE'[column - 3]}{index + 6}",
                            retrieved_at=info["retrieved_at"], published_at=info["published_at"], reference=reference))
        finally:
            wb.close()
    if sha256_file(report) != info["sha256"]:
        raise ValueError("report changed during parsing")
    return validate_bundle({"schema_version": SCHEMA, "generated_at": utc_now(), "records": records,
                            "import_notes": ["CONAB数值单元格截止日；上季及官方五季均值单独作为参考值，未写成历史观测。"]})


def merge_observations(baseline: dict, updates: list[dict], *, allow_revisions: bool = False) -> tuple[dict, dict]:
    validate_bundle(baseline)
    validate_bundle({"schema_version": SCHEMA, "generated_at": utc_now(), "records": updates, "import_notes": []})
    key = lambda row: tuple(row[k] for k in ("season", "region", "metric", "date"))
    rows = {key(row): dict(row) for row in baseline["records"]}
    stats = {"added": 0, "revised": 0, "unchanged": 0}
    for item in updates:
        old = rows.get(key(item))
        if old is None:
            stats["added"] += 1
        elif old["value"] == item["value"] and (item["reference"] is None or item["reference"] == old["reference"]):
            # A new report repeats the prior week with a new file SHA and no
            # reference columns. Keep the original verified record and its
            # official references instead of manufacturing a history revision.
            stats["unchanged"] += 1
            continue
        else:
            if not allow_revisions:
                raise ValueError(f"historical revision requires explicit choice: {key(item)}")
            stats["revised"] += 1
        rows[key(item)] = dict(item)
    if not stats["added"] and not stats["revised"]:
        return baseline, stats
    return validate_bundle({**baseline, "generated_at": utc_now(), "records": list(rows.values())}), stats


def prepare_update(context: RuntimeContext, baseline: Path, observations: Path, *, allow_revisions=False) -> Path:
    identity = sha256_file(baseline)
    current, updates = load_bundle(baseline), load_bundle(observations)
    sources = {}
    for metadata in (context.runtime_root / "raw/brazil_soy").glob("*/source.json"):
        info = strict_json(metadata)
        report = metadata.parent / "report.bin"
        if report.is_file() and sha256_file(report) == info["sha256"]:
            sources[(info["source_url"], info["sha256"])] = info
    for row in updates["records"]:
        info = sources.get((row["source_url"], row["source_sha256"]))
        if (row["date_basis"] != "report_cutoff" or info is None or row["retrieved_at"] != info["retrieved_at"]
                or row["published_at"] != info["published_at"]):
            raise ValueError("new observation lacks verified archived report evidence")
    merged, stats = merge_observations(current, updates["records"], allow_revisions=allow_revisions)
    if sha256_file(baseline) != identity:
        raise ValueError("baseline changed during preparation")
    return save_candidate(context, merged, identity, stats)


def activate_local(context: RuntimeContext, candidate: Path) -> Path:
    root = (context.runtime_root / "candidates/brazil_soy").resolve()
    if candidate.is_symlink() or root not in candidate.resolve().parents:
        raise ValueError("candidate must be an ordinary local candidate")
    metadata = strict_json(candidate.parent / "candidate.json")
    identity = sha256_file(candidate)
    bundle = load_bundle(candidate)
    if metadata.get("schema_version") != "brazil-soy-candidate/1" or identity != metadata["payload_sha256"]:
        raise ValueError("candidate identity mismatch")
    stable = local_write(context, context.runtime_root / STABLE_RELATIVE_PATH)
    stable.parent.mkdir(parents=True, exist_ok=True)
    with FileLock(str(local_write(context, context.runtime_root / "brazil-soy.lock")), timeout=0):
        baseline = sha256_file(stable) if stable.is_file() else None
        if baseline == identity:
            return stable
        if baseline != metadata["baseline_sha256"]:
            raise ValueError("stable data moved after preparation")
        if stable.exists():
            backup = local_write(context, context.runtime_root / "backups/brazil_soy" / f"{uuid.uuid4().hex}.json")
            backup.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(stable, backup)
            if sha256_file(backup) != baseline:
                raise ValueError("backup identity mismatch")
        if sha256_file(candidate) != identity:
            raise ValueError("candidate changed before activation")
        atomic_write_json(local_write(context, stable), bundle)
        if sha256_file(stable) != identity:
            raise ValueError("stable serialized identity differs from candidate")
    return stable
