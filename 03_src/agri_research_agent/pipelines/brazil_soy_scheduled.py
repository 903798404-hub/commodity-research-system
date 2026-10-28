"""Discover official CONAB weeks and prepare additive, evidence-bound updates."""
from __future__ import annotations

import base64
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import json
from pathlib import Path
import re
from urllib.parse import urljoin

from bs4 import BeautifulSoup
from openpyxl import load_workbook
import requests

from .brazil_soy import SOURCE_URL, check_source_url, load_bundle, normalize_season, observation, utc_now, validate_bundle
from .brazil_soy_update import parse_progress, merge_observations, latest_dates
from .conab_growth_chart import read_growth_chart

MONTHS = {'janeiro':1,'fevereiro':2,'março':3,'abril':4,'maio':5,'junho':6,
          'julho':7,'agosto':8,'setembro':9,'outubro':10,'novembro':11,'dezembro':12}

GROWTH_STAGES = (('Emergência','EMERGENCE'),('Desenvolvimento vegetativo','VEGETATIVE'),
    ('Floração','FLOWERING'),('Enchimento de grãos','GRAIN_FILLING'),
    ('Maturação','MATURING'),('Colheita','STAGE_HARVEST'))


def native_growth_values(sheet, start):
    """Read the official chart's bound cells at its declared display precision.

    Hidden labels stay absent, as with the older printed-image reader. Never
    derive values from chart geometry, formulas or an unverified cache alone.
    """
    charts=[c for c in sheet._charts if start+4 <= c.anchor._from.row+1 < start+23]
    if len(charts)!=1 or charts[0].tagname!='doughnutChart' or len(charts[0].series)!=1:
        raise ValueError('CONAB native soybean chart ambiguous')
    series=charts[0].series[0]
    prefix=f"'{sheet.title}'!"
    cat=series.cat.strRef if series.cat else None
    val=series.val.numRef if series.val else None
    if (cat is None or val is None or cat.f!=prefix+f'$J${start}:$J${start+5}'
        or val.f!=prefix+f'$K${start}:$K${start+5}' or cat.strCache is None
        or val.numCache is None or val.numCache.formatCode!='0.0%'):
        raise ValueError('CONAB native soybean chart bindings changed')
    if cat.strCache.ptCount!=6 or val.numCache.ptCount!=6:
        raise ValueError('CONAB native soybean chart cache incomplete')
    labels={p.idx:p.v for p in cat.strCache.pt};numbers={p.idx:p.v for p in val.numCache.pt}
    if set(labels)!=set(range(6)) or set(numbers)!=set(range(6)):
        raise ValueError('CONAB native soybean chart cache indices changed')
    display=series.dLbls
    if display is None or display.showVal is not True:
        raise ValueError('CONAB native soybean chart value labels changed')
    overrides={label.idx:label for label in display.dLbl}
    if len(overrides)!=len(display.dLbl) or not set(overrides)<=set(range(6)):
        raise ValueError('CONAB native soybean chart label indices changed')
    values={};locations={}
    for index,(label,metric) in enumerate(GROWTH_STAGES):
        name=sheet.cell(start+index,10);cell=sheet.cell(start+index,11)
        if (name.value!=label or labels[index]!=label or cell.data_type!='n'
            or type(cell.value) not in (int,float) or cell.number_format!='0.0%'):
            raise ValueError('CONAB native soybean source cell changed')
        value=Decimal(str(cell.value));cached=Decimal(str(numbers[index]))
        if (not value.is_finite() or not cached.is_finite() or not 0<=value<=1
            or abs(value-cached)>Decimal('0.000000000001')):
            raise ValueError('CONAB native soybean source/cache mismatch')
        override=overrides.get(index)
        if override is not None and override.showVal is False:
            continue
        values[metric]=float((value*100).quantize(Decimal('0.1'),rounding=ROUND_HALF_UP))
        locations[metric]=f'{sheet.title}!{name.coordinate}:{cell.coordinate};chart_row={charts[0].anchor._from.row+1};display_format=0.0%'
    if not values or not 99.5<=sum(values.values())<=100.5:
        raise ValueError('CONAB native soybean displayed labels incomplete')
    return values,locations


def download(url, session):
    check_source_url(url)
    for _ in range(6):
        response = session.get(url, timeout=(15, 45), allow_redirects=False)
        if response.is_redirect:
            url = urljoin(url, response.headers['Location']); check_source_url(url)
            response.close(); continue
        response.raise_for_status()
        raw = response.content; response.close()
        if not raw or len(raw) > 20_000_000:
            raise ValueError('CONAB response exceeds source bound')
        return raw
    raise ValueError('CONAB redirect limit exceeded')


def report_links(raw, base):
    soup = BeautifulSoup(raw, 'html.parser')
    found = {}
    for a in soup.find_all('a', href=True):
        label = a.get_text(' ', strip=True)
        if re.search(r'Acompanhamento das Lavouras\s*-', label, re.I):
            m = re.search(r'a\s*(\d{2}/\d{2}/\d{2})', label)
            if m:
                url = urljoin(base, a['href']); check_source_url(url)
                found[url] = datetime.strptime(m.group(1), '%d/%m/%y').date()
    return sorted(found.items(), key=lambda item: item[1], reverse=True)


def weekly_assets(raw, base):
    soup = BeautifulSoup(raw, 'html.parser')
    published = soup.select_one('.documentPublished')
    if published is None:
        raise ValueError('CONAB official publication timestamp missing')
    m = re.search(r'\d{2}/\d{2}/\d{4}', published.get_text(' ', strip=True))
    if m is None:
        raise ValueError('CONAB publication date layout changed')
    pub = datetime.strptime(m.group(), '%d/%m/%Y').date().isoformat()
    assets = {}
    for a in soup.find_all('a', href=True):
        label = a.get_text(' ', strip=True)
        kind = 'progress' if label.startswith('Plantio e Colheita') else 'growth' if label.startswith('Fenologia das Lavouras') else None
        if kind:
            url = urljoin(base, a['href']); check_source_url(url)
            if kind in assets and assets[kind] != url:
                raise ValueError('ambiguous CONAB weekly asset')
            assets[kind] = url
    if set(assets) != {'progress','growth'}:
        raise ValueError('CONAB weekly assets incomplete')
    return pub, assets


def parse_growth(raw, source):
    book = load_workbook(BytesIO(raw), data_only=False)
    try:
        if book.sheetnames not in (['Progresso de safra'], ['Fenologia das Lavouras']):
            raise ValueError('CONAB growth workbook layout changed')
        sheet = book.active
        headers = [(row[1].row, str(row[1].value)) for row in sheet.iter_rows()
                   if len(row)>1 and re.fullmatch(r'Soja\s*-\s*Safra\s+20\d{2}/(?:20)?\d{2}', str(row[1].value or '').strip())]
        if len(headers) != 1:
            raise ValueError('CONAB soybean growth section ambiguous')
        start, title = headers[0]; season = normalize_season(title.rsplit(' ',1)[1])
        text = ' '.join(str(sheet.cell(r,2).value or '') for r in range(start+1,start+5))
        m = re.search(r'Semana de\s+\d+\s+a\s+(\d+)\s+de\s+(\w+)\s+de\s+(20\d{2})', text)
        if m is None or m.group(2) not in MONTHS:
            raise ValueError('CONAB growth cutoff missing')
        cutoff = date(int(m.group(3)),MONTHS[m.group(2)],int(m.group(1)))
        if sheet.title=='Fenologia das Lavouras':
            values,locations=native_growth_values(sheet,start)
        else:
            charts = [img for img in sheet._images if start+4 <= img.anchor._from.row+1 < start+23 and img.width>900]
            if len(charts) != 1:
                raise ValueError('CONAB soybean growth image ambiguous')
            chart = charts[0]; values = read_growth_chart(chart._data())
            locations={metric:f'{sheet.title}!B{start}:B{start+4};chart_row={chart.anchor._from.row+1};printed_label={metric}' for metric in values}
        return [observation(season,'BR',metric,cutoff,value,source_url=source['source_url'],
            source_sha256=source['sha256'],locator=locations[metric],
            retrieved_at=source['retrieved_at'],published_at=source['published_at']) for metric,value in values.items()]
    finally:
        book.close()


def collect(baseline_path, output, *, session=None, today=None):
    """Owned output only; retain every published record and reject revisions."""
    import hashlib
    baseline = load_bundle(baseline_path); bundle = baseline; today = today or date.today()
    # Harvest and growth continue after planting has ended. Resume from the
    # newest national observation rather than repeatedly replaying planting.
    current = [r['date'] for r in baseline['records'] if r['region']=='BR']
    latest = date.fromisoformat(max(current))
    session = session or requests.Session(); output = Path(output); output.mkdir(exist_ok=True)
    links = report_links(download(SOURCE_URL, session), SOURCE_URL)
    selected = [(url,end) for url,end in links if latest <= end <= today]
    if len(selected)>8:
        selected = selected[-8:]
    if not links:
        raise ValueError('CONAB report index empty')
    sources = []; checks = []
    for index,(url,end) in enumerate(reversed(selected)):
        pub, assets = weekly_assets(download(url, session),url)
        for kind, asset in assets.items():
            raw = download(asset, session); source={'source_url':asset,'sha256':hashlib.sha256(raw).hexdigest(),'retrieved_at':utc_now(),'published_at':pub}
            archive=output/f'{index:02d}-{kind}';archive.mkdir()
            (archive/'report.bin').write_bytes(raw);(archive/'source.json').write_text(json.dumps(source),encoding='utf-8')
            if kind=='progress':
                updates=parse_progress(archive/'source.json')['records']
            else:
                updates=parse_growth(raw,source)
            bundle,changes=merge_observations(bundle,updates)
            if changes['added']:
                sources.append({'kind':'report',**source,'bytes_base64':base64.b64encode(raw).decode('ascii')})
            checks.append({'kind':kind,'url':asset,'published_at':pub,**changes})
    # A repeated report may be entirely redundant. It never rewrites original
    # source identities, official references, or the last successful snapshot.
    changed=len(bundle['records']) != len(baseline['records'])
    evidence={'schema_version':'brazil-soy-source-evidence/1','sources':sources}
    return bundle,evidence,{'business_changed':changed,'record_count':len(bundle['records']),
        'latest_dates':latest_dates(bundle),'revision_keys':[],'source_checks':checks,
        'added':len(bundle['records'])-len(baseline['records'])}
