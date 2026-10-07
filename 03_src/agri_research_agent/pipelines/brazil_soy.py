"""Brazil soybean crop-season observations and read-only comparisons."""
from __future__ import annotations

import calendar
import math
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .canada_canola import sha256_file, strict_json, utc_now
from agri_research_agent.shared.chart_style import (
    CURRENT_LINE_WIDTH, CURRENT_YEAR_COLOR, GRID_COLOR, HISTORY_LINE_WIDTH,
    HISTORY_OPACITY, MEAN_COLOR, historical_year_color,
)

SCHEMA = "brazil-soy-progress/1"
STABLE_RELATIVE_PATH = Path("processed/brazil_soy/soy_weekly.json")
SOURCE_URL = "https://www.gov.br/conab/pt-br/atuacao/informacoes-agropecuarias/safras/progresso-de-safra"
REGIONS = {"BR": "全国汇总（CONAB主要12州）", "TO": "托坎廷斯 TO", "MA": "马拉尼昂 MA",
           "PI": "皮奥伊 PI", "BA": "巴伊亚 BA", "MT": "马托格罗索 MT",
           "MS": "南马托格罗索 MS", "GO": "戈亚斯 GO", "MG": "米纳斯吉拉斯 MG",
           "SP": "圣保罗 SP", "PR": "巴拉那 PR", "SC": "圣卡塔琳娜 SC", "RS": "南里奥格兰德 RS"}
STATES = tuple(REGIONS)[1:]
STATE_NAMES = ("Tocantins", "Maranhão", "Piauí", "Bahia", "Mato Grosso", "Mato Grosso do Sul",
               "Goiás", "Minas Gerais", "São Paulo", "Paraná", "Santa Catarina", "Rio Grande do Sul")
METRICS = {"PLANTED": "播种进度", "HARVESTED": "收割进度"}
STAGES = {"EMERGENCE": "出苗", "VEGETATIVE": "营养生长", "FLOWERING": "开花",
          "GRAIN_FILLING": "籽粒灌浆", "MATURING": "成熟", "STAGE_HARVEST": "收割阶段"}
FIELDS = {"season", "region", "metric", "date", "value", "source_url", "source_sha256",
          "source_locator", "date_basis", "published_at", "retrieved_at", "reference"}


def check_source_url(url: str) -> None:
    if not isinstance(url, str):
        raise ValueError("official source URL must be a string")
    parsed = urlparse(url)
    if (parsed.scheme != "https" or parsed.hostname not in {"www.gov.br", "antigo.conab.gov.br", "www.conab.gov.br"}
            or parsed.username or parsed.password or parsed.port not in (None, 443)
            or (parsed.hostname == "www.gov.br" and not parsed.path.startswith("/conab/"))):
        raise ValueError("source must be an official CONAB HTTPS URL")


def season_start(season: str) -> int:
    if not isinstance(season, str) or not re.fullmatch(r"20\d{2}/20\d{2}", season):
        raise ValueError("crop season must be YYYY/YYYY")
    start, end = map(int, season.split("/"))
    if end != start + 1:
        raise ValueError("crop season years must be consecutive")
    return start


def season_name(start: int) -> str:
    return f"{start}/{start + 1}"


def current_season(today: date | None = None) -> str:
    today = today or date.today()
    return season_name(today.year if today.month >= 9 else today.year - 1)


def normalize_season(value: str) -> str:
    value = str(value).strip()
    if re.fullmatch(r"20\d{2}/\d{2}", value):
        value = value[:5] + "20" + value[5:]
    season_start(value)
    return value


def seasonal_date(day: date) -> date:
    """Use a leap reference axis spanning September through August."""
    return date(1999 if day.month >= 9 else 2000, day.month, day.day)


def validate_bundle(bundle: dict, *, today: date | None = None) -> dict:
    today = today or date.today()
    if not isinstance(bundle, dict) or set(bundle) != {"schema_version", "generated_at", "records", "import_notes"}:
        raise ValueError("invalid Brazil bundle fields")
    if bundle["schema_version"] != SCHEMA or not isinstance(bundle["records"], list) or not bundle["records"]:
        raise ValueError("unsupported or empty Brazil bundle")
    if not isinstance(bundle["generated_at"], str) or datetime.fromisoformat(bundle["generated_at"]).utcoffset() is None:
        raise ValueError("generated_at requires a timezone")
    if not isinstance(bundle["import_notes"], list) or not all(isinstance(x, str) for x in bundle["import_notes"]):
        raise ValueError("import_notes must be strings")
    keys = set()
    for item in bundle["records"]:
        if not isinstance(item, dict) or set(item) != FIELDS:
            raise ValueError("invalid observation fields")
        if any(not isinstance(item[k], str) for k in FIELDS - {"value", "published_at", "reference"}):
            raise ValueError("observation identifiers must be strings")
        start = season_start(item["season"])
        day = date.fromisoformat(item["date"])
        if item["date"] != day.isoformat() or not date(start, 9, 1) <= day <= min(today, date(start + 1, 8, 31)):
            raise ValueError("observation outside crop season or in the future")
        if item["region"] not in REGIONS or item["metric"] not in METRICS | STAGES:
            raise ValueError("unknown region or metric")
        if item["metric"] in STAGES and item["region"] != "BR":
            raise ValueError("growth stages are national only")
        value = item["value"]
        if type(value) not in (int, float) or not math.isfinite(value) or not 0 <= value <= 100:
            raise ValueError("observation must be a finite percentage between 0 and 100")
        check_source_url(item["source_url"])
        if not re.fullmatch(r"[0-9a-f]{64}", item["source_sha256"]) or not item["source_locator"].strip():
            raise ValueError("source identity or locator is missing")
        if datetime.fromisoformat(item["retrieved_at"]).utcoffset() is None:
            raise ValueError("retrieved_at requires a timezone")
        if item["date_basis"] not in {"workbook_date", "report_cutoff"}:
            raise ValueError("unknown date basis")
        published = item["published_at"]
        if published is not None:
            if (not isinstance(published, str) or date.fromisoformat(published).isoformat() != published
                    or not day <= date.fromisoformat(published) <= today):
                raise ValueError("invalid publication date")
        if item["date_basis"] == "report_cutoff" and published is None:
            raise ValueError("official observations require publication date")
        reference = item["reference"]
        if reference is not None:
            if (item["date_basis"] != "report_cutoff" or item["metric"] not in METRICS
                    or not isinstance(reference, dict) or set(reference) != {"last_season", "five_season_mean", "locator"}
                    or not isinstance(reference["locator"], str) or not reference["locator"].strip()):
                raise ValueError("invalid official reference fields")
            for number in (reference["last_season"], reference["five_season_mean"]):
                if number is not None and (type(number) not in (int, float) or not math.isfinite(number) or not 0 <= number <= 100):
                    raise ValueError("invalid official reference percentage")
        key = tuple(item[k] for k in ("season", "region", "metric", "date"))
        if key in keys:
            raise ValueError(f"duplicate observation: {key}")
        keys.add(key)
    return bundle


def load_bundle(path: Path) -> dict:
    return validate_bundle(strict_json(path))


def observation(season: str, region: str, metric: str, day: date, value: float, *,
                source_url: str, source_sha256: str, locator: str, retrieved_at: str,
                published_at: str | None = None, reference: dict | None = None) -> dict:
    if type(value) not in (int, float):
        raise ValueError("observation value must be numeric, not boolean")
    return {"season": season, "region": region, "metric": metric, "date": day.isoformat(),
            "value": round(value, 8), "source_url": source_url, "source_sha256": source_sha256,
            "source_locator": locator, "retrieved_at": retrieved_at, "published_at": published_at,
            "date_basis": "report_cutoff" if published_at else "workbook_date", "reference": reference}


def import_workbook(path: Path) -> dict:
    identity, retrieved = sha256_file(path), utc_now()
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    records, notes = [], []
    specs = [("巴西种植", 2, 2, 3, list(STATES) + ["BR"], "PLANTED", 1),
             ("巴西收割", 2, 3, 4, list(STATES) + ["BR"], "HARVESTED", 100),
             ("巴西大豆作物进度作图", 4, 3, 4, list(STAGES), None, 100)]
    try:
        for sheet, first_row, date_col, first_value, columns, metric, scale in specs:
            if sheet not in wb.sheetnames:
                raise ValueError(f"source layout changed: missing {sheet}")
            ws = wb[sheet]
            if metric and [str(ws.cell(1, c).value).strip() for c in range(first_value, first_value + 13)] != list(STATE_NAMES) + ["All States"]:
                raise ValueError(f"source layout changed: {sheet} region columns")
            if not metric:
                labels = [str(ws.cell(2, c).value).strip() for c in range(4, 10)]
                if not (labels[0].startswith("Emerg") and labels[1] == "Desenvolvimento Vegetativo"
                        and labels[2] == "Floração" and labels[3] == "Enchimento de Grãos"
                        and labels[4] == "Maturação" and labels[5] == "Colheita"):
                    raise ValueError("source layout changed: national growth columns")
            missing = 0
            # Stop at first empty raw season/date pair: formatting/pivots below are not observations.
            for row_index, row in enumerate(ws.iter_rows(min_row=first_row, max_col=first_value + len(columns) - 1), first_row):
                season_value, day = row[0].value, row[date_col - 1].value
                if season_value is None and day is None:
                    break
                if not isinstance(day, datetime):
                    raise ValueError(f"invalid raw date: {sheet} row {row_index}")
                if metric == "PLANTED" and row[0].data_type == "f":
                    expected = f'=IF(MONTH(B{row_index})>=9,YEAR(B{row_index})&"/"&YEAR(B{row_index})+1,YEAR(B{row_index})-1&"/"&YEAR(B{row_index}))'
                    if season_value != expected:
                        raise ValueError(f"unknown crop-season formula: {sheet}!A{row_index}")
                    season = current_season(day.date())
                else:
                    season = normalize_season(season_value)
                for offset, key in enumerate(columns):
                    cell = row[first_value + offset - 1]
                    value = cell.value
                    coordinate = f"{get_column_letter(first_value + offset)}{row_index}"
                    if value is None or cell.data_type == "e":
                        missing += 1
                        continue
                    if cell.data_type == "f":
                        notes.append(f"原始表含推算公式，未导入 {sheet}!{coordinate}：{value}")
                        continue
                    if type(value) not in (float, int):
                        raise ValueError(f"raw value is formula/non-numeric: {sheet}!{coordinate}")
                    if (metric == "PLANTED" and season == "2024/2025" and day.date() == date(2024, 9, 22)
                            and value == {"MT": 0.003, "PR": 0.01, "BR": 0.002}.get(key)):
                        notes.append(f"待核实单位，隔离 {sheet}!{coordinate}：原值{value}，未参与图表或均值。")
                        continue
                    records.append(observation(season, key if metric else "BR", metric or key, day.date(), value * scale,
                        source_url=SOURCE_URL, source_sha256=identity, locator=f"{sheet}!A{row_index},{get_column_letter(date_col)}{row_index},{coordinate}",
                        retrieved_at=retrieved))
                if not metric:
                    numbers = [r.value for r in row[3:9] if type(r.value) in (int, float)]
                    total = sum(numbers) * 100
                    if numbers and abs(total - 100) > 1.1:
                        notes.append(f"全国阶段 {day.date()}：有效项{len(numbers)}/6，合计{total:.2f}%，不补零或截断。")
            notes.append(f"{sheet}：{missing}个空值或Excel错误未导入。")
    finally:
        wb.close()
    if sha256_file(path) != identity:
        raise ValueError("source workbook changed during import")
    notes.extend(["历史日期来自用户文件，发布日期未经逐期核实。", "全国汇总沿用原始数据；官方更新覆盖主要12州，约96%面积。",
                  "全国阶段使用原始D:I分布；未导入累计公式，也未把缺失出苗转成零。"])
    return validate_bundle({"schema_version": SCHEMA, "generated_at": retrieved, "records": records, "import_notes": notes})


def observations_frame(bundle: dict) -> pd.DataFrame:
    frame = pd.DataFrame(validate_bundle(bundle)["records"])
    frame["date"] = pd.to_datetime(frame["date"])
    return frame.sort_values("date").reset_index(drop=True)


def match_history(series: pd.DataFrame, target: date, target_season: str, historical_season: str) -> dict | None:
    year = target.year + season_start(historical_season) - season_start(target_season)
    try:
        end = date(year, target.month, target.day)
    except ValueError:
        end = date(year, 2, 28)
    rows = series.loc[(series["season"] == historical_season) & (series["date"] <= pd.Timestamp(end))
                      & (series["date"] >= pd.Timestamp(end - timedelta(days=7)))]
    return rows.iloc[-1].to_dict() if not rows.empty else None


def compare_metric(frame: pd.DataFrame, region: str, metric: str, season: str) -> dict:
    series = frame.loc[(frame["region"] == region) & (frame["metric"] == metric)]
    rows = series.loc[series["season"] == season]
    result = dict(current=None, previous_change=None, previous_days=None, last_season=None,
                  mean=None, samples=[], reference_label="历史计算")
    if rows.empty:
        return result
    current = rows.iloc[-1].to_dict()
    result["current"] = current
    if len(rows) > 1:
        previous = rows.iloc[-2]
        result["previous_change"] = current["value"] - previous["value"]
        result["previous_days"] = (current["date"] - previous["date"]).days
    start = season_start(season)
    target = current["date"].date()
    result["last_season"] = match_history(series, target, season, season_name(start - 1))
    result["samples"] = [item for year in range(start - 5, start)
                         if (item := match_history(series, target, season, season_name(year))) is not None]
    if result["samples"]:
        result["mean"] = sum(x["value"] for x in result["samples"]) / len(result["samples"])
    reference = current["reference"]
    if reference:
        if reference["last_season"] is not None:
            result["last_season"] = {"value": reference["last_season"], "date": None}
        if reference["five_season_mean"] is not None:
            result["mean"] = reference["five_season_mean"]
            result["reference_label"] = "官方五季均值"
    return result


def comparison_table(frame: pd.DataFrame, region: str, season: str, *, stages: bool = False) -> pd.DataFrame:
    rows = []
    for metric, label in (STAGES if stages else METRICS).items():
        comparison = compare_metric(frame, region, metric, season)
        current, last, mean = comparison["current"], comparison["last_season"], comparison["mean"]
        rows.append({"指标": label, "日期": current["date"].strftime("%Y-%m-%d") if current else "—",
                     "最新": current["value"] if current else None, "较上次": comparison["previous_change"],
                     "上季同期": last["value"] if last else None, "同期均值": mean,
                     "较均值": current["value"] - mean if current and mean is not None else None,
                     "均值依据": "官方五季" if comparison["reference_label"] == "官方五季均值" else f"{len(comparison['samples'])}/5季"})
    return pd.DataFrame(rows)


def seasonal_figure(frame: pd.DataFrame, region: str, metric: str, season: str, extra: list[str]):
    import plotly.graph_objects as go
    series = frame.loc[(frame["region"] == region) & (frame["metric"] == metric)]
    start = season_start(season)
    figure = go.Figure()
    for selected in sorted(set(extra) | {season_name(start - 1)}):
        if selected == season:
            continue
        rows = series.loc[series["season"] == selected]
        if rows.empty:
            continue
        figure.add_trace(go.Scatter(x=[seasonal_date(x.date()) for x in rows["date"]], y=rows["value"],
            name=selected, mode="lines", opacity=HISTORY_OPACITY,
            line=dict(color=historical_year_color(season_start(selected)), width=HISTORY_LINE_WIDTH),
            customdata=rows["date"].dt.strftime("%Y-%m-%d"), hovertemplate="%{customdata} · %{y:.1f}%<extra>%{fullData.name}</extra>"))
    axis = sorted({seasonal_date(x.date()) for x in series.loc[series["season"].isin([season_name(y) for y in range(start - 5, start + 1)]), "date"]})
    means, counts = [], []
    for point in axis:
        target_year = start if point.month >= 9 else start + 1
        target = date(target_year, point.month, 28 if point.month == 2 and point.day == 29 and not calendar.isleap(target_year) else point.day)
        samples = [item for year in range(start - 5, start) if (item := match_history(series, target, season, season_name(year))) is not None]
        means.append(sum(x["value"] for x in samples) / len(samples) if samples else None)
        counts.append(len(samples))
    figure.add_trace(go.Scatter(x=axis, y=means, name="前五季同期均值", connectgaps=False,
        mode="lines", line=dict(color=MEAN_COLOR, width=2, dash="dash"), customdata=counts,
        hovertemplate="%{x|%m-%d} · %{y:.1f}%（%{customdata}/5季）<extra>历史计算均值</extra>"))
    current = series.loc[series["season"] == season]
    if not current.empty:
        figure.add_trace(go.Scatter(x=[seasonal_date(x.date()) for x in current["date"]], y=current["value"],
            name=f"{season}（当前季）", mode="lines+markers", opacity=1,
            line=dict(color=CURRENT_YEAR_COLOR, width=CURRENT_LINE_WIDTH), marker=dict(size=5),
            customdata=current["date"].dt.strftime("%Y-%m-%d"), hovertemplate="%{customdata} · %{y:.1f}%<extra>当前季</extra>"))
    figure.update_layout(height=300, margin=dict(l=8, r=8, t=45, b=10),
        legend=dict(orientation="h", y=1.1, yanchor="bottom", font=dict(size=10)),
        xaxis=dict(tickformat="%m-%d", title=None, nticks=6, gridcolor=GRID_COLOR),
        yaxis=dict(title="%", range=[0, 102], gridcolor=GRID_COLOR), hovermode="closest")
    return figure
