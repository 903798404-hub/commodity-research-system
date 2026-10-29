"""Prepare an offline regional candidate without altering the existing Liuzhou history."""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd

from agri_research_agent.data_sources.domestic_sugar_history import (
    SugarHistoryError, file_sha256, load_sugar_history,
)
from agri_research_agent.data_sources.domestic_sugar_regional_history import (
    REGIONS, SCHEMA, load_regional_history, read_json, validate_payload,
)


def _capture_rows(capture: dict, fields: set[str]) -> list[dict]:
    try:
        data = capture["data"]; names = data["fields"]; items = data["items"]
        if (capture.get("code") != 0 or not isinstance(names, list) or len(names) != len(set(names))
                or not fields.issubset(names) or not isinstance(items, list)
                or any(not isinstance(row, list) or len(row) != len(names) for row in items)):
            raise ValueError("capture schema")
        return [dict(zip(names, row)) for row in items]
    except (KeyError, TypeError, ValueError) as exc:
        raise SugarHistoryError("补充主力映射或原始收盘行情结构无效") from exc


def _extension_records(path: Path) -> list[dict]:
    capture = read_json(path)
    try:
        mapping = _capture_rows(capture["mapping"], {"ts_code", "trade_date", "mapping_ts_code"})
        quotes = _capture_rows(capture["daily"], {"ts_code", "trade_date", "close"})
    except KeyError as exc:
        raise SugarHistoryError("补充行情必须同时包含映射和原始收盘价") from exc
    daily = {}
    for row in quotes:
        key = row["trade_date"], row["ts_code"]
        if key in daily:
            raise SugarHistoryError("补充主力行情业务键重复")
        daily[key] = row["close"]
    records = []; seen = set()
    for row in mapping:
        if row["ts_code"] != "SR.ZCE" or row["trade_date"] in seen:
            raise SugarHistoryError("补充行情必须为唯一的SR主力映射")
        seen.add(row["trade_date"])
        try:
            if not isinstance(row["trade_date"], str) or not re.fullmatch(r"\d{8}", row["trade_date"]):
                raise ValueError("date format")
            date = dt.datetime.strptime(row["trade_date"], "%Y%m%d").date().isoformat()
        except (TypeError, ValueError) as exc:
            raise SugarHistoryError("补充主力日期无效") from exc
        code = row["mapping_ts_code"]
        records.append(dict(business_date=date, ts_code=code,
                            futures_price=daily.get((row["trade_date"], code))))
    return records


def prepare_regional_sugar_history(
    spot_workbook: Path, spot_sha256: str, main_history_root: Path,
    candidate_root: Path, *, main_extension: Path | None = None,
    extension_sha256: str | None = None,
) -> dict:
    spot_workbook = Path(spot_workbook).resolve(); candidate_root = Path(candidate_root).resolve()
    if file_sha256(spot_workbook) != spot_sha256:
        raise SugarHistoryError("钢联现货原表身份不一致")
    base = load_sugar_history(main_history_root)
    records = []
    workbook = openpyxl.load_workbook(spot_workbook, read_only=True, data_only=True)
    try:
        sheet = workbook["导出数据"]
        expected = [f"白糖：现货价：{region}（日）" for region in REGIONS]
        if ([sheet.cell(2, col).value for col in (2, 3, 4)] != expected
                or any(sheet.cell(3, col).value != "日" for col in (2, 3, 4))
                or any(sheet.cell(4, col).value != "当地贸易商报价平均值" for col in (2, 3, 4))):
            raise SugarHistoryError("钢联三地指标名称、频度或贸易商均价口径不符")
        for row_number, row in enumerate(sheet.iter_rows(min_row=5, max_col=4, values_only=True), 5):
            if all(value is None for value in row):
                continue
            date = row[0]
            if not isinstance(date, (dt.date, dt.datetime)):
                raise SugarHistoryError("钢联现货日期不是Excel日期")
            if isinstance(date, dt.datetime):
                if date.time() != dt.time():
                    raise SugarHistoryError("钢联现货日期包含非零时间")
                date = date.date()
            for col, region in enumerate(REGIONS, 1):
                value = row[col]
                if isinstance(value, bool) or (value is not None and not isinstance(value, (int, float))):
                    raise SugarHistoryError("钢联现货价格须为数值或空值")
                records.append(dict(business_date=date.isoformat(), region=region, cash_price=value,
                                    source_cell=f"{'BCD'[col-1]}{row_number}"))
    except KeyError as exc:
        raise SugarHistoryError("钢联现货原表缺少导出数据工作表") from exc
    finally:
        workbook.close()
    if file_sha256(spot_workbook) != spot_sha256:
        raise SugarHistoryError("钢联现货原表在读取期间发生变化")
    main = base.main_basis[["business_date", "ts_code", "futures_price"]].copy()
    main["business_date"] = main.business_date.dt.strftime("%Y-%m-%d")
    main_records = main.astype(object).where(main.notna(), None).to_dict("records")
    by_date = {row["business_date"]: row for row in main_records}
    if main_extension is not None:
        if not extension_sha256 or file_sha256(main_extension) != extension_sha256:
            raise SugarHistoryError("补充主力行情来源身份不一致")
        for row in _extension_records(main_extension):
            old = by_date.get(row["business_date"])
            if old is not None and old != row:
                raise SugarHistoryError("补充主力行情与既有历史冲突，拒绝替换")
            by_date[row["business_date"]] = row
        if file_sha256(main_extension) != extension_sha256:
            raise SugarHistoryError("补充主力行情在读取期间发生变化")
    payload = dict(schema_version=SCHEMA, currency="CNY", unit="CNY/tonne", price_field="close",
        spot_definition="当地贸易商报价平均值", source_workbook=dict(name=spot_workbook.name,
        sha256=spot_sha256, sheet="导出数据", columns=dict(zip(REGIONS, "BCD"))),
        main_source_snapshot_id=base.snapshot_id, main_extension_sha256=extension_sha256,
        refresh_policy="manual_candidate_no_scheduled_refresh", spot_records=records,
        main_records=sorted(by_date.values(), key=lambda row: row["business_date"]))
    daily = validate_payload(payload)
    content = json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False).encode("utf-8")
    sid = hashlib.sha256(content).hexdigest()
    if candidate_root.exists():
        existing = load_regional_history(candidate_root)
        if existing.snapshot_id != sid:
            raise SugarHistoryError("候选目录已有不同快照，请使用新的候选目录")
        return dict(status="NO_CHANGE", snapshot_id=sid, candidate_root=str(candidate_root))
    candidate_root.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="regional-sugar-", dir=candidate_root.parent) as folder:
        work = Path(folder); snapshot_dir = work / "snapshots" / sid; snapshot_dir.mkdir(parents=True)
        (snapshot_dir / "history.json").write_bytes(content)
        (work / "current.json").write_text(json.dumps(dict(schema_version=SCHEMA, snapshot_id=sid)), encoding="utf-8")
        load_regional_history(work)
        work.rename(candidate_root)
    common = daily.loc[daily.basis.notna(), "business_date"]
    return dict(status="PREPARED", snapshot_id=sid, candidate_root=str(candidate_root), rows=len(daily),
                latest_spot=daily.business_date.max().date().isoformat(),
                latest_basis=common.max().date().isoformat() if not common.empty else None)
