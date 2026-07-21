"""Extract approved static crop-weather 30-year normals from a read-only workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import sys
import uuid
from datetime import UTC, datetime
from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = PROJECT_ROOT / "03_src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from agri_research_agent.weather.crop_weather import load_weather_config  # noqa: E402


CONFIG_FILE = PROJECT_ROOT / "02_configs" / "soybean_weather_us.yaml"
CANDIDATE_ROOT = PROJECT_ROOT / "01_data" / "candidates" / "weather" / "soybean_us_30y_normal"
STABLE_FILE = PROJECT_ROOT / "01_data" / "processed" / "weather" / "soybean" / "us" / "soybean_weather_us_30y_normal.parquet"
BACKUP_ROOT = PROJECT_ROOT / "01_data" / "backups" / "weather" / "soybean_us_30y_normal"
STATUS_FILE = PROJECT_ROOT / "01_data" / "update_status" / "soybean_weather_us_30y_normal.json"
BASELINE_LABEL = "沿用旧项目30年历史同期基准"
TARGET_SHEETS = {
    "美国历史30年平均降雨": ("precipitation", "mm"),
    "美国历史30年平均最高气温": ("temperature_max", "degC"),
}
REQUIRED_COLUMNS = (
    "month_day", "country", "region", "metric", "normal_value", "unit", "baseline_label", "source_workbook_sha256", "source_sheet",
)
SCHEMA = pa.schema([
    pa.field("month_day", pa.string()), pa.field("country", pa.string()), pa.field("region", pa.string()),
    pa.field("metric", pa.string()), pa.field("normal_value", pa.float64()), pa.field("unit", pa.string()),
    pa.field("baseline_label", pa.string()), pa.field("source_workbook_sha256", pa.string()), pa.field("source_sheet", pa.string()),
])


class NormalImportError(ValueError):
    """Raised when the approved workbook cannot be safely extracted."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _expected_headers(config: dict[str, object]) -> list[str]:
    source_regions = config.get("normal_source_regions", config.get("source_regions", config["regions"]))
    return ["日期", *[str(region.get("normal_column", region["display_name"])) for region in source_regions]]


def target_sheets(config: dict[str, object]) -> dict[str, tuple[str, str]]:
    configured = config.get("normal_sheets")
    if not isinstance(configured, dict):
        return TARGET_SHEETS
    precipitation = configured.get("precipitation")
    temperature_max = configured.get("temperature_max")
    if not isinstance(precipitation, str) or not isinstance(temperature_max, str):
        raise NormalImportError("天气配置缺少30年基准工作表名")
    return {precipitation: ("precipitation", "mm"), temperature_max: ("temperature_max", "degC")}


def _validate_source_sheet(workbook: Path, sheet_name: str, expected_headers: list[str]) -> None:
    """Reject formulas in the approved date-and-state source columns."""

    wb = load_workbook(workbook, read_only=True, data_only=False)
    if sheet_name not in wb.sheetnames:
        raise NormalImportError(f"缺少授权30年基准工作表：{sheet_name}")
    ws = wb[sheet_name]
    headers = [ws.cell(1, index).value for index in range(1, len(expected_headers) + 1)]
    if headers != expected_headers:
        raise NormalImportError(f"工作表 {sheet_name} 的15州列与配置不完全一致")
    for row in ws.iter_rows(min_row=2, max_col=len(expected_headers), values_only=False):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                raise NormalImportError(f"工作表 {sheet_name} 的目标数据列包含公式，拒绝提取")


def extract_normals(workbook: Path, *, config_file: Path = CONFIG_FILE) -> tuple[pd.DataFrame, dict[str, object]]:
    if not workbook.is_file() or workbook.suffix.lower() != ".xlsx":
        raise NormalImportError("输入必须是可读的 .xlsx 工作簿")
    config = load_weather_config(config_file)
    expected_headers = _expected_headers(config)
    sheet_specs = target_sheets(config)
    source_hash = sha256_file(workbook)
    frames: list[pd.DataFrame] = []
    sheet_info: dict[str, object] = {}
    for sheet_name, (metric, unit) in sheet_specs.items():
        _validate_source_sheet(workbook, sheet_name, expected_headers)
        wb = load_workbook(workbook, read_only=True, data_only=True)
        ws = wb[sheet_name]
        rows: list[dict[str, object]] = []
        source_regions = list(config.get("normal_source_regions", config.get("source_regions", config["regions"])))
        source_index = {str(region["key"]): index + 1 for index, region in enumerate(source_regions)}
        for values in ws.iter_rows(min_row=2, max_col=len(expected_headers), values_only=True):
            raw_date = values[0]
            if raw_date is None and all(value is None for value in values[1:]):
                continue
            parsed_date = pd.to_datetime(raw_date, errors="coerce")
            if pd.isna(parsed_date):
                raise NormalImportError(f"工作表 {sheet_name} 存在无法解释的日期结构")
            month_day = parsed_date.strftime("%m-%d")
            for region in config["regions"]:
                value = values[source_index[str(region["key"])]]
                if value is None:
                    raise NormalImportError(f"工作表 {sheet_name} 的 {month_day} 存在缺州值")
                try:
                    normal_value = float(value)
                except (TypeError, ValueError) as exc:
                    raise NormalImportError(f"工作表 {sheet_name} 的 {month_day} 存在非数值基准") from exc
                rows.append({"month_day": month_day, "country": str(config["country"]), "region": str(region["key"]), "metric": metric, "normal_value": normal_value, "unit": unit, "baseline_label": BASELINE_LABEL, "source_workbook_sha256": source_hash, "source_sheet": sheet_name})
        frame = pd.DataFrame(rows)
        if frame.empty:
            raise NormalImportError(f"工作表 {sheet_name} 未提取到基准记录")
        if frame.duplicated(["month_day", "region", "metric"], keep=False).any():
            raise NormalImportError(f"工作表 {sheet_name} 存在重复月日或冲突值")
        if set(frame["region"]) != {str(region["key"]) for region in config["regions"]}:
            raise NormalImportError(f"工作表 {sheet_name} 未覆盖全部展示地区")
        sheet_info[sheet_name] = {"rows": int(len(frame)), "month_day_start": str(frame["month_day"].min()), "month_day_end": str(frame["month_day"].max())}
        frames.append(frame)
    result = pd.concat(frames, ignore_index=True)
    if tuple(result.columns) != REQUIRED_COLUMNS:
        raise NormalImportError("30年基准输出字段不符合契约")
    return result, {"source_sha256": source_hash, "sheets": sheet_info}


def validate_normals(frame: pd.DataFrame, config: dict[str, object]) -> dict[str, object]:
    if tuple(frame.columns) != REQUIRED_COLUMNS or frame[list(REQUIRED_COLUMNS)].isna().any().any():
        raise NormalImportError("30年基准存在缺失字段或Schema不正确")
    if set(frame["metric"]) != {"precipitation", "temperature_max"} or set(frame["country"]) != {str(config["country"])}:
        raise NormalImportError("30年基准指标或国家不符合契约")
    expected_regions = {str(region["key"]) for region in config["regions"]}
    if set(frame["region"]) != expected_regions:
        raise NormalImportError("30年基准地区集合不符合展示地区配置")
    if frame.duplicated(["month_day", "region", "metric"], keep=False).any():
        raise NormalImportError("30年基准存在重复或冲突月日")
    expected_days = 365
    counts = frame.groupby(["metric", "region"])["month_day"].nunique()
    if not (counts == expected_days).all():
        raise NormalImportError("30年基准不是每州365个可解释MM-DD日值")
    return {"row_count": int(len(frame)), "metric_counts": frame.groupby("metric").size().astype(int).to_dict(), "month_day_range": {"start": str(frame["month_day"].min()), "end": str(frame["month_day"].max())}}


def _atomic_json(path: Path, value: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    os.replace(temporary, path)


def promote(candidate: Path, status: dict[str, object], *, stable_file: Path | None = None, backup_root: Path | None = None, status_file: Path | None = None) -> str:
    stable_file = stable_file or STABLE_FILE
    backup_root = backup_root or BACKUP_ROOT
    status_file = status_file or STATUS_FILE
    if stable_file.exists():
        backup = backup_root / datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ") / stable_file.name
        backup.parent.mkdir(parents=True, exist_ok=False)
        shutil.copy2(stable_file, backup)
        if sha256_file(backup) != sha256_file(stable_file):
            raise NormalImportError("30年基准备份哈希不一致")
    stable_file.parent.mkdir(parents=True, exist_ok=True)
    temporary = stable_file.with_name(f".{stable_file.name}.{uuid.uuid4().hex}.tmp")
    shutil.copy2(candidate, temporary)
    os.replace(temporary, stable_file)
    stable_hash = sha256_file(stable_file)
    status.update({"status": "success", "stable_file_sha256": stable_hash})
    _atomic_json(status_file, status)
    return stable_hash


def _country_paths(config_file: Path, config: dict[str, object]) -> tuple[Path, Path, Path, Path, str]:
    slug = str(config.get("data_slug", "us"))
    if config_file == CONFIG_FILE:
        return CANDIDATE_ROOT, STABLE_FILE, BACKUP_ROOT, STATUS_FILE, slug
    crop = str(config["crop"])
    return (
        PROJECT_ROOT / "01_data" / "candidates" / "weather" / f"{crop}_{slug}_30y_normal",
        PROJECT_ROOT / "01_data" / "processed" / "weather" / crop / slug / f"{crop}_weather_{slug}_30y_normal.parquet",
        PROJECT_ROOT / "01_data" / "backups" / "weather" / f"{crop}_{slug}_30y_normal",
        PROJECT_ROOT / "01_data" / "update_status" / f"{crop}_weather_{slug}_30y_normal.json",
        slug,
    )


def import_workbook(workbook: Path, *, run_id: str, promote_stable: bool, config_file: Path = CONFIG_FILE) -> dict[str, object]:
    config = load_weather_config(config_file)
    candidate_root, stable_file, backup_root, status_file, slug = _country_paths(config_file, config)
    candidate_dir = candidate_root / run_id
    if candidate_dir.exists():
        raise NormalImportError(f"候选目录已存在：{candidate_dir}")
    candidate_dir.mkdir(parents=True)
    frame, source = extract_normals(workbook, config_file=config_file)
    quality = validate_normals(frame, config)
    candidate = candidate_dir / f"{config['crop']}_weather_{slug}_30y_normal.parquet"
    pq.write_table(pa.Table.from_pandas(frame, schema=SCHEMA, preserve_index=False), candidate, compression="zstd")
    status: dict[str, object] = {"status": "candidate_validated", "run_id": run_id, "source_filename": workbook.name, "source_sha256": source["source_sha256"], "source_sheets": list(target_sheets(config)), "generated_at": datetime.now(UTC).isoformat(), **quality}
    _atomic_json(candidate_dir / "candidate_report.json", {"status": status, "source": source})
    result: dict[str, object] = {"candidate": str(candidate), "candidate_sha256": sha256_file(candidate), "status": status}
    if promote_stable:
        result["stable"] = str(stable_file)
        result["stable_sha256"] = promote(candidate, status, stable_file=stable_file, backup_root=backup_root, status_file=status_file)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--config", type=Path, default=CONFIG_FILE)
    parser.add_argument("--run-id", default=None)
    parser.add_argument("--promote", action="store_true")
    args = parser.parse_args()
    run_id = args.run_id or datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    print(json.dumps(import_workbook(args.workbook, run_id=run_id, promote_stable=args.promote, config_file=args.config), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
