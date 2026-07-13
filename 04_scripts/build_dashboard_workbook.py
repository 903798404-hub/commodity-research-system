from __future__ import annotations

import datetime as dt
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DASHBOARD_NAME = "价差图看板.xlsx"
IMAGE_WIDTH_PX = 480
BLOCK_HEIGHT = 18

CALENDAR_SPREADS = [
    "M 1-5",
    "M 5-9",
    "M 9-1",
    "RM 1-5",
    "RM 5-9",
    "RM 9-1",
    "Y 1-5",
    "Y 5-9",
    "Y 9-1",
    "OI 1-5",
    "OI 5-9",
    "OI 9-1",
    "P 1-5",
    "P 5-9",
    "P 9-1",
]

COMMODITY_SPREADS = [
    "M-RM 1",
    "M-RM 5",
    "M-RM 9",
    "Y-P 1",
    "Y-P 5",
    "Y-P 9",
    "OI-Y 1",
    "OI-Y 5",
    "OI-Y 9",
    "OI-P 1",
    "OI-P 5",
    "OI-P 9",
]

RATIO_SPREADS: list[str] = []


def setup_logger(log_file: Path) -> logging.Logger:
    logger = logging.getLogger("build_dashboard_workbook")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    handler = logging.FileHandler(log_file, encoding="utf-8")
    handler.setFormatter(logging.Formatter("[%(asctime)s] %(levelname)s %(message)s"))
    logger.addHandler(handler)
    return logger


def safe_filename(value: str) -> str:
    return (
        value.strip()
        .replace(" ", "_")
        .replace("/", "-")
        .replace("\\", "-")
        .replace(":", "-")
    )


def style_merged_header(sheet, cell_range: str, value: str, fill_color: str) -> None:
    sheet.merge_cells(cell_range)
    cell = sheet[cell_range.split(":")[0]]
    cell.value = value
    cell.font = Font(name="Noto Sans CJK SC", bold=True, size=14, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def apply_medium_outline(sheet, start_col: int, end_col: int, start_row: int, end_row: int) -> None:
    side = Side(style="medium", color="808080")
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = sheet.cell(row=row, column=col)
            left = side if col == start_col else cell.border.left
            right = side if col == end_col else cell.border.right
            top = side if row == start_row else cell.border.top
            bottom = side if row == end_row else cell.border.bottom
            cell.border = Border(left=left, right=right, top=top, bottom=bottom)


def add_chart_image(sheet, png_file: Path, spread_name: str, start_cell: str, start_row: int, start_col: int) -> None:
    title_range = (
        sheet.cell(row=start_row, column=start_col).coordinate,
        sheet.cell(row=start_row, column=start_col + 7).coordinate,
    )
    sheet.merge_cells(f"{title_range[0]}:{title_range[1]}")
    title_cell = sheet.cell(row=start_row, column=start_col)
    title_cell.value = spread_name
    title_cell.font = Font(name="Noto Sans CJK SC", bold=True, size=12)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[start_row].height = 22

    image = ExcelImage(str(png_file))
    aspect_ratio = image.height / image.width
    image.width = IMAGE_WIDTH_PX
    image.height = int(IMAGE_WIDTH_PX * aspect_ratio)
    sheet.add_image(image, start_cell)


def add_missing_note(sheet, message: str, start_row: int, start_col: int) -> None:
    start = sheet.cell(row=start_row, column=start_col).coordinate
    end = sheet.cell(row=start_row + 2, column=start_col + 7).coordinate
    sheet.merge_cells(f"{start}:{end}")
    cell = sheet.cell(row=start_row, column=start_col)
    cell.value = message
    cell.font = Font(name="Noto Sans CJK SC", bold=True, size=12, color="666666")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def add_category(
    *,
    sheet,
    category: str,
    spreads: list[str],
    png_by_spread: dict[str, Path],
    start_col: int,
    start_row: int,
    directory_rows: list[dict[str, str]],
) -> int:
    inserted = 0
    row = start_row
    for spread_name in spreads:
        png_file = png_by_spread.get(safe_filename(spread_name))
        if png_file:
            image_cell = sheet.cell(row=row + 1, column=start_col).coordinate
            add_chart_image(sheet, png_file, spread_name, image_cell, row, start_col)
            status = "inserted"
            inserted += 1
        else:
            add_missing_note(sheet, f"未找到图片：{spread_name}", row, start_col)
            status = "missing"
        directory_rows.append(
            {
                "category": category,
                "spread_name": spread_name,
                "png_path": str(png_file) if png_file else "",
                "status": status,
            }
        )
        row += BLOCK_HEIGHT
    return inserted


def build_dashboard(current_dir: Path, output_file: Path) -> dict[str, object]:
    png_files = sorted(current_dir.glob("*.png"))
    png_by_spread = {png_file.stem: png_file for png_file in png_files}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "价差图看板"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A3"

    directory = workbook.create_sheet("目录")
    directory.sheet_view.showGridLines = False

    report_date = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    style_merged_header(sheet, "A1:H1", "油脂油料月间价差", "1F4E79")
    style_merged_header(sheet, "J1:Q1", "油脂油料品种间价差", "385723")
    style_merged_header(sheet, "S1:Z1", "油粕比", "7F6000")
    style_merged_header(sheet, "A2:H2", f"报告日期：{report_date}", "D9EAF7")
    style_merged_header(sheet, "J2:Q2", f"报告日期：{report_date}", "E2F0D9")
    style_merged_header(sheet, "S2:Z2", f"报告日期：{report_date}", "FFF2CC")

    for cell in ["A2", "J2", "S2"]:
        sheet[cell].font = Font(name="Noto Sans CJK SC", bold=True, size=14, color="000000")

    for column_index in list(range(1, 9)) + list(range(10, 18)) + list(range(19, 27)):
        sheet.column_dimensions[get_column_letter(column_index)].width = 10
    for column_index in [9, 18]:
        sheet.column_dimensions[get_column_letter(column_index)].width = 3

    max_blocks = max(len(CALENDAR_SPREADS), len(COMMODITY_SPREADS), max(len(RATIO_SPREADS), 1))
    end_row = 3 + max_blocks * BLOCK_HEIGHT
    apply_medium_outline(sheet, 1, 8, 1, end_row)
    apply_medium_outline(sheet, 10, 17, 1, end_row)
    apply_medium_outline(sheet, 19, 26, 1, end_row)

    directory_rows: list[dict[str, str]] = []
    calendar_count = add_category(
        sheet=sheet,
        category="油脂油料月间价差",
        spreads=CALENDAR_SPREADS,
        png_by_spread=png_by_spread,
        start_col=1,
        start_row=4,
        directory_rows=directory_rows,
    )
    commodity_count = add_category(
        sheet=sheet,
        category="油脂油料品种间价差",
        spreads=COMMODITY_SPREADS,
        png_by_spread=png_by_spread,
        start_col=10,
        start_row=4,
        directory_rows=directory_rows,
    )

    ratio_count = 0
    if RATIO_SPREADS:
        ratio_count = add_category(
            sheet=sheet,
            category="油粕比",
            spreads=RATIO_SPREADS,
            png_by_spread=png_by_spread,
            start_col=19,
            start_row=4,
            directory_rows=directory_rows,
        )
    else:
        add_missing_note(sheet, "当前未生成油粕比图", 4, 19)
        directory_rows.append(
            {
                "category": "油粕比",
                "spread_name": "油粕比",
                "png_path": "",
                "status": "当前未生成油粕比图",
            }
        )

    classified = {safe_filename(spread_name) for spread_name in CALENDAR_SPREADS + COMMODITY_SPREADS + RATIO_SPREADS}
    unclassified = []
    for png_file in png_files:
        if png_file.stem not in classified:
            unclassified.append(png_file.stem)
            directory_rows.append(
                {
                    "category": "未分类",
                    "spread_name": png_file.stem,
                    "png_path": str(png_file),
                    "status": "unclassified",
                }
            )

    directory.append(["category", "spread_name", "png_path", "status"])
    for cell in directory[1]:
        cell.font = Font(name="Noto Sans CJK SC", bold=True, size=11, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in directory_rows:
        directory.append([row["category"], row["spread_name"], row["png_path"], row["status"]])
    directory.column_dimensions["A"].width = 24
    directory.column_dimensions["B"].width = 18
    directory.column_dimensions["C"].width = 96
    directory.column_dimensions["D"].width = 24

    workbook.save(output_file)
    return {
        "calendar_count": calendar_count,
        "commodity_count": commodity_count,
        "ratio_count": ratio_count,
        "unclassified": unclassified,
        "png_count": len(png_files),
    }


def main() -> int:
    project_root = Path(__file__).resolve().parents[1]
    reports_dir = project_root / "06_outputs" / "reports" / "current"
    current_dir = reports_dir / "价差日报图片"
    logs_dir = project_root / "10_logs"
    timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
    log_file = logs_dir / f"build_dashboard_workbook_{timestamp}.log"
    output_file = reports_dir / DASHBOARD_NAME

    reports_dir.mkdir(parents=True, exist_ok=True)
    current_dir.mkdir(parents=True, exist_ok=True)
    logs_dir.mkdir(parents=True, exist_ok=True)
    logger = setup_logger(log_file)

    try:
        result = build_dashboard(current_dir, output_file)
        logger.info("dashboard generated: %s result=%s", output_file, result)
        print(f"dashboard_path: {output_file}")
        print(f"calendar_count: {result['calendar_count']}")
        print(f"commodity_count: {result['commodity_count']}")
        print(f"ratio_count: {result['ratio_count']}")
        print(f"unclassified: {', '.join(result['unclassified']) if result['unclassified'] else ''}")
        print(f"build_dashboard_workbook_log: {log_file}")
        return 0
    except Exception as exc:  # noqa: BLE001
        logger.exception("dashboard build failed: %s", exc)
        print(f"build_dashboard_workbook failed: {exc}")
        print(f"build_dashboard_workbook_log: {log_file}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

